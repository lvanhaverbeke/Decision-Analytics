from pathlib import Path
import os, json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
ROOT=Path(__file__).resolve().parent
research_path=ROOT/'research.json'
research=json.loads(research_path.read_text()) if research_path.exists() else {}
wb=Workbook(); wb.remove(wb.active)
for name in ['Read me','Inputs','Investment','Annual cash flow','Room sizing','Scenarios','Sources']: wb.create_sheet(name)
ws=wb['Read me']
rows=[['PROJECT XL — student housing near VUB, Brussels'],['Decision model • illustrative development, not a property valuation'],['Status',research.get('status','Usable initial version — assumptions require verification')],['How to use','Edit blue cells on Inputs; Excel recalculates all formulas on opening.'],['Rent rule','Affordable rent = minimum of student budget, market ceiling and applicable regulatory/comparator ceiling.'],['Legal caveat','A market average is not a legal rent cap. Reference rent must be obtained for the exact dwelling.'],['Scenario','Development/major conversion, with land acquisition entered separately.'],['Model boundaries','Unlevered annual cash flow; debt-service and equity cash flow shown separately; no sale proceeds in payback.'],['Room choice','Room sizing compares every integer from 12–96 rooms; selects highest recurring cash flow among feasible candidates.'],['Feasible means','Total cost within budget, positive equity cash flow, DSCR target and assumed site capacity; return target shown separately.'],['Exclusions','Personal/company income tax, capital gains tax and sale proceeds; acquisition costs and VAT must be validated.'],['Sources','Source links and dates in Sources; all cost and financing inputs are illustrative unless marked otherwise.']]
for r in rows: ws.append(r)
a=wb['Inputs']
inputs=[('Number of rooms',40,'rooms','Illustrative baseline'),('Student monthly housing budget incl. charges',600,'EUR/month','User can replace with surveyed student budget'),('Monthly charges paid by student',100,'EUR/month','Separate from rent; assumed to cover utilities only'),('Local market rent comparator excl. charges',500,'EUR/month','Bru Kot room median EUR500 excl. charges; Brussels-wide listing data'),('Conservative reference-rent planning ceiling',500,'EUR/month','UNVERIFIED placeholder, NOT a legal cap; replace with address-specific assessment'),('Chosen base rent excl. charges',475,'EUR/month','Investor proposal'),('Occupancy',0.95,'%','Assumption; all 12 months are billable'),('Billable months',12,'months/year','Adjust for 10-month contracts'),('Land / property acquisition',900000,'EUR','Land for new build; avoid double counting existing structure'),('Gross floor area per room incl. common areas',34,'m²/room','Planning assumption, not minimum legal room area'),('Build cost per gross m²',2400,'EUR/m²','Assumed VAT-inclusive; validate with contractor'),('Furnishing per room',5000,'EUR/room','Assumption'),('Fixed professional / permits / acquisition costs',250000,'EUR','Includes placeholder closing costs; obtain breakdown'),('Contingency on build + furnishing',0.10,'%','Assumption'),('Operating expenses / collected rent',0.25,'%','Management, maintenance, insurance, tax provision; excludes recoverable utilities'),('Annual rent escalation',0.02,'%','Assumption, subject to indexation conditions'),('Annual operating cost escalation',0.025,'%','Assumption'),('Debt share of total investment',0.35,'%','Assumption'),('Debt annual interest',0.045,'%','Assumption, fixed amortising loan'),('Debt term',25,'years','Assumption'),('Total investment budget',5000000,'EUR','Illustrative cap, replace with actual budget'),('Available equity',3250000,'EUR','Illustrative cap'),('Minimum DSCR',1.20,'ratio','Illustrative lender criterion'),('Required unlevered net yield',0.045,'%','Illustrative target'),('Discount rate',0.06,'%','Unlevered cash flow NPV; terminal value excluded'),('Operating horizon',30,'years','No sale proceeds modelled; supported range 1–38 years'),('Development years before rent',2,'years','All investment assumed paid at time zero'),('Site maximum rooms',60,'rooms','UNVERIFIED capacity assumption; architect must validate'),('Net private room area',18,'m²','Illustrative design; 12 m² student dwelling minimum is not full planning clearance')]
a.append(['Input','Value','Units','Basis / status'])
for r in inputs:a.append(list(r))
i=wb['Investment']; i.append(['Measure','Value','Meaning'])
metrics=[('Affordable base rent','=MAX(0,MIN(Inputs!B3-Inputs!B4,Inputs!B5,Inputs!B6,Inputs!B7))','Student budget and comparator constrained'),('All-in student monthly payment','=B2+Inputs!B4','Rent + utility charge'),('Gross floor area','=Inputs!B2*Inputs!B11','m²'),('Build cost','=B4*Inputs!B12','EUR'),('Furnishing','=Inputs!B2*Inputs!B13','EUR'),('Contingency','=(B5+B6)*Inputs!B15','EUR'),('Total project investment','=Inputs!B10+B5+B6+B7+Inputs!B14','All-in development capital'),('Loan','=B8*Inputs!B19','EUR'),('Required equity','=B8-B9','EUR'),('Potential annual base rent','=Inputs!B2*B2*Inputs!B9','Excludes charges'),('Collected annual rent','=B11*Inputs!B8','Vacancy adjusted'),('Operating expenses','=B12*Inputs!B16','Year one'),('Net operating income (NOI)','=B12-B13','Before debt and income tax'),('Annual debt service','=IF(B9=0,0,IF(Inputs!B20=0,B9/Inputs!B21,-PMT(Inputs!B20,Inputs!B21,B9)))','Principal + interest'),('Annual equity cash flow','=B14-B15','Before investor income tax'),('Unlevered net yield','=B14/B8','NOI / all-in investment'),('Cash-on-cash equity return','=B16/B10','Annual equity cash flow / equity'),('DSCR','=IF(B15=0,999,B14/B15)','NOI / debt service'),('Static project payback years','=IF(B14>0,B8/B14,999)+Inputs!B28','Includes development delay; no escalation'),('Static equity payback years','=IF(B16>0,B10/B16,999)+Inputs!B28','Includes development delay; before income tax'),('Rent needed for required net yield','=B8*Inputs!B25/(Inputs!B2*Inputs!B9*Inputs!B8*(1-Inputs!B16))','Excl. charges; compare with affordable rent'),('Budget and debt feasibility','=IF(AND(B8<=Inputs!B22,B10<=Inputs!B23,B16>0,B19>=Inputs!B24,Inputs!B2<=Inputs!B29),"PASS","FAIL")','Yield target assessed separately'),('Required net yield met','=IF(B17>=Inputs!B25,"YES","NO")','Target is not a forecast'),('NPV excluding terminal value','=NPV(Inputs!B26,\'Annual cash flow\'!F3:F42)+\'Annual cash flow\'!F2','40-year cash flow range; inactive years set to zero'),('IRR excluding terminal value','=IFERROR(IRR(\'Annual cash flow\'!F2:F42),"No IRR")','Zero cash flow after horizon'),('Dynamic project payback','=IFERROR(MATCH(1,\'Annual cash flow\'!J2:J42,0)-1,"Beyond model horizon")','First end-year cumulative unlevered cash flow >= 0')]
for r in metrics:i.append(r)
c=wb['Annual cash flow']; c.append(['Year','Operating year','Rent','Operating costs','NOI','Project cash flow','Debt service','Equity cash flow','Cumulative project CF','Payback flag','Cumulative equity CF','Equity payback flag'])
for year in range(41):
 r=year+2;c.cell(r,1,year)
 c.cell(r,2,f'=MAX(0,A{r}-Inputs!B28)')
 c.cell(r,3,f'=IF(AND(B{r}>0,B{r}<=Inputs!B27),Investment!B12*(1+Inputs!B17)^(B{r}-1),0)')
 c.cell(r,4,f'=IF(AND(B{r}>0,B{r}<=Inputs!B27),Investment!B13*(1+Inputs!B18)^(B{r}-1),0)')
 c.cell(r,5,f'=C{r}-D{r}')
 c.cell(r,6,'=-Investment!B8' if year==0 else f'=E{r}')
 c.cell(r,7,f'=IF(AND(B{r}>0,B{r}<=Inputs!B21),Investment!B15,0)')
 c.cell(r,8,'=-Investment!B10' if year==0 else f'=E{r}-G{r}')
 c.cell(r,9,f'=F{r}' if year==0 else f'=I{r-1}+F{r}')
 c.cell(r,10,f'=IF(I{r}>=0,1,0)')
 c.cell(r,11,f'=H{r}' if year==0 else f'=K{r-1}+H{r}')
 c.cell(r,12,f'=IF(K{r}>=0,1,0)')
s=wb['Room sizing'];s.append(['Rooms','Floor area m²','Total cost','Equity','NOI','Debt service','Equity CF','Net yield','DSCR','Feasible','Rank value'])
for r,n in enumerate(range(12,97),2):
 s.cell(r,1,n)
 for col,f in {2:f'=A{r}*Inputs!B11',3:f'=Inputs!B10+Inputs!B14+(B{r}*Inputs!B12+A{r}*Inputs!B13)*(1+Inputs!B15)',4:f'=C{r}*(1-Inputs!B19)',5:f'=A{r}*Investment!B2*Inputs!B9*Inputs!B8*(1-Inputs!B16)',6:f'=IF(Inputs!B19=0,0,IF(Inputs!B20=0,C{r}*Inputs!B19/Inputs!B21,-PMT(Inputs!B20,Inputs!B21,C{r}*Inputs!B19)))',7:f'=E{r}-F{r}',8:f'=E{r}/C{r}',9:f'=IF(F{r}=0,999,E{r}/F{r})',10:f'=IF(AND(C{r}<=Inputs!B22,D{r}<=Inputs!B23,G{r}>0,I{r}>=Inputs!B24,A{r}<=Inputs!B29),"YES","NO")',11:f'=IF(J{r}="YES",G{r},-1)'}.items():s.cell(r,col,f)
s['A89']='Recommended rooms';s['B89']='=IF(MAX(K2:K86)<0,"No feasible candidate",INDEX(A2:A86,MATCH(MAX(K2:K86),K2:K86,0)))'
s['A90']='Scope';s['B90']='Best of tested sizes only; site capacity, planning consent and demand must be confirmed.'
t=wb['Scenarios'];t.append(['Scenario','Rent','Occupancy','Cost multiplier','NOI','Total investment','Net yield','Static project payback incl. delay'])
for r,(label,rent,occ,cost) in enumerate([('Downside',400,0.90,1.15),('Base',475,0.95,1),('Upside',500,0.98,0.95),('VUB all-in match',349.30,0.95,1)],2):
 t.append([label,rent,occ,cost,f'=Inputs!B2*MIN(B{r},Inputs!B3-Inputs!B4,Inputs!B5,Inputs!B6)*Inputs!B9*C{r}*(1-Inputs!B16)',f'=Investment!B8*D{r}',f'=E{r}/F{r}',f'=F{r}/E{r}+Inputs!B28'])
z=wb['Sources'];z.append(['Item','Source / URL','Finding','Accessed / status'])
for row in research.get('sources',[]):z.append(row)
if z.max_row==1:z.append(['All numerical inputs','Illustrative assumptions','Research pending; do not interpret as verified Brussels rent or cost data','2026-10-14'])
for sheet in wb:
 sheet.freeze_panes='B2';sheet.auto_filter.ref=sheet.dimensions
 for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='17365D');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True)
 sheet.row_dimensions[1].height=32
 for row in sheet.iter_rows(min_row=2):
  for cell in row:
   cell.alignment=Alignment(vertical='top',wrap_text=True)
   if cell.data_type=='f':cell.font=Font(color='111111');cell.number_format='#,##0.00'
 for col in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(col)].width=23 if col>1 else 43
 if sheet.title in ['Read me','Sources']:sheet.column_dimensions['B'].width=82;sheet.column_dimensions['C'].width=65
 if sheet.title=='Inputs':
  sheet.column_dimensions['D'].width=70
  for r in range(2,sheet.max_row+1):sheet.cell(r,2).font=Font(color='0000FF');sheet.cell(r,2).fill=PatternFill('solid',fgColor='EAF2F8')
  for r in [8,15,16,17,18,19,20,25,26]:sheet.cell(r,2).number_format='0.0%'
for r in [17,18]:i.cell(r,2).number_format='0.0%'
for r in range(2,87):s.cell(r,8).number_format='0.0%'
for r in range(2,6):t.cell(r,7).number_format='0.0%'
from openpyxl.worksheet.datavalidation import DataValidation
dv=DataValidation(type='whole',operator='between',formula1=1,formula2=38);dv.errorTitle='Horizon range';dv.error='Use 1 to 38 operating years';dv.showErrorMessage=True;a.add_data_validation(dv);dv.add(a['B27'])
wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
s['A91']='Minimum feasible rooms';s['B91']='=IF(COUNTIF(J2:J86,"YES")=0,"No feasible candidate",MATCH("YES",J2:J86,0)+11)'
s['A92']='Investment recommendation';s['B92']='Conditional screening result only: target yield, address rent and planning require separate approval.'
i['A28']='Maximum investment at target yield';i['B28']='=B14/Inputs!B25'
i['A29']='Cost reduction needed for target yield';i['B29']='=MAX(0,B8-B28)'
i['A30']='Unlevered yield at ten collected months';i['B30']='=B14*10/Inputs!B9/B8';i['B30'].number_format='0.0%'
i['A32']='Dynamic equity payback';i['B32']="=IFERROR(MATCH(1,'Annual cash flow'!L2:L42,0)-1,\"Beyond model horizon\")"
i['A31']='Legal reference verified?';i['B31']='NO — exact address and dwelling details missing'
tmp=ROOT/'model.tmp.xlsx';wb.save(tmp);load_workbook(tmp);os.replace(tmp,ROOT/'model.xlsx')
print('Saved model.xlsx atomically');print('SAVE_NOW:',(ROOT/'SAVE_NOW.txt').exists())
