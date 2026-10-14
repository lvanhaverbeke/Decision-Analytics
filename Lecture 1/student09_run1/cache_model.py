from pathlib import Path
import re, math, json, zipfile, os
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
P=Path(__file__).resolve().parent
wb=load_workbook(P/'model.xlsx');memo={}
def flat(x):
 return [v for item in x for v in (item if isinstance(item,list) else [item])]
def npv(rate,values):return sum(v/(1+rate)**(j+1) for j,v in enumerate(values))
def irr(values):
 lo,hi=-0.5,1.0
 def value(rate):return sum(v/(1+rate)**j for j,v in enumerate(values))
 if value(lo)*value(hi)>0:return 'No IRR'
 for _ in range(150):
  mid=(lo+hi)/2
  if value(mid)>0:lo=mid
  else:hi=mid
 return (lo+hi)/2
def calc(sheet,addr):
 key=(sheet,addr)
 if key in memo:return memo[key]
 v=wb[sheet][addr].value
 if not isinstance(v,str) or not v.startswith('='):return 0 if v is None else v
 def rr(sh,start,end):
  from openpyxl.utils.cell import range_boundaries
  a,b,c,d=range_boundaries(start+':'+end)
  return [calc(sh,wb[sh].cell(r,col).coordinate) for r in range(b,d+1) for col in range(a,c+1)]
 def repl(m):
  sh=m.group(1) or m.group(2) or sheet;start=m.group(3);end=m.group(4)
  return f'R({sh!r},{start!r},{end!r})' if end else f'C({sh!r},{start!r})'
 expr=v[1:];parts=re.split(r'("[^"]*")',expr)
 pat=r"(?<![A-Za-z0-9_])(?:(?:'([^']+)'|([A-Za-z][A-Za-z ]*))!)?(\$?[A-Z]{1,3}\$?\d+)(?::(\$?[A-Z]{1,3}\$?\d+))?"
 for j in range(0,len(parts),2):
  parts[j]=re.sub(pat,repl,parts[j]).replace('^','**').replace('<>','!=')
  parts[j]=re.sub(r'(?<![<>=!])=(?!=)','==',parts[j])
 expr=''.join(parts)
 env={'C':calc,'R':rr,'MAX':lambda *v:max(flat(v)),'MIN':lambda *v:min(flat(v)),'IF':lambda a,b,c:b if a else c,'AND':lambda *v:all(v),'PMT':lambda r,n,p:-p/n if r==0 else -p*r/(1-(1+r)**(-n)),'NPV':npv,'IRR':irr,'IFERROR':lambda a,b:a,'MATCH':lambda val,arr,mode:arr.index(val)+1,'INDEX':lambda arr,index:arr[index-1],'COUNTIF':lambda arr,v:arr.count(v)}
 try: result=eval(expr,{'__builtins__':{}},env)
 except Exception as e:
  fallback=re.search(r',"([^"]*)"\)$',v) if v.startswith('=IFERROR(') else None
  if fallback:result=fallback.group(1)
  else:raise RuntimeError(f'{sheet}!{addr} {v} -> {expr}: {e}')
 memo[key]=result;return result
for sh in wb:
 for row in sh:
  for c in row:
   if c.data_type=='f':calc(sh.title,c.coordinate)
ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
tmp=P/'model.cached.tmp.xlsx'
with zipfile.ZipFile(P/'model.xlsx') as src,zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as dst:
 for info in src.infolist():
  data=src.read(info.filename)
  match=re.fullmatch(r'xl/worksheets/sheet(\d+)\.xml',info.filename)
  if match:
   sh=wb.worksheets[int(match.group(1))-1].title;root=ET.fromstring(data)
   for cell in root.findall('.//m:c',ns):
    if cell.find('m:f',ns) is not None:
     value=memo[(sh,cell.attrib['r'])];cv=cell.find('m:v',ns)
     if cv is None:cv=ET.SubElement(cell,'{'+ns['m']+'}v')
     if isinstance(value,str):cell.set('t','str');cv.text=value
     else:cell.pop('t',None) if False else None;cell.attrib.pop('t',None);cv.text=str(value)
   data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
  dst.writestr(info,data)
load_workbook(tmp,data_only=True);os.replace(tmp,P/'model.xlsx')
summary={c[0]:calc('Investment',f'B{r}') for r,c in enumerate(list(wb['Investment'].values)[1:],2)}
summary['Recommended rooms']=calc('Room sizing','B89');summary['Minimum feasible rooms']=calc('Room sizing','B91')
(P/'results.json').write_text(json.dumps(summary,indent=2))
print(json.dumps(summary,indent=2));print('SAVE_NOW:',(P/'SAVE_NOW.txt').exists())
