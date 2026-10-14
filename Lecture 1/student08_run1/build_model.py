from pathlib import Path
import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'model.xlsx'

def save_atomic(wb):
    tmp = ROOT / 'model.tmp.xlsx'
    wb.save(tmp)
    check = load_workbook(tmp)
    assert 'Invoer' in check.sheetnames
    check.close()
    os.replace(tmp, OUT)

def base():
    wb = Workbook(); w = wb.active; w.title = 'Invoer'
    rows = [
        ['Studentenhuisvesting – investeringsmodel'],
        ['Alle bedragen zijn fictieve, aanpasbare voorbeelden in EUR. Geen marktinschatting.'],
        ['Parameter','Waarde','Toelichting'],
        ['Aantal kamers',10,'Voorbeeld'],
        ['Studentenhuur per kamer per maand',600,'Voorbeeld'],
        ['Jaarlijkse huurverhoging',0.02,'Onbekend in de notities; fictieve aanname'],
        ['Maanden studentenhuur',8,'Gebaseerd op notities'],
        ['Maanden Airbnb',2,'Gebaseerd op notities'],
        ['Bezetting studentenhuur',0.95,'Inclusief leegstand bij wissels / exchange'],
        ['Airbnb prijs per nacht',75,'Voorbeeld'],
        ['Airbnb bezetting',0.6,'Voorbeeld'],
        ['Dagen per maand',30,'Rekenconventie'],
        ['Grondkosten',200000,'Eenmalig'],
        ['Constructiekosten',400000,'Eenmalig, exclusief aankoop gebouw'],
        ['Aankoop bestaand gebouw',0,'Voorkom dubbeltelling met constructie'],
        ['Opstart exploitatie',10000,'Eenmalig'],
        ['Belastingen aankoop / bouw',20000,'Eenmalig'],
        ['Onderhoud bij start',5000,'Eenmalig'],
        ['Jaarlijkse exploitatiekosten',12000,'Exclusief onderhoud en belasting'],
        ['Jaarlijks onderhoud',6000,'Terugkerend'],
        ['Jaarlijkse vastgoedbelasting',2000,'Terugkerend'],
        ['Overige diensten: jaaromzet',2000,'Was, schoonmaak, afhaalpunt en bagage'],
    ]
    for row in rows: w.append(row)
    s=wb.create_sheet('Samenvatting')
    for row in [['Resultaat – basisversie'],['Initiële investering','=SUM(Invoer!B13:B18)'],['Studentenhuur jaar 1','=Invoer!B4*Invoer!B5*Invoer!B7*Invoer!B9'],['Airbnb jaar 1','=Invoer!B4*Invoer!B8*Invoer!B12*Invoer!B10*Invoer!B11'],['Totale inkomsten','=SUM(B3:B4)+Invoer!B22'],['Exploitatiekosten','=SUM(Invoer!B19:B21)'],['Operationele kasstroom','=B5-B6']]:s.append(row)
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
    for sh in wb:
        sh.column_dimensions['A'].width=42;sh.column_dimensions['B'].width=23;sh.column_dimensions['C'].width=76
        sh.freeze_panes='B4'
        for cell in sh[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17365D')
        for row in sh.iter_rows(min_row=4):
            for c in row:
                if c.column==2:c.number_format='#,##0.00'
    save_atomic(wb)
    return wb

if __name__=='__main__':
    base()
    print('Basisversie atomair opgeslagen als model.xlsx')
    print('SAVE_NOW:',(ROOT/'SAVE_NOW.txt').exists())
