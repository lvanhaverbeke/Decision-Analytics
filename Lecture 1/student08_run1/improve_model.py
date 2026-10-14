"""Maak het aanpasbare investeringsmodel; uitsluitend standaardbibliotheek + openpyxl."""
from pathlib import Path
import os, math, zipfile, xml.etree.ElementTree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties
ROOT=Path(__file__).resolve().parent
if (ROOT/'SAVE_NOW.txt').exists():
    print('SAVE_NOW aanwezig; bestaande basisversie behouden.'); raise SystemExit
wb=Workbook(); wb.remove(wb.active)
for name in ['Leeswijzer','Invoer','Samenvatting','Kasstromen','Scenariovergelijking']:
    wb.create_sheet(name)
cache={}
def put(s,cell,v,result=None):
    wb[s][cell]=v
    if isinstance(v,str) and v.startswith('=') and result is not None:cache[(s,cell)]=result

def title(s,t):put(s,'A1',t)
notes=[
 'Studentenhuisvesting | haalbaarheid en rendement',
 'Doel: transparant rekenmodel op basis van de aangeleverde notities.',
 'Alle bedragen en percentages behalve de genoemde 8 + 2 maanden zijn fictieve voorbeelden.',
 'Valuta: EUR. Bedragen moeten consequent exclusief of inclusief niet-recupereerbare btw worden ingevoerd.',
 'Blauwe cellen op Invoer zijn aanpasbaar. Zwarte cellen bevatten formules.',
 'Keuze 1 = 12 maanden studentenhuur; keuze 2 = 8 maanden studentenhuur + 2 maanden Airbnb; keuze 3 = vrije mix.',
 'De resterende maanden leveren geen kamerinkomsten op; bezettingsgraden verwerken leegstand binnen actieve maanden.',
 'Exchange en wissels: verwerk leegstand in studentenbezetting; tel dezelfde leegstand niet nogmaals als kosten.',
 'Grond, aankoop bestaand gebouw en nieuwbouw zijn afzonderlijke posten. Vul overlappende bedragen niet dubbel in.',
 'Opstart en aanvangsonderhoud zijn eenmalig. Jaarlijkse exploitatie, onderhoud en vastgoedbelasting zijn terugkerend.',
 'Diensten: was, schoonmaak, afhaalpunt (interpretatie onduidelijke tekst) en bagageopslag.',
 'Diensten worden per verhuurde kamermaand geraamd, ook bij Airbnb; pas besteding aan de werkelijke vraag aan.',
 'Schoonmaakomzet betreft een verkochte dienst. Airbnb-wissels hebben daarnaast eigen schoonmaakkosten.',
 'Airbnb-platformkosten gelden alleen voor overnachtingsomzet; beheer geldt voor alle inkomsten.',
 'Kasstromen: ongeleend, vóór inkomsten-/vennootschapsbelasting; geen lening, rente, afschrijving of btw-afrekening.',
 'Vastgoedbelasting en eenmalige aankoop-/bouwbelasting zijn expliciete invoer, zonder wettelijke tariefaanname.',
 'Horizon: 20 jaar. Investering op t=0; exploitatie en verkoop aan het einde van elk jaar.',
 'Restwaarde = grond + aankoop gebouw + constructie, met jaarlijkse waardegroei; verkoopkosten worden afgetrokken.',
 'NCW/NPV gebruikt de discontovoet. IRR is het projectrendement inclusief verkoop; terugverdientijd telt volle jaren.',
 'Scenario’s gebruiken dezelfde investering en vaste kosten; verschillen komen door verhuurmaanden en bezetting.',
 'Geen markt- of juridisch onderzoek verricht. Vergunningen, contractregels en fiscale behandeling zijn niet beoordeeld.',
 'Bron: gebruikerstekst met onduidelijke namen, huurverhoging en omschrijving van het afhaalpunt.',
 'Code: build_model.py bewaart de basis; improve_model.py genereert deze uitgebreide versie.',
]
for r,t in enumerate(notes,1):put('Leeswijzer',f'A{r}',t)
title('Invoer','Aanpasbare aannames | fictief voorbeeld')
for c,v in zip('ABCD',['Parameter','Waarde','Eenheid','Toelichting']):put('Invoer',f'{c}3',v)
items=[
 ('keuze','Verhuurscenario',2,'1 / 2 / 3','1: student 12m; 2: student 8m + Airbnb 2m; 3: vrije mix'),
 ('kamers','Aantal kamers',10,'kamers','Identieke kamers als vereenvoudiging'),
 ('huur','Studentenhuur',600,'EUR/kamer/maand','Fictief startbedrag'),
 ('huurgroei','Jaarlijkse huurverhoging',.02,'%','Niet gespecificeerd in notities'),
 ('studentmaand','Studentmaanden vrije mix',8,'maanden','Alleen gebruikt bij keuze 3'),
 ('airmaand','Airbnb-maanden vrije mix',2,'maanden','Alleen gebruikt bij keuze 3'),
 ('studentbezet','Bezetting studentenhuur',.95,'%','Inclusief exchange en gaten tussen huurders'),
 ('nacht','Airbnb-nachtprijs',75,'EUR/kamer/nacht','Fictief'),
 ('airbezet','Airbnb-bezetting',.60,'%','Binnen voor Airbnb beschikbare maanden'),
 ('dagen','Dagen per maand',30,'dagen','Vaste rekenconventie'),
 ('airgroei','Jaarlijkse groei Airbnb-prijs',.02,'%','Fictief'),
 ('verblijf','Gemiddeld Airbnb-verblijf',4,'nachten','Bepaalt aantal schoonmaakwissels'),
 ('platform','Airbnb-platformkosten',.03,'% omzet Airbnb','Fictief'),
 ('wisselkost','Schoonmaakkost per Airbnb-wissel',20,'EUR/wissel','Apart van verkochte schoonmaakdienst'),
 ('land','Grond / grondkosten',200000,'EUR','Eenmalig'),
 ('bouw','Constructie / nieuwbouw',400000,'EUR','Eenmalig; exclusief bestaand gebouw'),
 ('gebouw','Aankoop bestaand gebouw',0,'EUR','Eenmalig; voorkom overlap met constructie'),
 ('opstart','Opstart exploitatie',10000,'EUR','Eenmalige operating fees'),
 ('starttax','Aankoop-/bouwbelasting',20000,'EUR','Eenmalig; zelf invullen'),
 ('startonderhoud','Onderhoud bij aanvang',5000,'EUR','Eenmalig'),
 ('vast','Vaste exploitatiekosten per jaar',12000,'EUR/jaar','Energie, verzekering etc.; exclusief apart genoemde kosten'),
 ('onderhoud','Jaarlijks onderhoud',6000,'EUR/jaar','Terugkerend'),
 ('vastgoedtax','Jaarlijkse vastgoedbelasting',2000,'EUR/jaar','Terugkerend; fictief bedrag'),
 ('beheer','Vastgoedbeheer',.08,'% totale omzet','Vergoeding beheerder'),
 ('kostgroei','Jaarlijkse kostenstijging',.02,'%','Geldt ook voor wisselkost'),
 ('was','Was: omzet per bezette kamermaand',10,'EUR','Gemiddelde besteding'),
 ('wasmarge','Was: directe kosten',.40,'% dienstomzet','Fictief'),
 ('schoon','Schoonmaak: omzet per bezette kamermaand',15,'EUR','Verkochte extra dienst'),
 ('schoonmarge','Schoonmaak: directe kosten',.60,'% dienstomzet','Fictief'),
 ('post','Afhaalpunt: omzet per bezette kamermaand',2,'EUR','Interpretatie van onduidelijke notitie'),
 ('postmarge','Afhaalpunt: directe kosten',.20,'% dienstomzet','Fictief'),
 ('bagage','Bagage: omzet per bezette kamermaand',3,'EUR','Extra dienst uit notities'),
 ('bagagemarge','Bagage: directe kosten',.10,'% dienstomzet','Fictief'),
 ('dienstgroei','Jaarlijkse groei diensttarieven',.02,'%','Fictief'),
 ('discount','Discontovoet',.07,'%','Vereist projectrendement; fictief'),
 ('waardegroei','Jaarlijkse vastgoedwaardegroei',.01,'%','Fictief; geen taxatie'),
 ('verkoop','Verkoopkosten',.03,'% verkoopwaarde','Fictief'),
]
row={};v={}
for r,(key,label,value,unit,note) in enumerate(items,4):
    row[key]=r;v[key]=value
    for c,x in zip('ABCD',[label,value,unit,note]):put('Invoer',f'{c}{r}',x)
    wb['Invoer'][f'B{r}'].font=Font(color='0000FF')
def ref(k):return f'Invoer!$B${row[k]}'
R=ref
for key in v:
    typ='whole' if key in ['keuze','kamers','studentmaand','airmaand','dagen'] else 'decimal'
    lo=1 if key in ['keuze','kamers','dagen','verblijf'] else (-.99 if key in ['huurgroei','airgroei','kostgroei','dienstgroei','waardegroei'] else 0)
    hi=3 if key=='keuze' else (12 if key in ['studentmaand','airmaand'] else (1 if items[row[key]-4][3].startswith('%') and key not in ['huurgroei','airgroei','kostgroei','dienstgroei','waardegroei'] else 1000000000))
    dv=DataValidation(type=typ,operator='between',formula1=lo,formula2=hi);dv.errorTitle='Ongeldige invoer';dv.error='Controleer het toegestane bereik.';dv.showErrorMessage=True;dv.errorStyle='stop';wb['Invoer'].add_data_validation(dv);dv.add(f'B{row[key]}')
    wb['Invoer'][f'B{row[key]}'].number_format='0.0%' if items[row[key]-4][3].startswith('%') else '#,##0.00'
inv=sum(v[k] for k in ['land','bouw','gebouw','opstart','starttax','startonderhoud'])
property0=sum(v[k] for k in ['land','bouw','gebouw'])
def calc(sm,am):
    nights=v['kamers']*am*v['dagen']*v['airbezet'];cm=v['kamers']*(sm*v['studentbezet']+am*v['airbezet'])
    out=[]
    for y in range(1,21):
        rent=v['kamers']*sm*v['studentbezet']*v['huur']*(1+v['huurgroei'])**(y-1)
        air=nights*v['nacht']*(1+v['airgroei'])**(y-1)
        services=[cm*v[k]*(1+v['dienstgroei'])**(y-1) for k in ['was','schoon','post','bagage']]
        revenue=rent+air+sum(services)
        cost=(v['vast']+v['onderhoud']+v['vastgoedtax']+nights/v['verblijf']*v['wisselkost'])*(1+v['kostgroei'])**(y-1)+air*v['platform']+revenue*v['beheer']+sum(s*v[k] for s,k in zip(services,['wasmarge','schoonmarge','postmarge','bagagemarge']))
        sale=property0*(1+v['waardegroei'])**20*(1-v['verkoop']) if y==20 else 0
        out.append((revenue,cost,revenue-cost,sale))
    flows=[-inv]+[x[2]+x[3] for x in out]
    npv=sum(x/(1+v['discount'])**i for i,x in enumerate(flows))
    low,high=-.99,10
    for _ in range(150):
        mid=(low+high)/2;pv=sum(x/(1+mid)**i for i,x in enumerate(flows))
        if pv>0:low=mid
        else:high=mid
    return out,flows,npv,(low+high)/2
sm=12 if v['keuze']==1 else 8 if v['keuze']==2 else v['studentmaand']
am=0 if v['keuze']==1 else 2 if v['keuze']==2 else v['airmaand']
data,flows,npv,irr=calc(sm,am)
title('Samenvatting','Studentenhuisvesting | investeringsbeslissing')
summary=[('Scenario',f'={R("keuze")}',v['keuze']),('Studentmaanden',f'=IF(B3=1,12,IF(B3=2,8,{R("studentmaand")}))',sm),('Airbnb-maanden',f'=IF(B3=1,0,IF(B3=2,2,{R("airmaand")}))',am),('Niet ingeplande maanden','=12-B4-B5',12-sm-am),('Controle kalender','=IF(AND(B4>=0,B5>=0,B4+B5<=12),"OK","FOUT: totaal > 12 maanden")','OK'),('Initiële investering','='+ '+'.join(R(k) for k in ['land','bouw','gebouw','opstart','starttax','startonderhoud']),inv),('Omzet jaar 1','=Kasstromen!C13',data[0][0]),('Kosten jaar 1','=Kasstromen!C24',data[0][1]),('Operationele kasstroom jaar 1','=Kasstromen!C25',data[0][2]),('Aanvangsrendement','=IFERROR(B11/B8,0)',data[0][2]/inv),('Netto contante waarde',f'=NPV({R("discount")},Kasstromen!C28:V28)+Kasstromen!B28',npv),('Intern rendement (IRR)','=IFERROR(IRR(Kasstromen!B28:V28),"Niet berekenbaar")',irr),('Terugverdientijd incl. verkoop (jaar)','=IFERROR(MATCH(1,Kasstromen!B31:V31,0)-1,"Niet binnen 20 jaar")',next((i for i in range(21) if sum(flows[:i+1])>=0),'Niet binnen 20 jaar')),('Netto verkoopopbrengst jaar 20','=Kasstromen!V27',data[-1][3]),('Rekenhorizon',20,20)]
for rr,(label,f,val) in enumerate(summary,3):put('Samenvatting',f'A{rr}',label);put('Samenvatting',f'B{rr}',f,val)
put('Samenvatting','A21','Let op: fictieve invoer. NCW en IRR omvatten de verkoop aan het einde van jaar 20.')
put('Samenvatting','A22','Leegstand verlaagt omzet; vaste kosten lopen door. Kosten staan positief, investering negatief.')
put('Samenvatting','A23','Bij een kalenderfout eerst de vrije mix corrigeren; resultaten zijn dan niet bruikbaar.')
title('Kasstromen','Kasstromen | EUR | eindejaarsconventie')
labels={3:'Jaar',4:'Bezet: student-kamermaanden',5:'Bezet: Airbnb-kamernachten',6:'Bezet: totale kamermaanden',7:'Studentenhuur',8:'Airbnb-overnachtingen',9:'Wasdiensten',10:'Schoonmaakdiensten',11:'Afhaalpunt',12:'Bagageopslag',13:'Totale omzet',15:'Vaste exploitatiekosten',16:'Onderhoud',17:'Vastgoedbelasting',18:'Vastgoedbeheer',19:'Airbnb-platform',20:'Airbnb-wissels / schoonmaak',21:'Was: directe kosten',22:'Schoonmaak: directe kosten',23:'Afhaalpunt + bagage: directe kosten',24:'Totale kosten',25:'Operationele kasstroom',26:'Investering',27:'Netto verkoopopbrengst',28:'Netto kasstroom',29:'Contante waarde kasstroom',30:'Cumulatieve kasstroom',31:'Terugverdiend (1=ja)'}
for rr,label in labels.items():put('Kasstromen',f'A{rr}',label)
from openpyxl.utils import get_column_letter as L
cumulative=-inv
for y in range(21):
    c=L(y+2);prev=L(y+1);put('Kasstromen',f'{c}3',y)
    if y==0:
        for rr in [4,5,6,7,8,9,10,11,12,13,15,16,17,18,19,20,21,22,23,24,25,27]:put('Kasstromen',f'{c}{rr}',0)
        for rr in [26,28,29,30]:put('Kasstromen',f'{c}{rr}','=-Samenvatting!B8',-inv)
        put('Kasstromen',f'{c}31',f'=IF({c}30>=0,1,0)',0);continue
    cm=v['kamers']*(sm*v['studentbezet']+am*v['airbezet']);nights=v['kamers']*am*v['dagen']*v['airbezet'];g=(1+v['kostgroei'])**(y-1)
    forms={4:(f'={R("kamers")}*Samenvatting!$B$4*{R("studentbezet")}',v['kamers']*sm*v['studentbezet']),5:(f'={R("kamers")}*Samenvatting!$B$5*{R("dagen")}*{R("airbezet")}',nights),6:(f'={c}4+{c}5/{R("dagen")}',cm),7:(f'={c}4*{R("huur")}*(1+{R("huurgroei")})^({c}$3-1)',v['kamers']*sm*v['studentbezet']*v['huur']*(1+v['huurgroei'])**(y-1)),8:(f'={c}5*{R("nacht")}*(1+{R("airgroei")})^({c}$3-1)',nights*v['nacht']*(1+v['airgroei'])**(y-1))}
    for rr,key in zip(range(9,13),['was','schoon','post','bagage']):forms[rr]=(f'={c}6*{R(key)}*(1+{R("dienstgroei")})^({c}$3-1)',cm*v[key]*(1+v['dienstgroei'])**(y-1))
    forms[13]=(f'=SUM({c}7:{c}12)',data[y-1][0])
    for rr,key in zip([15,16,17],['vast','onderhoud','vastgoedtax']):forms[rr]=(f'={R(key)}*(1+{R("kostgroei")})^({c}$3-1)',v[key]*g)
    forms[18]=(f'={c}13*{R("beheer")}',data[y-1][0]*v['beheer']);forms[19]=(f'={c}8*{R("platform")}',forms[8][1]*v['platform']);forms[20]=(f'={c}5/{R("verblijf")}*{R("wisselkost")}*(1+{R("kostgroei")})^({c}$3-1)',nights/v['verblijf']*v['wisselkost']*g)
    for rr,sr,key in [(21,9,'wasmarge'),(22,10,'schoonmarge')]:forms[rr]=(f'={c}{sr}*{R(key)}',forms[sr][1]*v[key])
    forms[23]=(f'={c}11*{R("postmarge")}+{c}12*{R("bagagemarge")}',forms[11][1]*v['postmarge']+forms[12][1]*v['bagagemarge'])
    forms[24]=(f'=SUM({c}15:{c}23)',data[y-1][1]);forms[25]=(f'={c}13-{c}24',data[y-1][2]);forms[26]=('=0',0)
    forms[27]=(f'=IF({c}$3=20,({R("land")}+{R("bouw")}+{R("gebouw")})*(1+{R("waardegroei")})^{c}$3*(1-{R("verkoop")}),0)',data[y-1][3]);forms[28]=(f'=SUM({c}25:{c}27)',flows[y]);forms[29]=(f'={c}28/(1+{R("discount")})^{c}$3',flows[y]/(1+v['discount'])**y)
    cumulative+=flows[y];forms[30]=(f'={prev}30+{c}28',cumulative);forms[31]=(f'=IF({c}30>=0,1,0)',int(cumulative>=0))
    for rr,(f,val) in forms.items():put('Kasstromen',f'{c}{rr}',f,val)
if (ROOT/'SAVE_NOW.txt').exists():
    print('SAVE_NOW aanwezig; basisversie behouden.');raise SystemExit
title('Scenariovergelijking','Vergelijking | identieke investering en aannames')
for rr,label in {3:'Scenario',4:'Studentmaanden',5:'Airbnb-maanden',6:'Niet ingeplande maanden',7:'Omzet jaar 1',8:'Kosten jaar 1',9:'Operationele kasstroom jaar 1',10:'Netto contante waarde',11:'IRR incl. verkoop',12:'Cumulatieve kasstroom na 20 jaar',13:'Kalendercontrole',16:'Jaar / netto kasstroom',40:'Jaar / cumulatieve kasstroom'}.items():put('Scenariovergelijking',f'A{rr}',label)
for col,sm2,am2,name in [('B',12,0,'12 maanden student'),('C',8,2,'8 student + 2 Airbnb'),('D',v['studentmaand'],v['airmaand'],'Vrije mix')]:
    dd,ff,nn,ii=calc(sm2,am2);put('Scenariovergelijking',f'{col}3',name)
    put('Scenariovergelijking',f'{col}4',f'={R("studentmaand")}' if col=='D' else sm2,sm2)
    put('Scenariovergelijking',f'{col}5',f'={R("airmaand")}' if col=='D' else am2,am2)
    put('Scenariovergelijking',f'{col}6',f'=12-{col}4-{col}5',12-sm2-am2)
    cmexpr=f'{R("kamers")}*({col}$4*{R("studentbezet")}+{col}$5*{R("airbezet")})'
    nightsexpr=f'{R("kamers")}*{col}$5*{R("dagen")}*{R("airbezet")}'
    def scenario_components(yr):
        student=f'{R("kamers")}*{col}$4*{R("studentbezet")}*{R("huur")}*(1+{R("huurgroei")})^({yr}-1)'
        air=f'{nightsexpr}*{R("nacht")}*(1+{R("airgroei")})^({yr}-1)'
        svc=f'({cmexpr})*({R("was")}+{R("schoon")}+{R("post")}+{R("bagage")})*(1+{R("dienstgroei")})^({yr}-1)'
        rev=f'({student}+{air}+{svc})'
        servicecost=f'({cmexpr})*({R("was")}*{R("wasmarge")}+{R("schoon")}*{R("schoonmarge")}+{R("post")}*{R("postmarge")}+{R("bagage")}*{R("bagagemarge")})*(1+{R("dienstgroei")})^({yr}-1)'
        cost=f'(({R("vast")}+{R("onderhoud")}+{R("vastgoedtax")}+({nightsexpr})/{R("verblijf")}*{R("wisselkost")} )*(1+{R("kostgroei")})^({yr}-1)+({air})*{R("platform")}+{rev}*{R("beheer")}+{servicecost})'
        return rev,cost
    re,co=scenario_components('1')
    put('Scenariovergelijking',f'{col}7','='+re,dd[0][0]);put('Scenariovergelijking',f'{col}8','='+co,dd[0][1]);put('Scenariovergelijking',f'{col}9',f'={col}7-{col}8',dd[0][2])
    put('Scenariovergelijking',f'{col}10',f'=NPV({R("discount")},{col}18:{col}37)+{col}17',nn)
    put('Scenariovergelijking',f'{col}11',f'=IFERROR(IRR({col}17:{col}37),"Niet berekenbaar")',ii)
    put('Scenariovergelijking',f'{col}12',f'=SUM({col}17:{col}37)',sum(ff))
    put('Scenariovergelijking',f'{col}13',f'=IF({col}4+{col}5<=12,"OK","FOUT: totaal > 12 maanden")','OK' if sm2+am2<=12 else 'FOUT: totaal > 12 maanden')
    cum=0
    for y in range(21):
        rr=17+y;put('Scenariovergelijking',f'A{rr}',y)
        if y==0:f='=-Samenvatting!B8'
        else:
            re,co=scenario_components(f'$A{rr}')
            f=f'={re}-{co}+IF($A{rr}=20,Kasstromen!$V$27,0)'
        put('Scenariovergelijking',f'{col}{rr}',f,ff[y]);cum+=ff[y]
        put('Scenariovergelijking',f'A{41+y}',y)
        put('Scenariovergelijking',f'{col}{41+y}',f'=SUM({col}$17:{col}{rr})',cum)
# Grafieken voor snelle interpretatie.
chart=LineChart();chart.title='Cumulatieve kasstroom per scenario';chart.y_axis.title='EUR';chart.x_axis.title='Jaar'
for col in range(2,5):
    chart.add_data(Reference(wb['Scenariovergelijking'],min_col=col,min_row=41,max_row=61))
    chart.series[-1].tx=__import__('openpyxl').chart.series.SeriesLabel(strRef=__import__('openpyxl').chart.data_source.StrRef(f=f"'Scenariovergelijking'!${L(col)}$3"))
chart.set_categories(Reference(wb['Scenariovergelijking'],min_col=1,min_row=41,max_row=61));chart.width=25;chart.height=11
wb['Samenvatting'].add_chart(chart,'D3')
bar=BarChart();bar.title='Operationele kasstroom jaar 1';bar.y_axis.title='EUR'
bar.add_data(Reference(wb['Scenariovergelijking'],min_col=2,max_col=4,min_row=9),from_rows=True,titles_from_data=False)
bar.set_categories(Reference(wb['Scenariovergelijking'],min_col=2,max_col=4,min_row=3));bar.width=25;bar.height=8
wb['Samenvatting'].add_chart(bar,'D25')
# Vormgeving.
navy='18354A';teal='DCEEF0';gray='EEF2F6'
for sh in wb:
    sh.sheet_view.showGridLines=False;sh.freeze_panes='B4';sh.sheet_properties.pageSetUpPr.fitToPage=True
    sh.page_setup.orientation='landscape';sh.page_setup.paperSize=sh.PAPERSIZE_A3;sh.page_setup.fitToWidth=1;sh.page_setup.fitToHeight=0
    sh.sheet_properties.outlinePr.summaryRight=False
    sh.column_dimensions['A'].width=45
    for rowcells in sh:
        for cell in rowcells:
            cell.font=Font(name='Calibri',size=11,color='17354A') if not (sh.title=='Invoer' and cell.column==2 and cell.row>=4) else Font(name='Calibri',size=11,color='0000FF')
            if cell.row>3 and cell.row%2==0:cell.fill=PatternFill('solid',fgColor='F5F8FA')
            if cell.column>1 and cell.row>3:cell.number_format='#,##0.00;[Red](#,##0.00);–'
    for cell in sh[1]:cell.fill=PatternFill('solid',fgColor=navy);cell.font=Font(name='Calibri',size=16,bold=True,color='FFFFFF')
    sh.row_dimensions[1].height=32
    for cell in sh[3]:cell.fill=PatternFill('solid',fgColor=teal);cell.font=Font(name='Calibri',bold=True,color=navy)
    sh.print_options.horizontalCentered=True
    sh.print_title_rows='1:3'
    for c in range(2,sh.max_column+1):sh.column_dimensions[L(c)].width=19
wb['Leeswijzer'].column_dimensions['A'].width=125
for rr in range(1,len(notes)+1):wb['Leeswijzer'].row_dimensions[rr].height=30;wb['Leeswijzer'][f'A{rr}'].alignment=Alignment(wrap_text=True,vertical='center')
wb['Invoer'].column_dimensions['A'].width=49;wb['Invoer'].column_dimensions['C'].width=25;wb['Invoer'].column_dimensions['D'].width=70
for key in v:
    wb['Invoer'][f'B{row[key]}'].number_format='0.0%' if items[row[key]-4][3].startswith('%') else '#,##0.00'
wb['Samenvatting'].column_dimensions['B'].width=30
for rr in [12,14]:wb['Samenvatting'][f'B{rr}'].number_format='0.0%'
for rr in [3,4,5,6,15,17]:wb['Samenvatting'][f'B{rr}'].number_format='0'
for col in 'BCD':wb['Scenariovergelijking'][f'{col}11'].number_format='0.0%'
for rr in [13,24,25,28,30]:
    for cell in wb['Kasstromen'][rr]:cell.fill=PatternFill('solid',fgColor=teal);cell.font=Font(name='Calibri',bold=True,color=navy)
from openpyxl.formatting.rule import CellIsRule
red=PatternFill('solid',fgColor='FADBD8');green=PatternFill('solid',fgColor='D5F5E3')
wb['Samenvatting'].conditional_formatting.add('B13',CellIsRule(operator='lessThan',formula=['0'],fill=red))
wb['Samenvatting'].conditional_formatting.add('B13',CellIsRule(operator='greaterThanOrEqual',formula=['0'],fill=green))
wb['Kasstromen'].conditional_formatting.add('B28:V30',CellIsRule(operator='lessThan',formula=['0'],fill=red))
wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
wb.active=2
# Sla op, valideer en voeg berekende beginwaarden toe voor viewers zonder rekenmachine.
tmp=ROOT/'model.tmp.xlsx';packed=ROOT/'model.cached.tmp.xlsx';wb.save(tmp)
probe=load_workbook(tmp);assert len(probe.sheetnames)==5;probe.close()
ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with zipfile.ZipFile(tmp) as zin,zipfile.ZipFile(packed,'w',zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        blob=zin.read(item.filename)
        if item.filename.startswith('xl/worksheets/sheet') and item.filename.endswith('.xml'):
            idx=int(item.filename.split('sheet')[-1].split('.')[0])-1;name=wb.sheetnames[idx]
            root=ET.fromstring(blob)
            for node in root.findall('.//m:c',ns):
                key=(name,node.attrib['r'])
                if key not in cache:continue
                value=cache[key];vn=node.find('m:v',ns)
                if vn is None:vn=ET.SubElement(node,'{'+ns['m']+'}v')
                if isinstance(value,str):node.set('t','str');vn.text=value
                else:node.attrib.pop('t',None);vn.text=str(value)
            blob=ET.tostring(root,encoding='utf-8',xml_declaration=True)
        zout.writestr(item,blob)
probe=load_workbook(packed,data_only=True)
assert abs(probe['Samenvatting']['B13'].value-npv)<.001
assert abs(probe['Kasstromen']['V30'].value-sum(flows))<.001
assert probe['Samenvatting']['B7'].value=='OK'
assert all(not (c.data_type=='e') for sh in probe for rr in sh for c in rr)
probe.close();os.replace(packed,ROOT/'model.xlsx');tmp.unlink(missing_ok=True)
print('Uitgebreid model atomair opgeslagen.')
print(f'Investering EUR {inv:,.2f}; omzet jaar 1 EUR {data[0][0]:,.2f}; operationele kasstroom EUR {data[0][2]:,.2f}; NCW EUR {npv:,.2f}; IRR {irr:.2%}')
print('SAVE_NOW:',(ROOT/'SAVE_NOW.txt').exists())
