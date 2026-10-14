from pathlib import Path
import os, zipfile, xml.etree.ElementTree as ET, json
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parent
p=ROOT/'model.xlsx'; tmp=ROOT/'model.tmp.xlsx'
# Preserve cached calculation results and all formulas while improving row visibility.
stop=(ROOT/'SAVE_NOW.txt').exists()
ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}; ET.register_namespace('',ns['m'])
w=load_workbook(p)
with zipfile.ZipFile(p) as z,zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as out:
 for item in z.infolist():
  data=z.read(item.filename)
  if not stop:
   for i,name in enumerate(w.sheetnames,1):
    if item.filename==f'xl/worksheets/sheet{i}.xml':
     root=ET.fromstring(data)
     for row in root.findall('.//m:sheetData/m:row',ns):
      r=int(row.get('r')); height=None
      if r==1: height=44 if name=='Model' else 32
      elif name=='Read me': height=42
      elif name=='Inputs': height=46
      elif name=='Cost summary': height=34
      elif name=='Results' and r<=19: height=38
      elif name=='Locations': height=42
      elif name=='Sources': height=78
      if height: row.set('ht',str(height)); row.set('customHeight','1')
     data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
  out.writestr(item,data)
with zipfile.ZipFile(tmp) as z: assert z.testzip() is None
formulas=load_workbook(tmp,data_only=False); values=load_workbook(tmp,data_only=True)
count=0
for ws in formulas:
 for row in ws:
  for c in row:
   if c.data_type=='f':
    count+=1
    assert values[ws.title][c.coordinate].value is not None,(ws.title,c.coordinate,'missing cache')
    assert '#REF!' not in c.value,(ws.title,c.coordinate)
assert formulas.sheetnames[-1]=='Results'
assert values['Model']['P22'].value==values['Results']['B16'].value
assert abs(values['Model']['S22'].value-values['Results']['B2'].value)<1e-7
assert values['Model']['A22'].value==20
os.replace(tmp,p)
print('Workbook validation passed:',count,'formulas with populated caches;',len(formulas.sheetnames),'tabs')
print('SAVE_NOW:',(ROOT/'SAVE_NOW.txt').exists())
if not (ROOT/'SAVE_NOW.txt').exists():
 (ROOT/'LLM_prompt.txt').write_text(formulas['LLM prompt']['A2'].value+'\n',encoding='utf-8')
 (ROOT/'README.txt').write_text('Student Housing Investment — VUB Decision Analytics\n\nOpen model.xlsx in Excel. Blue cells in Inputs are editable; Excel recalculates green formulas.\nFinancial inputs and travel times are illustrative. Official sources and unresolved project requirements are in Sources and Read me.\n\nRebuild in order (Python 3 with openpyxl):\npython3 build_model.py\npython3 improve_model.py\npython3 finalize_model.py\n\nThe workbook is saved via a temporary .xlsx and atomic os.replace. Numeric caches allow immediate baseline viewing. The Results tab is last.\nLLM_prompt.txt contains the reusable prompt also included in the workbook.\n',encoding='utf-8')
 report={'formula_cells':count,'tabs':formulas.sheetnames,'NPV_EUR':values['Results']['B2'].value,'IRR':values['Results']['B3'].value,'payback_years':values['Results']['B8'].value,'discounted_payback_years':values['Results']['B9'].value,'NPV_without_exit_EUR':values['Results']['B7'].value,'checks':['xlsx archive integrity','all formula caches populated','20-year horizon','cash-flow totals reconcile','IRR root checked by improve_model.py'],'caveat':'Independent numeric baseline checks; no Excel calculation engine available in environment.'}
 (ROOT/'validation.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print('Final saved model size',p.stat().st_size)
