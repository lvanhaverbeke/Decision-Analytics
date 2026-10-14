from pathlib import Path
import os, math, zipfile, xml.etree.ElementTree as ET
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
ROOT=Path(__file__).resolve().parent
p=ROOT/'model.xlsx'; w=load_workbook(p)
if (ROOT/'SAVE_NOW.txt').exists():
 t=ROOT/'model.tmp.xlsx'; w.save(t); os.replace(t,p); print('SAVE_NOW: saved and stopped'); raise SystemExit
idx={w['Inputs'].cell(r,1).value:r for r in range(2,w['Inputs'].max_row+1)}
def I(n): return "'Inputs'!$B$"+str(idx[n])
w['Inputs'].cell(idx['Required social share'],2,.20)
w['Inputs'].cell(idx['Required social share'],4,'Conservative screening assumption, inspired by City of Brussels guidance (30+ rooms); applicability to Etterbeek/Ixelles unconfirmed.')
w['Inputs'].cell(idx['Social share'],4,'20% affordable/social room target; eligible tenants and rent mechanism require agreement.')
w['Inputs'].cell(idx['Acquisition tax proxy'],4,'Regional standard 12.5% registration duty; actual asset transfer may instead trigger VAT. No personal residence abatement assumed.')
w['Inputs'].cell(idx['Horizon'],4,'Fixed 20-year workbook horizon; do not change.')
dv=DataValidation(type='whole',operator='equal',formula1=20); dv.showErrorMessage=True; dv.error='This model has a fixed 20-year horizon.'; w['Inputs'].add_data_validation(dv); dv.add(w['Inputs'].cell(idx['Horizon'],2))
for n,lo,hi in [('Rooms',1,10000),('Billable months',1,12),('Replacement year',1,20)]:
 dv=DataValidation(type='whole',operator='between',formula1=lo,formula2=hi); dv.showErrorMessage=True; w['Inputs'].add_data_validation(dv); dv.add(w['Inputs'].cell(idx[n],2))
w['Read me']['B7']='20% chosen target. City of Brussels guidance specifies 20% for 30+ student units; other municipalities and site conditions must be confirmed.'
w['Read me']['B13']='Enhanced version: official sources reviewed 14 October 2026; numeric results cached for immediate viewing.'
w['Read me'].append(['Social definition','Reduced rent alone does not establish legal social status. Confirm eligibility, allocation, affordable rents, duration and university/AISE management.'])
w['Read me'].append(['Taxes and exit','Default tax rate 0 gives pre-tax results. Nonzero rate applies only to positive operating NOI, not taxable accounting profit or capital gains.'])
w['Read me'].append(['Cost completeness','Works and furniture inputs must include nonrecoverable VAT. Fire safety, accessibility, permits, EPB and communal facilities belong in works/fees; no grant assumed.'])
s=w['Sources']
rows=[('Social room quota','https://www.bruxelles.be/sites/default/files/bxl/Recommandations_logements_etudiants_VBX_17juin2021.pdf','City guidance, 17 June 2021; search-index excerpt; full PDF fetch blocked','20% social student units from 30 units; City of Brussels scope. Confirm current permit conditions locally.'),('Registration duty','https://be.brussels/fr/impots-financement/impots-et-taxes/fiscalite-immobiliere/droits-denregistrement','Official Brussels Region; reviewed 2026-10-14','Standard rate 12.5%; applied to acquisition proxy; transaction basis needs notary confirmation.'),('VAT alternative','https://financien.belgium.be/fr/particuliers/habitation/acheter-vendre/droits-enregistrement/bruxelles','Federal Finance; reviewed 2026-10-14','New property purchase may trigger VAT instead of registration duty. Avoid double counting.'),('Student quality label','https://be.brussels/en/housing/rental/conformity-for-rental-properties/ensuring-quality-student-housing','Official Brussels Region; reviewed 2026-10-14','Housing Code conformity, compliant student lease and valid EPB required for label application; label lasts five years.'),('Student social eligibility','https://www.vub.be/sites/default/files/2026-04/2026_Huurinformatie_VUB_Koten_Handleiding_aanvraag_SVKS_ENG_0.pdf','VUB 2026 SVKS guidance; reviewed search excerpt','SVKS route requires student enrolment and social income eligibility; use VUB/AISE to validate allocation rules.'),('Institution collaboration','https://www.ulb.be/medias/fichier/20250725-ulb-environnement-et-mobilite-rapport-de-durabilite-en_1753692398626-pdf?ID_FICHE=1020633&INLINE=FALSE','ULB sustainability report 2025; reviewed search excerpt','ULB describes AISE partnership. Candidate counterpart, not a secured agreement for this project.'),('AISE contact','https://www.aise-svks.brussels/proprietaires/','Site under construction when checked 2026-10-14','Published contact info@aise-svks.brussels; negotiate landlord contract, allocation and rent terms.'),('Municipal housing tax','https://www.brussels.be/student-housing','City of Brussels; reviewed 2026-10-14','City distinguishes student second-home exemption from furnished letting tax on landlord. Check municipality of actual site.'),('Travel / financial market data','Analyst scenario assumptions','Not independently measured or quoted','Rents, capital costs and travel times are illustrative; require property quotations and timetable measurements.')]
for row in rows:
 s.append(row); c=s.cell(s.max_row,2)
 if str(c.value).startswith('http'): c.hyperlink=c.value; c.font=Font(color='0563C1',underline='single')
s.column_dimensions['B'].width=65; s.column_dimensions['C'].width=55
for r in range(2,s.max_row+1): s.row_dimensions[r].height=64
m=w['Model']; m['T1']='Simple recovery year'; m['U1']='Discounted recovery year'
for y in range(21):
 r=y+2
 m.cell(r,20,'=999' if y==0 else f'=IF(AND(P{r}>=0,P{r-1}<0),A{r}-1-P{r-1}/O{r},999)')
 m.cell(r,21,'=999' if y==0 else f'=IF(AND(S{r}>=0,S{r-1}<0),A{r}-1-S{r-1}/R{r},999)')
 for col in [20,21]: m.cell(r,col).number_format='0.00'; m.cell(r,col).font=Font(color='00804A'); m.column_dimensions[m.cell(1,col).column_letter].width=24
rs=w['Results']; rs['B8']='=IF(MIN(Model!T2:T22)=999,"Not reached in 20 years",MIN(Model!T2:T22))'; rs['B9']='=IF(MIN(Model!U2:U22)=999,"Not reached in 20 years",MIN(Model!U2:U22))'
rs['B16']='=SUM(Model!O2:O22)'; rs['A16']='Total undiscounted net cash'; rs['C16']='Includes initial investment and disposal'
rs['A17']='Social rooms'; rs['B17']='=Model!B3'; rs['C17']='Rounded up to whole rooms'
rs['A18']='Cumulative operations at year 20'; rs['B18']='=Model!P22-Model!M22-Model!N22'; rs['C18']='Excludes sale proceeds and working capital release'
rs['A19']='Model checks'; rs['B19']=f'=IF(AND({I("Horizon")}=20,Locations!E7=1,Model!B3+Model!C3={I("Rooms")}),"OK","CHECK INPUTS")'; rs['C19']='20-year horizon, student weights sum to 1, room counts reconcile'
for chart in rs._charts: chart.anchor='A22'
w['LLM prompt']['A2']=w['LLM prompt']['A2'].value+' Note the City of Brussels 2021 student housing recommendations specify a 20% social share from 30 units; verify current applicability to the selected municipality and permit. Supply a numeric baseline while keeping formulas editable, show the impact of excluding terminal sale, and provide a source and status for each legal assumption.'
w['LLM prompt'].row_dimensions[2].height=280
for ws in [w['Read me'],s,rs]:
 for row in ws.iter_rows(min_row=2):
  for c in row:
   c.alignment=Alignment(wrap_text=True,vertical='top')
   if c.data_type=='f': c.font=Font(color='00804A'); c.number_format='#,##0.00;[Red](#,##0.00)'
rs['B3'].number_format='0.00%'
t=ROOT/'model.tmp.xlsx'; w.save(t); load_workbook(t); os.replace(t,p)
print('Enhanced model saved atomically; SAVE_NOW:',(ROOT/'SAVE_NOW.txt').exists())
# Independently compute baseline and embed Excel formula caches, keeping every formula.
a={n:w['Inputs'].cell(r,2).value for n,r in idx.items()}; cache={}
def put(sheet,cell,val): cache.setdefault(sheet,{})[cell]=val
N=a['Rooms']; social=math.ceil(N*a['Social share']); market=N-social
initial=[a['Acquisition price'],a['Acquisition price']*a['Acquisition tax proxy'],a['Legal and transaction fees'],a['Construction / renovation'],a['Construction / renovation']*a['Professional fees'],a['Construction / renovation']*(1+a['Professional fees'])*a['Works contingency'],N*a['Furniture per room'],a['Initial working capital']]
for r,val in enumerate(initial,2): put('Cost summary',f'B{r}',val)
initial_total=sum(initial); put('Cost summary','B10',initial_total)
fixed=N*(a['Maintenance per room']+a['Landlord utilities per room'])+a['Insurance']+a['Property tax']+a['Administration / partnership']
flows=[]; discounted=[]; cum=0; dcum=0; records=[]; pay=[]; dpay=[]
for y in range(21):
 mr=market*a['Market room rent']*a['Billable months']*a['Occupancy']*(1+a['Rent growth'])**(y-1) if y else 0
 sr=social*a['Social room rent']*a['Billable months']*a['Occupancy']*(1+a['Rent growth'])**(y-1) if y else 0
 rev=mr+sr; fee=rev*a['Management fee']; op=fixed*(1+a['Operating cost growth'])**(y-1) if y else 0; noi=rev-fee-op; tax=max(noi,0)*a['Effective cash tax rate']; cap=a['Furniture replacement cost']*(1+a['Operating cost growth'])**(y-1) if y==a['Replacement year'] else 0
 sale=((rev*(1+a['Rent growth']))*(1-a['Management fee'])-op*(1+a['Operating cost growth']))/a['Exit capitalization rate']*(1-a['Sale costs']) if y==20 else 0
 wc=a['Initial working capital'] if y==20 else 0; inv=initial_total if y==0 else 0; cf=noi-tax-inv-cap+sale+wc; factor=(1+a['Discount rate'])**(-y); dcf=cf*factor
 pb=y-1-cum/cf if y and cum<0 and cum+cf>=0 else 999; dpb=y-1-dcum/dcf if y and dcum<0 and dcum+dcf>=0 else 999
 cum+=cf; dcum+=dcf; vals=[y,social,market,mr,sr,rev,fee,op,noi,tax,inv,cap,sale,wc,cf,cum,factor,dcf,dcum,pb,dpb]; records.append(vals); flows.append(cf); discounted.append(dcf); pay.append(pb); dpay.append(dpb)
 for col,val in enumerate(vals,1):
  from openpyxl.utils import get_column_letter
  if m.cell(y+2,col).data_type=='f': put('Model',f'{get_column_letter(col)}{y+2}',val)
for r,val in enumerate([N*a['Maintenance per room'],N*a['Landlord utilities per room'],a['Insurance'],a['Property tax'],a['Administration / partnership'],records[1][6],a['Furniture replacement cost'],a['Sale costs']],11): put('Cost summary',f'B{r}',val)
trav=[sum(v*weight for v,weight in zip(times,[.45,.45,.10])) for times in [[10,15,50],[25,30,30],[50,55,10]]]
for r,val in enumerate(trav,2): put('Locations',f'E{r}',val)
put('Locations','E7',1); put('Locations','B8',min(trav)); put('Locations','B9','Etterbeek / Ixelles')
def npv(rate): return sum(c/(1+rate)**y for y,c in enumerate(flows))
lo=-.9; hi=1
for _ in range(150):
 mid=(lo+hi)/2
 if npv(mid)>0: lo=mid
 else: hi=mid
irr=(lo+hi)/2
results={2:sum(discounted),3:irr,4:initial_total,5:records[1][8],6:records[-1][12],7:sum(discounted)-records[-1][12]*records[-1][16],8:min(pay) if min(pay)!=999 else 'Not reached in 20 years',9:min(dpay) if min(dpay)!=999 else 'Not reached in 20 years',10:'Positive NPV' if sum(discounted)>0 else 'Negative / zero NPV',11:'Pending verification / agreement',12:'Etterbeek / Ixelles',13:'Meets assumed target',14:'Revise / verify before committing',16:sum(flows),17:social,18:cum-records[-1][12]-records[-1][13],19:'OK'}
for r,val in results.items(): put('Results',f'B{r}',val)
ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}; ET.register_namespace('',ns['m'])
with zipfile.ZipFile(p) as zin, zipfile.ZipFile(t,'w',zipfile.ZIP_DEFLATED) as zout:
 for item in zin.infolist():
  data=zin.read(item.filename)
  for k,sheet in enumerate(w.sheetnames,1):
   if item.filename==f'xl/worksheets/sheet{k}.xml' and sheet in cache:
    root=ET.fromstring(data)
    for c in root.findall('.//m:sheetData/m:row/m:c',ns):
     address=c.get('r')
     if address in cache[sheet] and c.find('m:f',ns) is not None:
      value=cache[sheet][address]; v=c.find('m:v',ns)
      if v is None: v=ET.SubElement(c,'{'+ns['m']+'}v')
      if isinstance(value,str): c.set('t','str'); v.text=value
      else: c.pop('t',None) if hasattr(c,'pop') else c.attrib.pop('t',None); v.text=str(value)
    data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
  zout.writestr(item,data)
check=load_workbook(t,data_only=True); assert abs(check['Results']['B2'].value-sum(discounted))<.01
assert abs(npv(irr))<.01
assert abs(sum(initial)-initial_total)<.01
assert len(records)==21 and all(r[1]+r[2]==N for r in records)
os.replace(t,p)
print('Cached results saved and checked:',results)
print('SAVE_NOW:',(ROOT/'SAVE_NOW.txt').exists())
