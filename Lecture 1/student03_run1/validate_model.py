from pathlib import Path
from openpyxl import load_workbook
import json, math, zipfile
p=Path(__file__).resolve().parent
w=load_workbook(p/'model.xlsx',data_only=True)
f=load_workbook(p/'model.xlsx',data_only=False)
r=w['Executive Report']; c=w['Cash Flow']; a=w['Inputs']
flows=[c.cell(i,10).value for i in range(2,63)]
rate=a['B17'].value
npv=sum(v/(1+rate)**t for t,v in enumerate(flows))
assert math.isclose(npv,r['B2'].value,abs_tol=.001)
assert abs(sum(v/(1+r['B3'].value)**t for t,v in enumerate(flows)))<.001
assert flows[0]==-7500000 and flows[1]==-7500000 and flows[2]==0
assert flows[3]==6669000 and c['D5'].value==450
assert math.isclose(c['D6'].value,472.5,abs_tol=1e-7)
assert sum(c.cell(i,8).value for i in range(2,63))==15000000
assert r['B4'].value<10 and r['B10'].value=='PASS'
assert w['Scenarios']['G2'].value==r['B2'].value
assert w['Scenarios']['H2'].value==r['B3'].value
assert w['Scenarios']['I2'].value==r['B4'].value
assert w['Scenarios']['K7'].value=='FAIL'
assert w['Scenarios']['J7'].value=='Not reached'
assert math.isclose(w['Sensitivity']['B5'].value,npv,abs_tol=.001)
count=0
for sh in f:
 for row in sh:
  for cell in row:
   if cell.data_type=='f':
    assert w[sh.title][cell.coordinate].value is not None,(sh.title,cell.coordinate)
    count+=1
with zipfile.ZipFile(p/'model.xlsx') as z: assert z.testzip() is None
results={'validated_formulas_with_cached_results':count,'npv_reconciled':True,'irr_residual_under_one_cent':True,'construction_and_rent_timing_verified':True,'scenario_and_sensitivity_reconciled':True,'save_now':(p/'SAVE_NOW.txt').exists()}
(p/'validation_results.json').write_text(json.dumps(results,indent=2))
print(json.dumps(results,indent=2))
