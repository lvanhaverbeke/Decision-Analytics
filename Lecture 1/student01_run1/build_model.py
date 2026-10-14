from pathlib import Path
import os, math, json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, Reference
ROOT=Path(__file__).resolve().parent
wb=Workbook(); info=wb.active; info.title='Read me'
info.append(['BRUSSELS STUDENT HOUSING — DECISION MODEL'])
info.append(['Purpose','Compare affordable student rent with investor project returns. All starting values are illustrative assumptions, not quotations.'])
info.append(['Use','Edit blue cells on Inputs. Read Results, then Cash flow. Monetary amounts are nominal EUR.'])
info.append(['Scope','Unlevered, pre-tax project cash flows. No debt, tax or depreciation. All investment paid at time 0, a conservative simplification.'])
info.append(['Timing','Construction months rounded UP to full years. First full rental year follows construction; no partial-year income.'])
info.append(['Investment','Includes land, construction, facilities, professional fees, contingency and explicit compliance cost. Subsidy is received at time 0.'])
info.append(['Student price','Monthly all-in rent: utilities included. Separate affordability cap is an editable policy target.'])
info.append(['Return','NPV at entered discount rate; IRR uses yearly project cash flows; PBP is simple payback from investment, excluding sale proceeds.'])
info.append(['Risk','Force-majeure contingency is a scenario allowance, not a probability forecast. Illegal payments are not budgeted; use compliance and delay scenarios.'])
a=wb.create_sheet('Inputs'); a.append(['Input','Value','Unit','Basis / assumption'])
rows=[('Rooms',100,'rooms','Illustrative project; minimum capacity constraint'),('Gross area per room',30,'m²/room','Includes allocation of common areas'),('Land / acquisition',1200000,'EUR','Site-specific quotation needed'),('Construction cost',2000,'EUR/m²','Illustrative; VAT inclusion must be checked'),('Facilities',400000,'EUR','Furniture and common equipment'),('Professional fees',.10,'share','Of construction + facilities'),('Force-majeure contingency',.10,'share','Of construction + facilities + fees'),('Compliance / audit cost',0,'EUR','Legal control costs; no unlawful payments'),('Initial subsidy',0,'EUR','Only committed funding should be used'),('Construction duration',18,'months','Illustrative, including permitting'),('Operating useful life',40,'years','Subject to engineering assessment and refurbishment'),('Discount rate',.06,'annual','Nominal unlevered required return'),('Occupancy',.95,'share','Average paying occupancy'),('Rent / billing months',12,'months/year','Check lease structure'),('Student monthly all-in rent',600,'EUR/month','Illustrative initial rent'),('Affordability cap',650,'EUR/month','Policy assumption; validate student budget'),('Rent escalation',.02,'annual','Scenario assumption, not entitlement to index'),('Operating cost escalation',.025,'annual','Nominal assumption'),('Fixed operations',30000,'EUR/year','Excludes wages, maintenance, utilities'),('Wages incl. employer costs',60000,'EUR/year','Illustrative annual staffing'),('Variable operations',300,'EUR/occupied room/year','Cleaning/administration, excludes utilities'),('Utilities',900,'EUR/occupied room/year','All-in rent model'),('Maintenance reserve',.01,'share/year','Of construction + facilities; treated as annual cash expense'),('Refurbishment interval',15,'operating years','Every N years, before final operating year'),('Refurbishment cost',300000,'EUR/event','Today-price cost, escalates with operating costs'),('Terminal gross asset value',0,'EUR at end','Default zero; enter defensible net-of-tax valuation assumption'),('Sale costs',.02,'share','Of gross terminal value')]
for r in rows:a.append(r)
def I(r):return f"Inputs!$B${r}"
c=wb.create_sheet('Cash flow'); c.append(['Year','Operating year','All-in rent / month','Revenue','Operating costs','Refurbishment','Terminal proceeds','Net cash flow','Discount factor','PV cash flow','Cumulative operating cash','PV revenue per EUR rent'])
r=wb.create_sheet('Results'); r.append(['Metric','Value','Interpretation'])
metrics=[('Gross initial investment',f'={I(4)}+{I(2)}*{I(3)}*{I(5)}+{I(6)}+({I(2)}*{I(3)}*{I(5)}+{I(6)})*{I(7)}+({I(2)}*{I(3)}*{I(5)}+{I(6)})*(1+{I(7)})*{I(8)}+{I(9)}','Includes contingency; no subsidy deduction'),('Investor initial cash',f'=B2-{I(10)}','Gross investment less committed upfront subsidy'),('Construction years',f'=ROUNDUP({I(11)}/12,0)','Conservative rounding; full-year operating model'),('Final cash-flow year',f'=B4+{I(12)}','Construction + operating life'),('NPV','=SUM(\'Cash flow\'!J2:J62)','Project NPV including terminal value'),('IRR','=IFERROR(IRR(\'Cash flow\'!H2:H62),"No IRR")','Annual project IRR; check sign changes'),('Simple PBP (years)','=IFERROR(MATCH(TRUE,INDEX(\'Cash flow\'!K3:K62>=0,0),0)-1+INDEX(-\'Cash flow\'!K2:K61,MATCH(TRUE,INDEX(\'Cash flow\'!K3:K62>=0,0),0))/INDEX(\'Cash flow\'!H3:H62,MATCH(TRUE,INDEX(\'Cash flow\'!K3:K62>=0,0),0)),"Not recovered")','From time 0; interpolated annual cash; excludes terminal sale'),('Student price P',f'={I(16)}','Initial monthly all-in rent'),('Student affordability cap',f'={I(17)}','Editable student budget threshold'),('Required initial rent for NPV = 0',f'=MAX(0,({I(16)}*SUM(\'Cash flow\'!L3:L62)-B6)/SUM(\'Cash flow\'!L3:L62))','Investor break-even price given all other inputs'),('NPV at affordability cap',f'=B6+({I(17)}-{I(16)})*SUM(\'Cash flow\'!L3:L62)','Feasibility under student cap'),('Upfront subsidy needed at cap','=MAX(0,-B12)','Additional time-0 grant / cost reduction required'),('Maximum investor capital at cap','=B3+B12','PV of future net cash flows at capped rent'),('Joint feasibility','=IF(B11<=B10,"Feasible under assumptions","Funding / cost gap")','A positive gap cannot be solved by choosing an arbitrary rent'),('Recommended student price','=IF(B11<=B10,B11,"No feasible price under cap")','Lowest price meeting NPV >= 0 at selected discount rate'),('First full operating year','=B4+1','Minimum safe planning time requires site and permit evidence')]
for row in metrics:r.append(row)
for year in range(61):
 j=year+2
 c.append([year, f'=IF(AND(A{j}>Results!$B$4,A{j}<=Results!$B$5),A{j}-Results!$B$4,0)',f'=IF(B{j}>0,{I(16)}*(1+{I(18)})^(B{j}-1),0)',f'=IF(B{j}>0,{I(2)}*{I(14)}*{I(15)}*C{j},0)',f'=IF(B{j}>0,({I(20)}+{I(21)}+{I(2)}*{I(14)}*({I(22)}+{I(23)})+({I(2)}*{I(3)}*{I(5)}+{I(6)})*{I(24)})*(1+{I(19)})^(B{j}-1),0)',f'=IF(AND(B{j}>0,B{j}<{I(12)},MOD(B{j},{I(25)})=0),{I(26)}*(1+{I(19)})^(B{j}-1),0)',f'=IF(A{j}=Results!$B$5,{I(27)}*(1-{I(28)}),0)', '= -Results!B3' if year==0 else f'=D{j}-E{j}-F{j}+G{j}',f'=1/(1+{I(13)})^A{j}',f'=H{j}*I{j}','=-Results!B3' if year==0 else f'=K{j-1}+H{j}-G{j}',f'=IF(B{j}>0,{I(2)}*{I(14)}*{I(15)}*(1+{I(18)})^(B{j}-1)*I{j},0)'])
for s in wb:
 s.freeze_panes='B2'; s.auto_filter.ref=s.dimensions
 s.row_dimensions[1].height=30
 for cell in s[1]:cell.fill=PatternFill('solid',fgColor='17365D');cell.font=Font(color='FFFFFF',bold=True)
 for row in s.iter_rows(min_row=2):
  for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
 s.column_dimensions['A'].width=39
 for col in 'BCDEFGHIJKL':s.column_dimensions[col].width=23 if s==c else 27
 if s!=c:s.column_dimensions['C'].width=65;s.column_dimensions['D'].width=65
 for row in range(2,s.max_row+1):s.row_dimensions[row].height=32 if s!=c else 20
for row in range(2,29):
 a.cell(row,2).font=Font(color='0070C0');a.cell(row,2).fill=PatternFill('solid',fgColor='E7F1FB');a.cell(row,2).number_format='0.00% ' if row in [7,8,13,14,18,19,24,28] else '#,##0.00'
for row in range(2,18):r.cell(row,2).number_format='0.00%' if row==7 else '#,##0.00'
for row in c.iter_rows(min_row=2,min_col=3):
 for cell in row:cell.number_format='#,##0.00;[Red](#,##0.00)'
wb.calculation.fullCalcOnLoad=True

def save_atomic():
 tmp=ROOT/'model.tmp.xlsx';wb.save(tmp);load_workbook(tmp).close();os.replace(tmp,ROOT/'model.xlsx')
save_atomic()
print('Saved basic model.xlsx', (ROOT/'model.xlsx').stat().st_size)
print('SAVE_NOW:', (ROOT/'SAVE_NOW.txt').exists())
if (ROOT/'SAVE_NOW.txt').exists():
 print('Stop requested; basic workbook saved.');raise SystemExit
# Validation limits keep every operating year inside the 60-year cash-flow horizon.
for rng,lo,hi,typ in [('B2',1,1000,'whole'),('B11',1,60,'whole'),('B12',1,50,'whole'),('B15',1,12,'whole'),('B25',1,50,'whole')]:
 dv=DataValidation(type=typ,operator='between',formula1=lo,formula2=hi);dv.errorTitle='Input outside model range';dv.error='Use a value within the permitted range.';dv.showErrorMessage=True;dv.errorStyle='stop';a.add_data_validation(dv);dv.add(rng)
for rr in [7,8,13,14,18,19,24,28]:
 dv=DataValidation(type='decimal',operator='between',formula1=0,formula2=1);dv.showErrorMessage=True;dv.errorStyle='stop';a.add_data_validation(dv);dv.add(f'B{rr}')
for rr in [3,4,5,6,9,10,16,17,20,21,22,23,26,27]:
 dv=DataValidation(type='decimal',operator='greaterThanOrEqual',formula1=0);dv.showErrorMessage=True;dv.errorStyle='stop';a.add_data_validation(dv);dv.add(f'B{rr}')
r.append(['Input check','=IF(AND(Inputs!B10<=B2,Inputs!B14>0,B5<=60),"OK","CHECK subsidy, occupancy or horizon")','Subsidy must not exceed gross investment; positive occupancy required'])
info.append(['Affordability over time','Cap tests initial rent only. Annual escalated rent appears in Cash flow; recheck student budgets each year.'])
info.append(['Optimization','Exact analytical minimum rent at NPV = 0 for a fixed design. Minimum investment is meaningful only with room capacity, quality and useful-life constraints. No global construction optimum is claimed.'])
info.append(['Model limits','60-year annual horizon; construction <=60 months and life <=50 years. No financial close recommendation without local quotes, tax review and financing terms.'])
info.append(['Double counting','Maintenance is expensed cash each year; refurbishment is separate major work. Do not also count a funded reserve withdrawal as a new expense.'])
info.append(['PBP limitation','Annual interpolation approximates receipt timing. Terminal proceeds excluded. Late negative cash flows can reverse an earlier payback.'])
src=wb.create_sheet('Sources');src.append(['Source','URL','Verified / published','Use and limits'])
for row in [
('Perspective monitoring student housing','https://beecole.brussels/sites/default/files/documents/fr_monitoring_logement_etudiant_final2026.pdf','2026-10-14 / January 2026','Official regional report. Reports rising student rents; supports affordability concern, not model cost assumptions.'),
('Brussels rental steps','https://be.brussels/fr/logement/location/les-differentes-etapes-de-la-location','2026-10-14','Official guidance: student leases are 12 months. 12 billed months still requires occupancy and lease validation.'),
('Brussels student housing label','https://be.brussels/fr/logement/location/conformite-logement/label-logement-etudiant','2026-10-14','Quality label lasts 5 years; quality declaration, student lease and valid energy certificate requested. Does not certify model compliance.'),
('Required project evidence','Not yet provided','Unverified','Site offer; VAT-inclusive bill of quantities; architect/engineer design; planning/fire/energy requirements; staffing and utility quotes; discount rate rationale; student budget survey; insurance; grants; tax and financing terms.')]:src.append(row)
for rr in range(2,5):src.cell(rr,2).hyperlink=src.cell(rr,2).value;src.cell(rr,2).style='Hyperlink'
p=wb.create_sheet('Prompt');p.append(['Step','Reusable prompt / deliverable'])
texts=[('Role and goal','Act as a decision analyst for student housing in Brussels. Build an auditable Excel project model balancing student affordability and investor returns. Do not invent verified inputs or promise a single correct forecast.'),('Gather inputs','For each input give value, unit, currency year, source URL/date, confidence and low/base/high range. Gather land price, room count and gross area, construction and facilities quotes, fees, VAT, utilities, fixed and variable operations, fully loaded wages, maintenance, major refurbishment, occupancy, billing months, nominal discount rate, escalation, construction schedule, useful life, subsidy, terminal value and student all-in budget cap. Mark missing project inputs and use clearly labelled illustrative assumptions.'),('Financial conventions','State unlevered/levered and pre/post-tax basis. Keep real/nominal discounting consistent. Avoid counting utilities, maintenance and renovation twice. Explain whether acquisition and construction are staged; if simplified upfront, say so. Use full monthly or dated cash flows if schedule detail permits.'),('Required outputs','Compute project NPV, annual IRR with sign-change warning, simple and discounted payback excluding residual sale; report initial and future student all-in price P, gross capital and investor funding. Show no-payback and no-feasible-price outcomes explicitly.'),('Optimization','For each technically feasible design hold minimum room capacity, code compliance and service quality fixed. Minimize student price subject to NPV >= 0 and P <= affordability cap. Compare design alternatives for lowest feasible capital, fastest evidenced delivery and longest evidenced useful life, including lifecycle refurbishment. Present Pareto tradeoffs rather than independently minimizing everything. If infeasible, quantify subsidy/cost-reduction gap and explain why.'),('Risk and corruption','Use low/base/high cost, occupancy, timing and discount-rate scenarios. Separate contingency from expected-event loss; apply probability only with evidence. Treat corruption as a governance risk via legal audit, procurement controls and delay/overrun stress, never as a planned illegal payment.'),('Results and pitfalls','Provide a conditional recommendation, assumptions that drive it, sensitivity, sources and unanswered questions. Warn about optimistic occupancy, nominal/real mismatch, VAT, rent/charge mixing, unsupported terminal values, planning delays, missed refurbishment, IRR ambiguity and optimization without constraints. Never infer engineering lifetime or legal compliance from a spreadsheet.'),('Delivery','Deliver model.xlsx with editable inputs, live formulas, annual cash flows, financial results, scenarios, checks and source register. Include enough code to reproduce and audit calculations.')]
for row in texts:p.append(row)
p.column_dimensions['B'].width=110
for rr in range(2,10):p.row_dimensions[rr].height=100
risk=wb.create_sheet('Risks');risk.append(['Pitfall / missing evidence','Effect','Action'])
for row in [('Site and construction quotes missing','Capital estimate may be materially wrong','Obtain tender with VAT, fees, land and utility connections'),('Force majeure','Cost increase and opening delay','Stress cost +20%, delay +12 months; review insurance and contingency'),('Corruption / procurement misconduct','Delay, legal cost, contract failure','Transparent tenders, independent audit, conflicts register; no illicit payment provision'),('Life and schedule unverified','Longer use or early opening can inflate valuation','Engineering survey and permit critical path; compare scenario bounds, not factual min/max'),('Discount rate / inflation','Incorrect return hurdle','Use matching nominal cash flows and nominal project discount rate'),('Terminal value','Sale assumption can dominate return','Base case zero; sensitivity only with evidence'),('Affordability cap','Policy target can differ from student capacity','Survey disposable budget and utilities; track future rent'),('Tax and debt excluded','Project IRR differs from equity IRR','Extend after structure is known; no investment advice implied')]:risk.append(row)
# Scenario drivers are editable assumptions and linked to the base model.
s=wb.create_sheet('Scenarios');s.append(['Scenario','Construction factor','Occupancy','Duration months','Discount rate','Life years','Gross capital','NPV','IRR','Required P','NPV at cap','Extra subsidy','PBP'])
configs=[('Base',1,'=Inputs!B14','=Inputs!B11','=Inputs!B13','=Inputs!B12'),('Lower-cost / fast',.85,.98,12,.05,50),('Delay / overrun',1.2,.85,30,.08,30)]
engines=[]
for idx,conf in enumerate(configs,2):
 s.append(list(conf));ss=wb.copy_worksheet(a);ss.title=f'S{idx-1} inputs';engines.append(ss)
 ss['B5']=f'=Inputs!B5*Scenarios!B{idx}';ss['B14']=f'=Scenarios!C{idx}';ss['B11']=f'=Scenarios!D{idx}';ss['B13']=f'=Scenarios!E{idx}';ss['B12']=f'=Scenarios!F{idx}'
 for rr in range(2,29):
  if rr not in [5,14,11,13,12]:ss.cell(rr,2,f'=Inputs!B{rr}')
 rrw=wb.copy_worksheet(r);rrw.title=f'S{idx-1} results'
 cc=wb.copy_worksheet(c);cc.title=f'S{idx-1} cash flow'
 for sh in [rrw,cc]:
  for row in sh:
   for cell in row:
    if cell.data_type=='f':cell.value=cell.value.replace('Inputs!',f"'{ss.title}'!").replace('Results!',f"'{rrw.title}'!").replace("'Cash flow'!",f"'{cc.title}'!")
 for col,ref in [('G',2),('H',6),('I',7),('J',11),('K',12),('L',13),('M',8)]:s[f'{col}{idx}']=f"='{rrw.title}'!B{ref}"
 for sh in [ss,rrw,cc]:sh.sheet_state='hidden'
s.append(['Scenario interpretation','All three are illustrative. Fastest = 12 months and longest = 50 years ONLY among these assumptions; not engineering estimates.'])
s.append(['Decision','Base capital and minimum rent depend on a fixed 100-room design. A cheaper scenario is not proven technically achievable.'])
chart=LineChart();chart.title='Cumulative operating cash (excluding sale)';chart.y_axis.title='EUR';chart.x_axis.title='Years from investment';chart.add_data(Reference(c,min_col=11,min_row=1,max_row=62),titles_from_data=True);chart.set_categories(Reference(c,min_col=1,min_row=2,max_row=62));r.add_chart(chart,'E2')
for sh in [src,p,risk,s]:
 sh.freeze_panes='B2';sh.auto_filter.ref=sh.dimensions
 for cell in sh[1]:cell.fill=PatternFill('solid',fgColor='17365D');cell.font=Font(color='FFFFFF',bold=True)
 for row in sh.iter_rows(min_row=2):
  for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
 sh.column_dimensions['A'].width=35
 if sh==s:
  for col in 'BCDEFGHIJKLM':sh.column_dimensions[col].width=20
  for rr in range(2,5):
   for col in range(2,7):sh.cell(rr,col).font=Font(color='0070C0')
   for col in [2,3,5,9]:sh.cell(rr,col).number_format='0.00%'
   for col in [7,8,10,11,12,13]:sh.cell(rr,col).number_format='#,##0.00'
 else:
  if sh!=p:sh.column_dimensions['B'].width=65
  sh.column_dimensions['C'].width=45;sh.column_dimensions['D'].width=75
  if sh!=p:
   for rr in range(2,sh.max_row+1):sh.row_dimensions[rr].height=65
save_atomic();print('Saved scenarios, sources, prompt and risk improvements.');print('SAVE_NOW:',(ROOT/'SAVE_NOW.txt').exists())
if (ROOT/'SAVE_NOW.txt').exists():raise SystemExit
# Scalar helper formulas make payback transparent and avoid array formulas.
for cs,rs in [(c,r)]+[(wb[f'S{i} cash flow'],wb[f'S{i} results']) for i in range(1,4)]:
 cs['M1']='Payback crossing year';cs['N1']='Cumulative discounted cash';cs['O1']='Discounted crossing year'
 cs['M2']='';cs['N2']='=J2';cs['O2']=''
 for j in range(3,63):
  cs[f'M{j}']=f'=IF(AND(K{j}>=0,K{j-1}<0),A{j}-1-K{j-1}/(H{j}-G{j}),"")'
  cs[f'N{j}']=f'=N{j-1}+J{j}-G{j}*I{j}'
  cs[f'O{j}']=f'=IF(AND(N{j}>=0,N{j-1}<0),A{j}-1-N{j-1}/((H{j}-G{j})*I{j}),"")'
 rs['B8']=f'=IF(COUNT(\'{cs.title}\'!M3:M62)>0,MIN(\'{cs.title}\'!M3:M62),"Not recovered")'
 rs.append(['Discounted PBP (years)',f'=IF(COUNT(\'{cs.title}\'!O3:O62)>0,MIN(\'{cs.title}\'!O3:O62),"Not recovered")','Discounted at project hurdle; excludes terminal sale'])
 rs.append(['Capital per room',f'=B2/\'{"Inputs" if rs==r else rs.title.replace("results","inputs")}\'!B2','Gross investment / room, including allocated common area and land'])
 for col in 'MNO':cs.column_dimensions[col].width=25
 for row in cs.iter_rows(min_row=2,min_col=13,max_col=15):
  for cell in row:cell.number_format='#,##0.00'
r.append(['Shortest construction scenario','=MIN(Scenarios!D2:D4)','Months, only within the assumed scenarios; not a validated project schedule'])
r.append(['Longest useful-life scenario','=MAX(Scenarios!F2:F4)','Operating years, only within scenarios; needs engineering validation'])
r.append(['Investment recommendation','=IF(B12>=0,"Proceed to project-specific due diligence","Revise cost / secure subsidy before proceeding")','Conditional result, not a verified investment decision'])
# First sheet directs readers to the financial decision and reusable prompt.
info.append(['Workbook map','Inputs → Results → Cash flow → Scenarios. Sources, Risks and Prompt explain evidence and next steps. Hidden scenario calculation sheets can be unhidden to audit.'])
wb.active=wb.sheetnames.index('Results')
save_atomic();print('Saved payback helpers, discounted PBP and planning bounds.');print('SAVE_NOW:',(ROOT/'SAVE_NOW.txt').exists())
if (ROOT/'SAVE_NOW.txt').exists():raise SystemExit
from openpyxl.formatting.rule import CellIsRule
r.append(['Lowest feasible scenario capital','=IF(MAX(Scenarios!K2:K4)<0,"No feasible scenario",MIN(IF(Scenarios!K2>=0,Scenarios!G2,1E+99),IF(Scenarios!K3>=0,Scenarios!G3,1E+99),IF(Scenarios!K4>=0,Scenarios!G4,1E+99)))','Minimum only among illustrative scenarios meeting NPV >= 0 at affordability cap. Design feasibility and quality must be independently confirmed.'])
for target in ['B6','B12']:
 r.conditional_formatting.add(target,CellIsRule(operator='lessThan',formula=['0'],fill=PatternFill('solid',fgColor='FCE4D6')))
 r.conditional_formatting.add(target,CellIsRule(operator='greaterThanOrEqual',formula=['0'],fill=PatternFill('solid',fgColor='E2F0D9')))
for rr in [2,3,6,9,10,11,12,13,14,20,24]:r.cell(rr,2).number_format='€ #,##0.00;[Red](€ #,##0.00)'
for rr in [4,5,8,17,19,21,22]:r.cell(rr,2).number_format='0.00'
r.column_dimensions['B'].width=34;r.column_dimensions['C'].width=75
r.row_dimensions[24].height=60
for rr in range(2,16):info.row_dimensions[rr].height=48
info.column_dimensions['B'].width=110
for rr in range(2,5):
 for col,lo,hi,typ in [('B',.1,3,'decimal'),('C',.01,1,'decimal'),('D',1,60,'whole'),('E',0,1,'decimal'),('F',1,50,'whole')]:
  dv=DataValidation(type=typ,operator='between',formula1=lo,formula2=hi);dv.showErrorMessage=True;dv.errorStyle='stop';s.add_data_validation(dv);dv.add(f'{col}{rr}')
save_atomic();print('Saved final presentation and scenario capital comparison.');print('SAVE_NOW:',(ROOT/'SAVE_NOW.txt').exists())
