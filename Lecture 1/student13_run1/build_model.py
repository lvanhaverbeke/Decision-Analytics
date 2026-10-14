from pathlib import Path
import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference
from openpyxl.workbook.properties import CalcProperties
ROOT=Path(__file__).resolve().parent
OUT=ROOT/'model.xlsx'
NAVY='142D4E'; BLUE='1B65A0'; TEAL='007F78'; LIGHT='EAF0F6'
wb=Workbook(); wb.remove(wb.active)
for name in ['Report','Guide','Assumptions','Opportunities','Compliance','Cash flows','Portfolio','Sources']:
 wb.create_sheet(name)
wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
def title(ws,text,subtitle):
 ws.append([text]);ws.append([subtitle]);ws.row_dimensions[1].height=30
 ws['A1'].font=Font(name='Calibri',size=21,bold=True,color=NAVY)
 ws['A2'].font=Font(name='Calibri',size=11,color='555555')
def header(ws,row,values):
 for c,v in enumerate(values,1):
  cell=ws.cell(row,c,v);cell.fill=PatternFill('solid',fgColor=NAVY);cell.font=Font(bold=True,color='FFFFFF');cell.alignment=Alignment(wrap_text=True)
 ws.row_dimensions[row].height=32
A=wb['Assumptions'];title(A,'Investment assumptions','Blue cells are editable. Unlevered, pre-tax EUR model; all capital committed at time zero.')
ass=[('Demo mode (YES / NO)','YES','YES accepts hypothetical demonstration checks; NO requires Verified evidence.'),('Available capital',3500000,'EUR; includes all-in acquisition and construction costs'),('Maximum projects',3,'Portfolio cap; eight opportunity slots'),('Discount rate',.09,'Required annual unlevered return'),('Holding period (years)',10,'Fixed 10 years; cash flow grid is ten years'),('Annual rent growth',.02,'Applied to base rent from acquisition date'),('Exit capitalization rate',.055,'Year 11 stabilized NOI / cap rate'),('Exit selling costs',.03,'Fraction of exit value'),('Acquisition costs',.14,'Planning estimate; replace with transaction-specific tax/legal fees'),('Construction contingency',.15,'Fraction of base renovation budget'),('Stress rent reduction',.1,'Rent haircut for downside case'),('Stress delay (months)',6,'Additional delay in downside case'),('Stress construction uplift',.2,'Applied on top of contingency'),('Maximum acceptable delay',36,'Months from acquisition to legal rental readiness')]
header(A,4,['Parameter','Value','Interpretation'])
for r,row in enumerate(ass,5):
 for c,v in enumerate(row,1): A.cell(r,c,v)
 A.cell(r,2).font=Font(color=BLUE)
for r in [8,10,11,12,13,14,15,17]:A.cell(r,2).number_format='0.0%'
A['B6'].number_format='#,##0 "EUR"'
O=wb['Opportunities'];title(O,'Opportunity universe | eight investment slots','HYPOTHETICAL EXAMPLES ONLY. Replace with your spreadsheet data; no Brussels properties have been sourced.')
cols=['ID','Project / location','Rooms','Monthly rent / room','Purchase price','Base construction cost','Other capital costs','Months to ready','Occupancy','Operating cost / revenue','Annual fixed costs','All-in capital','Compliance gate','Year 1 NOI','Year 11 NOI','Base NPV','Base IRR','Downside NPV','Decision','Rank','Portfolio selected']
header(O,7,cols)
data=[['DEMO01','Illustrative conversion A',14,640,1050000,230000,35000,12,.95,.24,14000],['DEMO02','Illustrative renovation B',10,620,650000,140000,25000,6,.95,.25,10500],['DEMO03','Illustrative new build C',24,680,1750000,620000,70000,30,.94,.23,23500],['DEMO04','Illustrative small asset D',6,590,430000,75000,18000,4,.94,.26,7000],['DEMO05','Illustrative delayed asset E',16,650,1100000,360000,45000,42,.95,.24,16000],['DEMO06','Illustrative compliance failure F',12,660,820000,175000,28000,8,.95,.24,12000]]
for i in range(8):
 r=8+i
 vals=data[i] if i<len(data) else [f'SLOT{i+1}','',0,0,0,0,0,0,.95,.25,0]
 for c,v in enumerate(vals,1):O.cell(r,c,v).font=Font(color=BLUE)
 O.cell(r,12,f'=IF(B{r}="",0,E{r}*(1+Assumptions!$B$13)+F{r}*(1+Assumptions!$B$14)+G{r})')
 cr=5+i*12
 O.cell(r,13,f'=IF(B{r}="","EMPTY",IF(COUNTIF(Compliance!D{cr}:D{cr+11},"Fail")>0,"FAIL",IF(AND(COUNTIF(Compliance!D{cr}:D{cr+11},"Verified")=12,COUNTBLANK(Compliance!E{cr}:E{cr+11})=0),"PASS",IF(AND(Assumptions!$B$5="YES",COUNTIF(Compliance!D{cr}:D{cr+11},"Demo pass")=12),"DEMO PASS","UNVERIFIED"))))')
 O.cell(r,14,f"='Cash flows'!C{5+i}");O.cell(r,15,f'=C{r}*D{r}*12*I{r}*(1-J{r})*(1+Assumptions!$B$10)^10-K{r}*(1+Assumptions!$B$10)^10')
 O.cell(r,16,f"=NPV(Assumptions!$B$8,'Cash flows'!C{5+i}:L{5+i})+'Cash flows'!B{5+i}")
 O.cell(r,17,f"=IFERROR(IRR('Cash flows'!B{5+i}:L{5+i}),0)")
 O.cell(r,18,f"=NPV(Assumptions!$B$8,'Cash flows'!C{18+i}:L{18+i})+'Cash flows'!B{18+i}")
 O.cell(r,19,f'=IF(B{r}="","EMPTY",IF(AND(OR(M{r}="PASS",M{r}="DEMO PASS"),H{r}<=Assumptions!$B$18,C{r}>0,D{r}>0,E{r}>0,H{r}>=0,I{r}>0,I{r}<=1,J{r}>=0,J{r}<1,F{r}>=0,G{r}>=0,K{r}>=0,P{r}>0),"ELIGIBLE","EXCLUDE"))')
 O.cell(r,20,f'=IF(S{r}="ELIGIBLE",1+COUNTIFS($S$8:$S$15,"ELIGIBLE",$P$8:$P$15,">"&P{r}),"")')
 O.cell(r,21,f'=INDEX(Portfolio!B$5:B$260,Report!$B$7)') if i==0 else O.cell(r,21,f'=INDEX(Portfolio!{chr(66+i)}$5:{chr(66+i)}$260,Report!$B$7)')
C=wb['Compliance'];title(C,'Mandatory requirements | project evidence register','Unknown evidence blocks eligibility. Demo pass is permitted only in explicit demonstration mode.')
header(C,4,['ID','Control','Minimum / validation needed','Status','Evidence / authority / date'])
controls=[('Planning use','Written confirmation that student housing use and works are lawful'),('Rental authorization','Confirm permits, registrations and required approvals before letting'),('Minimum room area','Verify room and dwelling area requirements for actual configuration'),('Occupancy limits','Verify maximum occupants and shared facility capacity'),('Sanitation','Compliant WC, washing and bathing facilities'),('Kitchen / cooking','Compliant cooking facilities and equipment'),('Daylight / ventilation','Compliant natural light, ventilation and air quality'),('Electrical / gas safety','Required inspection certificates and safe installations'),('Fire safety','Fire authority requirements, escape routes and detection'),('Water / heating / moisture','Safe water, heating, weatherproofing and absence of dangerous damp'),('Energy certification','Applicable PEB certificate and required disclosures'),('Accessibility / other rules','Case-specific access, municipal, co-ownership and student rental rules')]
for i in range(8):
 for j,(label,rule) in enumerate(controls):
  r=5+i*12+j;C.append([f'DEMO{i+1:02}' if i<6 else f'SLOT{i+1}',label,rule,'Demo pass' if i<5 else ('Fail' if i==5 and j==0 else 'Unknown'),'Illustrative only — no evidence' if i<6 else ''])
dv=DataValidation(type='list',formula1='"Unknown,Verified,Fail,Demo pass"');C.add_data_validation(dv);dv.add('D5:D100')
F=wb['Cash flows'];title(F,'Ten-year cash flows | base and downside','Annual approximation: delay reduces operating months. Capital paid at acquisition; terminal sale at year 10.')
header(F,4,['Project','Year 0']+[f'Year {y}' for y in range(1,11)])
header(F,17,['Downside','Year 0']+[f'Year {y}' for y in range(1,11)])
for i in range(8):
 r=8+i
 for stress,fr in [(False,5+i),(True,18+i)]:
  F.cell(fr,1,f'=Opportunities!B{r}')
  extra=f'+Opportunities!F{r}*(1+Assumptions!$B$14)*Assumptions!$B$17' if stress else ''
  F.cell(fr,2,f'=-Opportunities!L{r}{"-Opportunities!F"+str(r)+"*(1+Assumptions!$B$14)*Assumptions!$B$17" if stress else ""}')
  for year in range(1,11):
   delay=f'Opportunities!H{r}'+('+Assumptions!$B$16' if stress else '')
   haircut='*(1-Assumptions!$B$15)' if stress else ''
   noi=f'(Opportunities!C{r}*Opportunities!D{r}*12*Opportunities!I{r}*(1-Opportunities!J{r}){haircut}-Opportunities!K{r})*(1+Assumptions!$B$10)^{year-1}'
   formula=f'={noi}*MAX(0,MIN(1,({year}*12-({delay}))/12))'
   if year==10:formula+=f'+MAX(0,(Opportunities!O{r}{haircut})/Assumptions!$B$11)*(1-Assumptions!$B$12)'
   F.cell(fr,year+2,formula)
P=wb['Portfolio'];title(P,'Portfolio optimization | exact enumeration','256 subsets of eight projects. Maximize aggregate NPV subject to capital, project count and compliance.')
header(P,4,['Subset']+[f'Select {i+1}' for i in range(8)]+['Capital','Projects','NPV','Feasible','Objective'])
for mask in range(256):
 r=5+mask;P.cell(r,1,mask)
 for i in range(8):P.cell(r,2+i,(mask>>i)&1)
 P.cell(r,10,f'=SUMPRODUCT(B{r}:I{r},Opportunities!$L$8:$L$15)')
 # SUMPRODUCT needs same orientation, use explicit sums for vertical input ranges
 P.cell(r,10,'='+'+'.join(f'{chr(66+i)}{r}*Opportunities!L{8+i}' for i in range(8)))
 P.cell(r,11,f'=SUM(B{r}:I{r})')
 P.cell(r,12,'='+'+'.join(f'{chr(66+i)}{r}*Opportunities!P{8+i}' for i in range(8)))
 eligible='+'.join(f'{chr(66+i)}{r}*IF(Opportunities!S{8+i}="ELIGIBLE",0,1)' for i in range(8))
 P.cell(r,13,f'=IF(AND(J{r}<=Assumptions!$B$6,K{r}<=Assumptions!$B$7,{eligible}=0),1,0)')
 P.cell(r,14,f'=IF(M{r}=1,L{r},-1E99)')
R=wb['Report'];title(R,'BRUSSELS | STUDENT HOUSING','Investment committee report • capital allocation • model date 14 October 2026')
R['A4']='Use / evidence status';R['B4']='=IF(Assumptions!B5="YES","DEMONSTRATION — hypothetical opportunities","LIVE INPUTS — evidence required")'
R['A5']='Mandate';R['B5']='Maximize 10-year discounted unlevered profit within budget; legal readiness is mandatory.'
R['A6']='Optimal portfolio NPV';R['B6']='=MAX(Portfolio!N5:N260)'
R['A7']='Winning subset position';R['B7']='=MATCH(B6,Portfolio!N5:N260,0)'
R['A8']='Committed capital';R['B8']='=INDEX(Portfolio!J5:J260,B7)'
R['A9']='Remaining capital';R['B9']='=Assumptions!B6-B8'
R['A10']='Projects selected';R['B10']='=INDEX(Portfolio!K5:K260,B7)'
R['A11']='Recommendation';R['B11']='=IF(B10=0,"No eligible positive-NPV portfolio — retain capital",IF(Assumptions!B5="YES","Illustration only; obtain actual opportunities and evidence before investment","Selected projects require final due diligence and transaction approval"))'
header(R,14,['Order','Project','NPV','IRR','All-in capital','Ready (months)','Downside NPV','Compliance','Selection'])
for i in range(8):
 rr=15+i;orr=8+i
 R.cell(rr,1,f'=IF(Opportunities!S{orr}="ELIGIBLE",Opportunities!T{orr},"—")')
 for c,oc in [(2,'B'),(3,'P'),(4,'Q'),(5,'L'),(6,'H'),(7,'R'),(8,'M')]:R.cell(rr,c,f'=Opportunities!{oc}{orr}')
 R.cell(rr,9,f'=IF(Opportunities!U{orr}=1,"SELECTED",IF(Opportunities!S{orr}="ELIGIBLE","WATCHLIST","EXCLUDED"))')
R['A25']='Portfolio interpretation';R['A26']='Only compliant, positive-NPV projects within the delay limit enter the portfolio. A high individual rank does not guarantee selection.'
R['A28']='Key risks';R['A29']='Rent and occupancy, construction overrun, permit delay, terminal valuation and legal compliance. Downside is a sensitivity, not a probability forecast.'
R['A31']='Scope';R['A32']='No source spreadsheet was supplied. Examples are fictional. NPV measures discounted profit; IRR is supplemental. Financing and investor-specific taxes are excluded.'
G=wb['Guide'];title(G,'How to use this workbook','Editable blue cells • formula-driven analysis • no macros or external links')
notes=[('1. Import','Replace the eight input rows in Opportunities with the supplied spreadsheet. Preserve calculated columns L:U.'),('2. Set assumptions','Change budget, required return, delay limit and scenario assumptions. The horizon is fixed at ten years.'),('3. Validate compliance','Replace Demo pass by Verified / Unknown / Fail, and cite actual evidence in each evidence row. Switch Demo mode to NO for real evaluation.'),('4. Review','Report gives NPV, annual IRR, downside and exact optimal portfolio. Excluded assets remain visible for audit.'),('5. Rank','Rank is descending base NPV among eligible projects; identical values share a rank. Portfolio selection maximizes total NPV.'),('Method','Acquisition plus contingency capital is committed immediately. Rent and operating cost inflation use the same rate; operating costs begin at rental readiness.'),('Limitations','Annual timing is approximate. Add pre-opening carrying costs to Other capital costs. No leverage, tax, capex phasing or replacement capex is modeled.'),('Capacity','The model supports exactly eight projects; enumeration covers every subset, including retaining all capital.'),('Saving','Run python3 build_model.py to rebuild the template. This overwrites user edits; retain a copy before rerunning.'),('Source status','Legal controls are validation categories, not proof of compliance. Follow official sources and obtain building-specific authority confirmation.')]
header(G,4,['Step','Instruction'])
for r,v in enumerate(notes,5):G.cell(r,1,v[0]);G.cell(r,2,v[1])
S=wb['Sources'];title(S,'Sources and verification','Base version: official legal references will be added after initial save. All financial inputs are assumptions.')
header(S,4,['Topic','Source / URL','Use / caveat','Checked'])
for ws in wb:
 ws.freeze_panes='C8' if ws==O else 'B5';ws.sheet_view.showGridLines=False
 ws.column_dimensions['A'].width=25
 for col in range(2,ws.max_column+1):ws.column_dimensions[__import__('openpyxl').utils.get_column_letter(col)].width=20
 for row in ws.iter_rows(min_row=3):
  for cell in row:
   cell.alignment=Alignment(vertical='top',wrap_text=True)
   if cell.data_type=='f':cell.font=Font(color='222222')
 ws.print_options.horizontalCentered=True;ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
 ws.print_title_rows='1:4'
for ws in [G,S,C]:ws.column_dimensions['B'].width=48;ws.column_dimensions['C'].width=75
G.column_dimensions['B'].width=120;A.column_dimensions['C'].width=90;R.column_dimensions['B'].width=42;C.column_dimensions['E'].width=55;O.column_dimensions['B'].width=35
for ws in [O,F,P,R]:
 for row in ws.iter_rows(min_row=5):
  for cell in row:
   if cell.data_type=='f' or isinstance(cell.value,(int,float)):cell.number_format='#,##0;[Red](#,##0);—'
for r in range(8,16):
 O.cell(r,17).number_format='0.0%'
 for c in [9,10]:O.cell(r,c).number_format='0.0%'
for r in range(15,23):R.cell(r,4).number_format='0.0%'
R.row_dimensions[26].height=36;R.row_dimensions[29].height=36;R.row_dimensions[32].height=48
for r in [26,29,32]:R.merge_cells(start_row=r,start_column=1,end_row=r,end_column=9)
O.auto_filter.ref='A7:U15';C.auto_filter.ref='A4:E100'
demo=DataValidation(type='list',formula1='"YES,NO"');A.add_data_validation(demo);demo.add('B5')
A['B9'].protection=__import__('openpyxl').styles.Protection(locked=True)
from enhance_model import enhance
enhance(wb)
# Atomic replacement: never leave a half-written target workbook.
def save_atomic():
 tmp=ROOT/'model.pending.xlsx';wb.save(tmp)
 from cache_model import cache_workbook
 cache_workbook(wb,tmp)
 load_workbook(tmp).close();os.replace(tmp,OUT)
save_atomic()
print('SAVED',OUT,'bytes',OUT.stat().st_size)
print('SAVE_NOW', (ROOT/'SAVE_NOW.txt').exists())
