from pathlib import Path
from openpyxl import load_workbook
from cache_model import cache_workbook
ROOT=Path(__file__).resolve().parent
results=[]
def scenario(label,edit,check):
 if (ROOT/'SAVE_NOW.txt').exists():return
 wb=load_workbook(ROOT/'model.xlsx');edit(wb)
 tmp=ROOT/'test.pending.xlsx';wb.save(tmp);cache_workbook(wb,tmp)
 result=load_workbook(tmp,data_only=True);check(result);results.append(label+' PASS');tmp.unlink()
def equal(a,b):assert a==b,(a,b)
scenario('Zero budget selects nothing',lambda w:setattr(w['Assumptions']['B6'],'value',0),lambda w:equal(w['Report']['B10'].value,0))
scenario('Live mode rejects demonstration evidence',lambda w:setattr(w['Assumptions']['B5'],'value','NO'),lambda w:equal(w['Report']['B10'].value,0))
scenario('Unknown evidence blocks project',lambda w:setattr(w['Compliance']['D5'],'value','Unknown'),lambda w:equal(w['Opportunities']['U8'].value,0))
scenario('One-project cap respected',lambda w:setattr(w['Assumptions']['B7'],'value',1),lambda w:equal(w['Report']['B10'].value,1))
base=load_workbook(ROOT/'model.xlsx',data_only=True)['Opportunities']['P8'].value
scenario('Longer completion lowers NPV',lambda w:setattr(w['Opportunities']['H8'],'value',18),lambda w:equal(w['Opportunities']['P8'].value<base,True))
(ROOT/'tests.txt').write_text('\n'.join(results)+'\n')
print('\n'.join(results))
