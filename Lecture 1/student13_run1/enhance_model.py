from openpyxl.styles import Font, PatternFill
from openpyxl.chart import BarChart, Reference

def enhance(wb):
 A,O,C,F,R,S,G=[wb[n] for n in ['Assumptions','Opportunities','Compliance','Cash flows','Report','Sources','Guide']]
 A['B6']=2000000;A['B6'].number_format='#,##0 "EUR"';A['B6'].font=Font(color='1B65A0')
 for r,p in enumerate([500000,350000,700000,200000,600000,500000],8):O.cell(r,5,p).font=Font(color='1B65A0')
 rules={2:'Student accommodation: 12 m² for one resident; 18 m² for two. Shared area counted pro rata. Validate legal classification, habitable share and height. Source S2, art. 4.',3:'Shared WC and shower: maximum six users each; kitchen sink and four-ring hob: maximum eight users. Use contractual maximum occupancy. Source S2, art. 5.',4:'Bath/shower with hot water and drainage; shared facilities accessible and lockable. Source S2, art. 5.',5:'Operational cooking equipment in communal kitchen; suitable supply. Source S2, art. 5.',6:'Daylight and permanent base ventilation; verify windows, airflow and configuration against S2, art. 3.',7:'Safe electrical/gas installation; no type B gas heater in bedroom. Source S1/S2.',9:'Water, installed heating, dry structure, sound drainage and no harmful infestation; source S2, arts. 2–5.',11:'Verify access, ceiling heights, structure, mailbox and doorbell, and all case-specific rules. Source S2; confirm municipal requirements.'}
 for i in range(8):
  for j,text in rules.items():C.cell(5+12*i+j,3,text)
  C.cell(5+12*i+1,3,'Check required planning/letting approvals; DIRL clearance required after a letting prohibition. Ordinary voluntary conformity certificate and quality label are not universal prerequisites. Source S3/S4.')
 for i in range(8):
  r=8+i;fr=18+i
  formula=F.cell(fr,12).value
  old=f'(Opportunities!O{r}*(1-Assumptions!$B$15))'
  new=f'((Opportunities!C{r}*Opportunities!D{r}*12*Opportunities!I{r}*(1-Opportunities!J{r})*(1-Assumptions!$B$15)-Opportunities!K{r})*(1+Assumptions!$B$10)^10)'
  F.cell(fr,12,formula.replace(old,new))
  # Unique ranking; stable row-order tie break.
  O.cell(r,20,f'=IF(S{r}="ELIGIBLE",1+COUNTIFS($S$8:$S$15,"ELIGIBLE",$P$8:$P$15,">"&P{r})+COUNTIFS($S$8:S{r},"ELIGIBLE",$P$8:P{r},P{r})-1,"")')
  rr=15+i;k=i+1
  R.cell(rr,1,f'=IF({k}<=COUNTIF(Opportunities!$S$8:$S$15,"ELIGIBLE"),{k},"")')
  for c,oc in [(2,'B'),(3,'P'),(4,'Q'),(5,'L'),(6,'H'),(7,'R'),(8,'M')]:
   R.cell(rr,c,f'=IF($A{rr}="","",INDEX(Opportunities!${oc}$8:${oc}$15,MATCH($A{rr},Opportunities!$T$8:$T$15,0)))')
  R.cell(rr,9,f'=IF($A{rr}="","",IF(INDEX(Opportunities!$U$8:$U$15,MATCH($A{rr},Opportunities!$T$8:$T$15,0))=1,"SELECTED","WATCHLIST"))')
 R['A25']='Report inclusion policy';R['A26']='Shortlist shows eligible positive-NPV investments in descending NPV order. Exclusions and raw inputs remain in Opportunities for audit. Portfolio selection respects budget and project count.'
 R['A34']='Excluded / incomplete opportunities';R['B34']='=COUNTIF(Opportunities!S8:S15,"EXCLUDE")';R['A35']='Portfolio downside NPV';R['B35']='='+'+'.join(f'Opportunities!U{r}*Opportunities!R{r}' for r in range(8,16))
 R['A36']='Undiscounted portfolio profit';R['B36']='='+'+'.join(f'Opportunities!U{8+i}*SUM(\'Cash flows\'!B{5+i}:L{5+i})' for i in range(8))
 R['A38']='Independent analysis • Institutional report format • No affiliation with Vanguard.'
 S['A2']='Official Brussels references checked 14 October 2026. Financial inputs remain illustrative assumptions.'
 sources=[('S1 • 2026 changes','https://be.brussels/en/housing/new-minimum-quality-standards-2026','Effective 1 January 2026; minimum area, heating and bathing requirements.'),('S2 • Consolidated decree','https://www.slrb-bghm.brussels/sites/default/files/2026-02/arr%C3%AAt%C3%A9%20exigences%20%C3%A9l%C3%A9mentaires%20s%C3%A9curit%C3%A9%20%20texte%20coordonn%C3%A9%20fr%2001012026.pdf','30 November 2023 decree, coordinated at 1 January 2026. Articles 2–5: safety, sanitation, area, facilities.'),('S3 • DIRL','https://be.brussels/en/about-region/structure-and-organisations/overview-administrations-and-institutions-region/brussels-housing/structure-brussels-housing/regional-housing-inspectorate-dirl','Inspection framework; voluntary conformity certificate and clearance after a letting ban.'),('S4 • Student quality label','https://be.brussels/en/housing/rental/conformity-for-rental-properties/ensuring-quality-student-housing','Label application, student lease and valid PEB. Label is separate from mandatory dwelling standards.'),('S5 • Official comparison','https://admin.be.brussels/en/bhg/document/download/4537','Official comparison of pre-2026 and 2026 minimum requirements.'),('Financial assumptions','Assumptions / Opportunities','No listings or market rents supplied. Prices, rent, tax/cost proxies and cap rate are fictional planning inputs.')]
 for r,row in enumerate(sources,5):
  for c,val in enumerate((*row,'2026-10-14'),1):S.cell(r,c,val)
  if row[1].startswith('https:'):S.cell(r,2).hyperlink=row[1];S.cell(r,2).style='Hyperlink'
  S.row_dimensions[r].height=65
 S.column_dimensions['B'].width=75;S.column_dimensions['C'].width=70
 G['B8']='Shortlist ranks eligible assets by NPV. Exact enumeration selects the best subset; excluded assets remain in Opportunities.'
 G['B9']='Rank uses descending base NPV; ties are resolved by input row order. Portfolio selection maximizes aggregate NPV.'
 G['B14']='See Sources for 2026 official requirements. Verify every room and shared facility; project-wide Verified means all spaces passed, with linked evidence.'
 for r in range(5,15):G.row_dimensions[r].height=42
 for r in range(5,101):C.row_dimensions[r].height=54
 for r in range(15,23):R.row_dimensions[r].height=36
 chart=BarChart();chart.title='Eligible opportunities | base and downside NPV';chart.y_axis.title='EUR';chart.x_axis.title='Opportunity';chart.add_data(Reference(R,min_col=3,max_col=3,min_row=14,max_row=22),titles_from_data=True);chart.add_data(Reference(R,min_col=7,max_col=7,min_row=14,max_row=22),titles_from_data=True);chart.set_categories(Reference(R,min_col=2,min_row=15,max_row=22));chart.width=24;chart.height=10;R.add_chart(chart,'A41')
 R.print_area='A1:I60';O.print_area='A1:U15';R.column_dimensions['A'].width=30
