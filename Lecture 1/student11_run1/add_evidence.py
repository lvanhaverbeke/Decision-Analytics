from pathlib import Path
import os,tempfile
from openpyxl import load_workbook
p=Path(__file__).resolve().parent
w=load_workbook(p/'model.xlsx');s=w.create_sheet('Sources');s.append(['Source','URL','Evidence / limitation','Accessed'])
rows=[('City Brussels 2021 recommendations','https://www.bruxelles.be/sites/default/files/bxl/Recommandations_logements_etudiants_VBX_17juin2021.pdf','20% social student units from 30 units; policy applicability must be confirmed','2026-10-14'),('Ixelles municipal magazine March–April 2025','https://www.ixelles.be/site/downloads/publications/infoxl-2014-/infoxl20250304.pdf','15% moderate-rent units for large kot developments; size trigger not specified in article','2026-10-14'),('Perspective monitoring January 2026','https://perspective.brussels/fr/actualites/un-etudiant-sur-trois-trouve-un-kot-dans-un-logement-qui-lui-est-destine','11,305 dedicated units in 2025; around 37,000 housed students in 2022–2023; different reference years','2026-10-14'),('Kotkompas pricing 2025–2026','https://www.diggitstudentlife.eu/sites/default/files/2025-11/Kotkompas%20Pricing%20Update%202025Q4%20-%20E_0.pdf','Brussels average base rent EUR570 + EUR125 costs = EUR695 total; indicative industry sample','2026-10-14')]
for r in rows:s.append(r)
for row in w['Municipalities'].iter_rows(min_row=2):
 if row[0].value=='Bruxelles-Ville':row[1].value=.20;row[3].value='Policy located: >=30 units; current application needs confirmation'
 if row[0].value=='Ixelles':row[1].value=.15;row[3].value='Policy located: large projects; trigger/current application needs confirmation'
w['Iterations'].append([2,'Primary policy and market evidence added','Different municipal policies found; other minima remain unknown','Build DCF and exact portfolio search'])
w['Executive Report']['B2']='1.1 – sourced baseline'
fd,t=tempfile.mkstemp(suffix='.xlsx',dir=p);os.close(fd);w.save(t);os.replace(t,p/'model.xlsx')
print('Evidence version saved. SAVE_NOW:',(p/'SAVE_NOW.txt').exists())
