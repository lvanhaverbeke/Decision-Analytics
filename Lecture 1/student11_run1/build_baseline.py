from pathlib import Path
from datetime import datetime
import os, tempfile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.properties import CalcProperties
ROOT=Path(__file__).resolve().parent
OUT=ROOT/'model.xlsx'
wb=Workbook(); report=wb.active; report.title='Executive Report'
report.append(['BRUSSELS STUDENT HOUSING | Decision Analytics'])
report.append(['Version','1.0 – usable baseline'])
report.append(['Decision','Compare investments in affordable student housing; inputs are illustrative.'])
report.append(['Important','Municipal legal thresholds are unverified; no legal compliance claim.'])
report.append(['Method','Iterate: baseline → evaluate → financial model → evaluate → recommendation.'])
s=wb.create_sheet('Inputs'); s.append(['Input','Value','Meaning'])
for r in [('Budget',15000000,'EUR; illustrative'),('Annual occupancy',.95,'Illustrative'),('Operating cost share',.30,'Share of collected rent'),('Planning affordable share',.25,'Assumption, NOT municipal law')]:s.append(r)
s=wb.create_sheet('Projects');s.append(['Project','Municipality','Beds','Affordable beds','Monthly affordable rent','Monthly market rent','Total investment','Select 0/1','Annual net operating income'])
for row in [('A','Bruxelles-Ville',100,40,450,700,4800000,1),('B','Ixelles',80,24,450,750,4300000,1),('C','Anderlecht',120,60,400,650,4900000,1)]:
 s.append(row);i=s.max_row;s.cell(i,9,f'=(D{i}*E{i}+(C{i}-D{i})*F{i})*12*Inputs!B3*(1-Inputs!B4)')
report.append(['Selected investment',"=SUMPRODUCT(Projects!G2:G4,Projects!H2:H4)"])
report.append(['Selected beds',"=SUMPRODUCT(Projects!C2:C4,Projects!H2:H4)"])
report.append(['Selected affordable beds',"=SUMPRODUCT(Projects!D2:D4,Projects!H2:H4)"])
report.append(['Annual NOI',"=SUMPRODUCT(Projects!I2:I4,Projects!H2:H4)"])
s=wb.create_sheet('Municipalities');s.append(['Municipality','Verified legal minimum','Planning share','Verification status'])
for n in ['Anderlecht','Auderghem','Berchem-Sainte-Agathe','Bruxelles-Ville','Etterbeek','Evere','Forest','Ganshoren','Ixelles','Jette','Koekelberg','Molenbeek-Saint-Jean','Saint-Gilles','Saint-Josse-ten-Noode','Schaerbeek','Uccle','Watermael-Boitsfort','Woluwe-Saint-Lambert','Woluwe-Saint-Pierre']:s.append([n,None,.25,'UNVERIFIED: check local/project rules'])
s=wb.create_sheet('Iterations');s.append(['Iteration','Change','Evaluation','Next investment of effort']);s.append([1,'Baseline inputs and project comparison','Usable formulas; legal minima not verified','Verify evidence; add cash-flow and scenarios'])
for sh in wb:
 sh.freeze_panes='B2';sh.auto_filter.ref=sh.dimensions
 for c in sh[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
 for col in sh.columns:sh.column_dimensions[col[0].column_letter].width=min(65,max(16,max(len(str(c.value or '')) for c in col)+2))
wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
fd,tmp=tempfile.mkstemp(suffix='.xlsx',dir=ROOT);os.close(fd)
try:wb.save(tmp);os.replace(tmp,OUT)
finally:
 if os.path.exists(tmp):os.unlink(tmp)
print('Saved baseline:',OUT,'bytes',OUT.stat().st_size)
print('SAVE_NOW:',(ROOT/'SAVE_NOW.txt').exists())
