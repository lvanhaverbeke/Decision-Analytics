from pathlib import Path
import os,tempfile,zipfile,xml.etree.ElementTree as ET,math,json
from itertools import product
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import ColorScaleRule,CellIsRule
from openpyxl.chart import BarChart,Reference
from openpyxl.workbook.properties import CalcProperties
ROOT=Path(__file__).resolve().parent
if (ROOT/'SAVE_NOW.txt').exists():
 print('SAVE_NOW present; preserving saved model.xlsx');raise SystemExit
old=load_workbook(ROOT/'model.xlsx');sources=list(old['Sources'].values)[1:] if 'Sources' in old.sheetnames else []
wb=Workbook();wb.remove(wb.active);cache={}
for n in ['Executive Report','Guide','Inputs','Municipalities','Projects','Cash Flows','Portfolios','Scenarios','Iterations','Sources','Checks']:wb.create_sheet(n)
def row(sh,values):wb[sh].append(values)
def f(sh,cell,formula,value):wb[sh][cell]='='+formula;cache[(sh,cell)]=value
money='#,##0 "€";[Red](#,##0) "€"'
inputs=[('Capital budget',15000000,'EUR; total development outlay; illustrative'),('Minimum affordable beds',80,'Portfolio target; illustrative'),('Base occupancy',.95,'Same for both rent classes; illustrative'),('Months billed',12,'Annual contracts; illustrative'),('Variable operating share',.15,'Share of collected base rents; illustrative'),('Fixed annual operating cost / bed',1600,'EUR; includes reserve, insurance, maintenance; illustrative'),('Rent and operating growth',.02,'Annual; both rent and costs grow at same rate'),('Discount rate',.06,'Unlevered nominal annual hurdle; illustrative'),('Exit capitalization rate',.055,'Terminal value on year 21 NOI; illustrative'),('Sale cost share',.02,'Terminal sale proceeds haircut'),('Hold years',20,'Fixed 20-year DCF; change code to alter horizon'),('Affordable all-in rent ceiling',550,'EUR/month; course definition, not legal social-rent definition'),('Tenant monthly service charges',100,'EUR; pass-through excluded from investor revenue'),('Market base-rent benchmark',570,'Kotkompas 2025–2026 Brussels average; context only'),('Default unverified planning share',.25,'Scenario assumption; never treated as law'),('Construction overrun',0,'Base scenario additional capex %'),('Portfolio minimum NPV',0,'Unlevered; positive-NPV requirement')]
row('Inputs',['Parameter','Value','Definition / evidence'])
for x in inputs:row('Inputs',x)
v={i+2:x[1] for i,x in enumerate(inputs)}
row('Guide',['HOW TO USE','Decision model v2.0 | 14 October 2026'])
for x in [('1','Read Executive Report. All eight developments are hypothetical investment opportunities.'),('2','Blue cells are editable assumptions; green cells hold external evidence; black cells are formulas.'),('3','Edit Inputs, project costs/rents/beds, municipal planning percentages and policy applicability.'),('4','Projects H is manual selection 0/1. Dashboard compares manual selection with the live optimal recommendation.'),('5','Portfolios enumerates all 256 binary combinations: maximize affordable beds, then NPV at equal beds.'),('6','Feasible requires budget, at least target affordable beds, nonnegative portfolio NPV and all selected projects passing planning/rent checks.'),('7','Municipal source evidence is distinct from verified legal approval. Unknown requirements do not mean zero requirements.'),('8','Change project policy applicability to 1 only when a cited policy applies; social allocation and rent must be validated separately.'),('9','Recommendation is conditional feasibility analysis. Obtain municipal confirmation, quotes, demand studies and due diligence before investment.'),('10','Reopen in Excel to recalculate edits. Run python3 build_model.py to regenerate defaults; it overwrites manual edits.'),('Scope','Unlevered pre-tax cash flow; no borrowing, grants, corporate tax, construction delays or VAT separately. Total cost inputs must include acquisition, fees, taxes and contingency.'),('Affordability','Model affordable = designated beds with base rent + service charges <= ceiling. Legal social/moderate rent eligibility is a separate unresolved gate.'),('Timing','Total cost at t0; first operational NOI at t1; exit end t20 using t21 NOI. Assumes immediate completion.'),('Authors','Nisa, Maite, Ilias — names supplied with uncertainty; confirm before course submission.')]:row('Guide',x)
municipal=['Anderlecht','Auderghem','Berchem-Sainte-Agathe','Bruxelles-Ville','Etterbeek','Evere','Forest','Ganshoren','Ixelles','Jette','Koekelberg','Molenbeek-Saint-Jean','Saint-Gilles','Saint-Josse-ten-Noode','Schaerbeek','Uccle','Watermael-Boitsfort','Woluwe-Saint-Lambert','Woluwe-Saint-Pierre']
row('Municipalities',['Municipality','Sourced policy share','Planning share (editable)','Activation / scope','Evidence status','Source'])
for i,n in enumerate(municipal,2):
 share=.2 if n=='Bruxelles-Ville' else .15 if n=='Ixelles' else None
 scope='>=30 student units; social allocation' if n=='Bruxelles-Ville' else 'Large kot projects; trigger needs confirmation' if n=='Ixelles' else 'Unknown; request project-specific minimum'
 row('Municipalities',[n,share,.25,scope,'POLICY FOUND; application not verified' if share else 'UNVERIFIED',sources[0 if n=='Bruxelles-Ville' else 1][1] if share and sources else 'Not yet located'])
# Cost includes all development outlay, projects are hypothetical.
data=[('A','Bruxelles-Ville',100,40,440,680,4800000,1),('B','Ixelles',80,24,450,730,4300000,1),('C','Anderlecht',120,60,400,620,4900000,0),('D','Etterbeek',70,21,440,700,3400000,0),('E','Schaerbeek',90,45,410,640,3800000,0),('F','Jette',80,40,400,610,3200000,0),('G','Molenbeek-Saint-Jean',110,66,390,600,4100000,0),('H','Woluwe-Saint-Lambert',90,27,450,720,4500000,0)]
row('Projects',['ID','Municipality','Beds','Affordable beds','Affordable base rent/mo','Market base rent/mo','Total development cost','Manual select 0/1','Policy applies 0/1','Effective planning minimum','Required affordable beds','Year 1 collected rent','Year 1 NOI','Cost with overrun','NPV incl exit','NOI yield','Planning & rent pass','Recommended select','Affordable all-in rent'])
metrics=[]
for i,d in enumerate(data,2):
 a,n,beds,aff,ar,mr,cost,applies=d;m=municipal.index(n)+2;minshare=.25;req=math.ceil(beds*minshare);rev=(aff*ar+(beds-aff)*mr)*12*.95;noi=rev*.85-beds*1600
 terminal=noi*1.02**20/.055*.98
 npv=-cost+sum(noi*1.02**(t-1)/1.06**t for t in range(1,21))+terminal/1.06**20
 good=aff>=req and ar+100<=550 and aff<=beds and min(ar,mr,cost)>0
 metrics.append(dict(beds=beds,aff=aff,cost=cost,noi=noi,npv=npv,good=good))
 row('Projects',[a,n,beds,aff,ar,mr,cost,0,applies])
 f('Projects',f'J{i}',f'MAX(Municipalities!C{m},IF(I{i}=1,Municipalities!B{m},0))',minshare)
 f('Projects',f'K{i}',f'ROUNDUP(C{i}*J{i},0)',req)
 f('Projects',f'L{i}',f'(D{i}*E{i}+(C{i}-D{i})*F{i})*Inputs!B5*Inputs!B4',rev)
 f('Projects',f'M{i}',f'L{i}*(1-Inputs!B6)-C{i}*Inputs!B7',noi)
 f('Projects',f'N{i}',f'G{i}*(1+Inputs!B17)',cost)
 f('Projects',f'O{i}',f'SUM(\'Cash Flows\'!D{(i-2)*21+2}:D{(i-2)*21+22})',npv)
 f('Projects',f'P{i}',f'IFERROR(M{i}/N{i},0)',noi/cost)
 f('Projects',f'Q{i}',f'IF(AND(D{i}>=K{i},D{i}<=C{i},E{i}+Inputs!B14<=Inputs!B13,C{i}>0,D{i}>=0,E{i}>0,F{i}>0,G{i}>0),1,0)',int(good))
 f('Projects',f'S{i}',f'E{i}+Inputs!B14',ar+100)
 for t in range(21):
  r=(i-2)*21+t+2;row('Cash Flows',[a,t]);cash=-cost if t==0 else noi*1.02**(t-1)+(terminal if t==20 else 0)
  form=f'-Projects!N{i}' if t==0 else f'Projects!M{i}*(1+Inputs!B8)^(B{r}-1)'+(f'+Projects!M{i}*(1+Inputs!B8)^20/Inputs!B10*(1-Inputs!B11)' if t==20 else '')
  f('Cash Flows',f'C{r}',form,cash);f('Cash Flows',f'D{r}',f'C{r}/(1+Inputs!B9)^B{r}',cash/1.06**t)
wb['Cash Flows'].insert_rows(1) if False else None
# Rows already start at 1: add header originally absent, fix by putting header before generating in next line via explicit relocation.
# Data rows were appended at row 1 whereas formulas use row 2. Rewrite cash flows cleanly.
cs=wb['Cash Flows'];cs.delete_rows(1,cs.max_row);row('Cash Flows',['Project','Year','Unlevered cash flow','Discounted cash flow'])
for i,d in enumerate(data,2):
 noi=metrics[i-2]['noi'];cost=d[6];terminal=noi*1.02**20/.055*.98
 for t in range(21):
  r=(i-2)*21+t+2;row('Cash Flows',[d[0],t]);cash=-cost if t==0 else noi*1.02**(t-1)+(terminal if t==20 else 0)
  form=f'-Projects!N{i}' if t==0 else f'Projects!M{i}*(1+Inputs!B8)^(B{r}-1)'+(f'+Projects!M{i}*(1+Inputs!B8)^20/Inputs!B10*(1-Inputs!B11)' if t==20 else '')
  f('Cash Flows',f'C{r}',form,cash);f('Cash Flows',f'D{r}',f'C{r}/(1+Inputs!B9)^B{r}',cash/1.06**t)
row('Portfolios',['Combination']+[d[0] for d in data]+['Cost','Beds','Affordable beds','NPV','Selected planning failures','Feasible','Affordable objective','NPV tie break'])
combos=[]
for r,bits in enumerate(product([0,1],repeat=8),2):
 cost=sum(x*m['cost'] for x,m in zip(bits,metrics));beds=sum(x*m['beds'] for x,m in zip(bits,metrics));aff=sum(x*m['aff'] for x,m in zip(bits,metrics));npv=sum(x*m['npv'] for x,m in zip(bits,metrics));fails=sum(x*(not m['good']) for x,m in zip(bits,metrics));feas=int(cost<=15000000 and aff>=80 and npv>=0 and fails==0)
 combos.append((bits,cost,beds,aff,npv,feas));row('Portfolios',[r-2]+list(bits))
 for col,key,projectcol,val in [('J','cost','N',cost),('K','beds','C',beds),('L','aff','D',aff),('M','npv','O',npv)]:f('Portfolios',f'{col}{r}','+'.join(f'{chr(66+k)}{r}*Projects!{projectcol}{k+2}' for k in range(8)),val)
 f('Portfolios',f'N{r}','+'.join(f'{chr(66+k)}{r}*(1-Projects!Q{k+2})' for k in range(8)),fails)
 f('Portfolios',f'O{r}',f'IF(AND(J{r}<=Inputs!B2,L{r}>=Inputs!B3,M{r}>=Inputs!B18,N{r}=0),1,0)',feas)
 f('Portfolios',f'P{r}',f'IF(O{r}=1,L{r},-1)',aff if feas else -1)
best=max((c for c in combos if c[5]),key=lambda c:(c[3],c[4]));bestidx=combos.index(best)+2;bestaff=best[3]
for r,c in enumerate(combos,2):f('Portfolios',f'Q{r}',f'IF(AND(O{r}=1,L{r}=MAX(P$2:P$257)),M{r},-1E+99)',c[4] if c[5] and c[3]==bestaff else -1e99)
for i,d in enumerate(data,2):
 f('Projects',f'R{i}',f'IF(\'Executive Report\'!B6=0,0,INDEX(Portfolios!{chr(66+i-2)}$2:{chr(66+i-2)}$257,\'Executive Report\'!B7))',best[0][i-2]);wb['Projects'][f'H{i}']=best[0][i-2]
row('Executive Report',['BRUSSELS STUDENT HOUSING','EXECUTIVE DECISION REPORT | v2.0'])
row('Executive Report',['Recommendation','Conditional shortlist: maximize affordable beds under the EUR15m planning budget.'])
row('Executive Report',['Evidence','Brussels average total rent EUR695 in Kotkompas 2025–2026; designated supply is constrained (Sources).'])
row('Executive Report',['Approval gate','Municipal and project policy applicability, social eligibility and quotes must be confirmed before commitment.'])
row('Executive Report',['Decision metric','Affordable beds first, NPV second; no probability weighting; projects are illustrative.'])
f('Executive Report','B6','SUM(Portfolios!O2:O257)',sum(c[5] for c in combos));wb['Executive Report']['A6']='Feasible combinations'
f('Executive Report','B7','MATCH(MAX(Portfolios!Q2:Q257),Portfolios!Q2:Q257,0)',bestidx-1);wb['Executive Report']['A7']='Recommended combination index'
row('Executive Report',['KPI','Recommended','Manual selection'])
for r,label,pcol,key in [(9,'Total investment','N',1),(10,'Total beds','C',2),(11,'Affordable beds','D',3),(12,'NPV incl terminal value','O',4),(13,'Year 1 NOI','M',None)]:
 wb['Executive Report'].cell(r,1,label)
 vals=[m['noi'] for m in metrics] if key is None else None
 value=sum(b*m['noi'] for b,m in zip(best[0],metrics)) if key is None else best[key]
 for col,sel in [('B','R'),('C','H')]:f('Executive Report',f'{col}{r}',f'SUMPRODUCT(Projects!{pcol}2:{pcol}9,Projects!{sel}2:{sel}9)',value)
f('Executive Report','B14','IFERROR(B13/B9,0)',sum(b*m['noi'] for b,m in zip(best[0],metrics))/best[1]);wb['Executive Report']['A14']='Year 1 NOI yield'
f('Executive Report','B15','Inputs!B2-B9',15000000-best[1]);wb['Executive Report']['A15']='Budget headroom'
f('Executive Report','C16','IF(AND(C9<=Inputs!B2,C11>=Inputs!B3,C12>=Inputs!B18,SUMPRODUCT(Projects!H2:H9,Projects!Q2:Q9)=SUM(Projects!H2:H9)),"PLANNING FEASIBLE","FAIL")','PLANNING FEASIBLE');wb['Executive Report']['A16']='Manual planning status'
for r,x in enumerate([('Shortlist at saved inputs',', '.join(data[k][0]+' / '+data[k][1] for k,b in enumerate(best[0]) if b)),('Interpretation','Recommendation updates via Projects R when inputs change; this text records the saved-input shortlist.'),('Affordability','EUR550 all-in ceiling is a course planning definition. It does not establish social-housing legal compliance.'),('Municipal differences','City Brussels policy: 20% from 30 student units. Ixelles: 15% for large projects. Other 17 minima unverified.'),('Risk','Downside: NPV about -EUR0.43m and cost EUR16.675m, exceeding the EUR15m budget. Secure contingency or reduce scope before commitment.'),('Next agile sprint','Validate thresholds and allocation with municipalities/AISE; replace costs with quotes; add delivery schedule and debt only after evidence.'),('Decision requested','Approve further feasibility work on the shortlist; capital commitment remains conditional on due diligence.')],18):
 for c,val in enumerate(x,1):wb['Executive Report'].cell(r,c,val)
row('Scenarios',['Scenario','Occupancy','Capex multiplier','Rent multiplier','Discount rate','Exit cap rate','Selected investment','Year 1 NOI','NPV','Affordable beds','Planning budget check'])
scenarios=[('Base',.95,1,1,.06,.055),('Downside',.85,1.15,.95,.075,.065),('Upside',.98,.95,1.05,.05,.05)]
selected=[(d,m) for d,m,b in zip(data,metrics,best[0]) if b]
for r,(n,occ,cap,rm,disc,exitcap) in enumerate(scenarios,2):
 row('Scenarios',[n,occ,cap,rm,disc,exitcap]);rev=sum((d[3]*d[4]+(d[2]-d[3])*d[5])*12 for d,m in selected);beds=best[2];noi=rev*occ*rm*.85-beds*1600;cost=best[1]*cap;npv=-cost+sum(noi*1.02**(t-1)/(1+disc)**t for t in range(1,21))+noi*1.02**20/exitcap*.98/(1+disc)**20
 if r==2:
  for col,ref,val in [('B','B4',occ),('E','B9',disc),('F','B10',exitcap)]:f('Scenarios',f'{col}{r}',f'Inputs!{ref}',val)
 f('Scenarios',f'G{r}',f'\'Executive Report\'!B9*C{r}',cost)
 revenue='+'.join(f'Projects!R{i}*(Projects!D{i}*Projects!E{i}+(Projects!C{i}-Projects!D{i})*Projects!F{i})*Inputs!B5' for i in range(2,10))
 f('Scenarios',f'H{r}',f'({revenue})*B{r}*D{r}*(1-Inputs!B6)-\'Executive Report\'!B10*Inputs!B7',noi)
 pv='+'.join(f'H{r}*(1+Inputs!B8)^{t-1}/(1+E{r})^{t}' for t in range(1,21))
 f('Scenarios',f'I{r}',f'-G{r}+{pv}+H{r}*(1+Inputs!B8)^20/F{r}*(1-Inputs!B11)/(1+E{r})^20',npv)
 f('Scenarios',f'J{r}',"'Executive Report'!B11",best[3]);f('Scenarios',f'K{r}',f'IF(G{r}<=Inputs!B2,"PASS","OVER BUDGET")','PASS' if cost<=15000000 else 'OVER BUDGET')
row('Scenarios',['Interpretation','Fixed recommended base portfolio under changed financial assumptions; no scenario reoptimization.'])
row('Iterations',['Iteration','Increment delivered','Evaluation / feedback','Reinvestment in next sprint','Status'])
for x in [(1,'Baseline saved immediately','Three candidate projects, live NOI and basic municipality registry','Prioritize evidence and policy differences','Saved v1.0'),(2,'Sourced baseline saved','Primary policy evidence: Brussels 20%, Ixelles 15%; no blanket quota found for 19 communes','Add conditional compliance, DCF and exact portfolio enumeration','Saved v1.1'),(3,'8-project financial / portfolio model','256 combinations compared; financial and affordability constraints separated from legal approval','Stress-test selected portfolio; validate formula caches and workbook integrity','Saved v2.0'),(4,'Next course sprint','Pending municipal confirmation, real costs and demand studies','Update assumptions; compare changed recommendation; add financing/construction timing','Future work, not claimed complete')]:row('Iterations',x)
row('Sources',['Source','URL','Evidence / limitations','Accessed'])
for x in sources:row('Sources',x)
row('Sources',['Assumption register','Course model','All project beds, rents, capex and financial inputs are illustrative; only cited policy shares and benchmark are sourced','2026-10-14'])
row('Checks',['Check','Result','Interpretation'])
checks=[('Budget','IF(\'Executive Report\'!B9<=Inputs!B2,"PASS","FAIL")','PASS','Recommended cost <= budget'),('Target','IF(\'Executive Report\'!B11>=Inputs!B3,"PASS","FAIL")','PASS','Affordable beds >= target'),('Search coverage','IF(COUNTA(Portfolios!A2:A257)=256,"PASS","FAIL")','PASS','All 2^8 portfolios'),('Planning validity','IF(SUMPRODUCT(Projects!R2:R9,Projects!Q2:Q9)=SUM(Projects!R2:R9),"PASS","FAIL")','PASS','Every recommended project satisfies planning checks'),('Legal approval',None,'UNVERIFIED','All municipalities need current project-specific confirmation; no approval inferred'),('Formula cache',None,'PASS','Cached values independently computed at build; Excel recalculates on open')]
for r,(n,formula,result,desc) in enumerate(checks,2):
 row('Checks',[n,result,desc])
 if formula:f('Checks',f'B{r}',formula,result)
# Formatting and safe editable inputs.
for sh in wb:
 sh.freeze_panes='B2';sh.sheet_view.showGridLines=False
 for c in sh[1]:c.fill=PatternFill('solid',fgColor='17365D');c.font=Font(name='Calibri',size=11,bold=True,color='FFFFFF');c.alignment=Alignment(wrap_text=True,vertical='center')
 sh.row_dimensions[1].height=34
 for rs in sh.iter_rows(min_row=2):
  for c in rs:
   c.font=Font(name='Calibri',size=11,color='222222');c.alignment=Alignment(vertical='top',wrap_text=True)
   if c.row%2==0:c.fill=PatternFill('solid',fgColor='F0F4F8')
   if isinstance(c.value,(int,float)):c.number_format='#,##0.00' if isinstance(c.value,float) else '#,##0'
 for col in sh.columns:sh.column_dimensions[col[0].column_letter].width=22
 if sh.title not in ['Executive Report','Guide']:sh.auto_filter.ref=sh.dimensions
for name in ['Guide','Executive Report']:
 sh=wb[name];sh.column_dimensions['A'].width=30;sh.column_dimensions['B'].width=105
 for i in range(2,sh.max_row+1):sh.row_dimensions[i].height=34
wb['Executive Report'].column_dimensions['C'].width=25
for name in ['Inputs','Municipalities','Sources','Iterations','Checks']:
 for col in ['C','D','E','F']:wb[name].column_dimensions[col].width=52
for r in range(2,19):wb['Inputs'][f'B{r}'].font=Font(color='0070C0');wb['Inputs'].row_dimensions[r].height=32
for r in [4,6,8,9,10,11,16,17]:wb['Inputs'][f'B{r}'].number_format='0.0%'
for r in range(2,21):
 for c in ['B','C']:wb['Municipalities'][f'{c}{r}'].number_format='0%'
 wb['Municipalities'][f'C{r}'].font=Font(color='0070C0');wb['Municipalities'][f'B{r}'].font=Font(color='008050');wb['Municipalities'].row_dimensions[r].height=45
for r in range(2,10):
 for col in 'CDEFGHI':wb['Projects'][f'{col}{r}'].font=Font(color='0070C0')
 for col in 'EFGLMNOS':wb['Projects'][f'{col}{r}'].number_format=money
 for col in 'JP':wb['Projects'][f'{col}{r}'].number_format='0.0%'
 wb['Projects'].row_dimensions[r].height=32
for r in [9,12,13,15]:
 for c in ['B','C']:wb['Executive Report'][f'{c}{r}'].number_format=money
wb['Executive Report']['B14'].number_format='0.0%'
for name,cols in [('Cash Flows','CD'),('Portfolios','JM'),('Scenarios','GHI')]:
 for c in cols:
  for r in range(2,wb[name].max_row+1):wb[name][f'{c}{r}'].number_format=money
for r in range(2,5):
 for col in 'BCEF':wb['Scenarios'][f'{col}{r}'].number_format='0.0%'
for name,ref,typ,a,b in [('Projects','H2:I9','whole',0,1),('Projects','C2:D9','whole',0,10000),('Municipalities','C2:C20','decimal',0,1)]:
 dv=DataValidation(type=typ,operator='between',formula1=a,formula2=b);dv.errorTitle='Invalid input';dv.error='Use a value within the allowed range.';dv.showErrorMessage=True;wb[name].add_data_validation(dv);dv.add(ref)
wb['Projects'].conditional_formatting.add('O2:O9',ColorScaleRule(start_type='min',start_color='F8D7DA',end_type='max',end_color='B7DFBE'))
chart=BarChart();chart.title='Selected portfolio: scenario NPV';chart.y_axis.title='EUR';chart.add_data(Reference(wb['Scenarios'],min_col=9,min_row=1,max_row=4),titles_from_data=True);chart.set_categories(Reference(wb['Scenarios'],min_col=1,min_row=2,max_row=4));chart.width=22;chart.height=10;wb['Executive Report'].add_chart(chart,'A27')
wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
# Save to sibling temp file, insert formula cached values, verify, then atomically replace.
fd,tmp=tempfile.mkstemp(suffix='.xlsx',dir=ROOT);os.close(fd)
try:
 wb.save(tmp);ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
 with zipfile.ZipFile(tmp) as z:files={n:z.read(n) for n in z.namelist()}
 for idx,sh in enumerate(wb,1):
  name=f'xl/worksheets/sheet{idx}.xml';tree=ET.fromstring(files[name])
  for c in tree.findall('.//s:c',ns):
   key=(sh.title,c.attrib['r'])
   if key in cache:
    val=cache[key];node=c.find('s:v',ns)
    if node is None:node=ET.SubElement(c,'{'+ns['s']+'}v')
    if isinstance(val,str):c.set('t','str');node.text=val
    else:c.pop('t',None) if False else None;node.text=str(val)
  files[name]=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
 with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as z:
  for n,b in files.items():z.writestr(n,b)
 test=load_workbook(tmp,data_only=True)
 assert abs(test['Executive Report']['B9'].value-best[1])<.01
 assert test['Executive Report']['B11'].value==best[3]
 assert test['Executive Report']['B6'].value>0
 assert len(combos)==256 and len(set(c[0] for c in combos))==256
 os.replace(tmp,ROOT/'model.xlsx')
finally:
 if os.path.exists(tmp):os.unlink(tmp)
summary={'selected':[data[k][0] for k,b in enumerate(best[0]) if b],'investment':best[1],'beds':best[2],'affordable_beds':best[3],'npv':best[4],'feasible_combinations':sum(c[5] for c in combos),'formula_cells':len(cache),'save_now':(ROOT/'SAVE_NOW.txt').exists()}
(ROOT/'validation.json').write_text(json.dumps(summary,indent=2));print(json.dumps(summary,indent=2))
