from pathlib import Path
from openpyxl import load_workbook
import json,math,zipfile
p=Path(__file__).resolve().parent
w=load_workbook(p/'model.xlsx',data_only=True);form=load_workbook(p/'model.xlsx',data_only=False)
errors=[]
def check(b,msg):
 if not b:errors.append(msg)
for r in range(2,10):
 begin=(r-2)*21+2
 npv=sum(w['Cash Flows'].cell(t,4).value for t in range(begin,begin+21))
 check(math.isclose(npv,w['Projects'].cell(r,15).value,abs_tol=.01),f'Project {r} DCF reconciliation')
for r in range(2,258):
 bits=[w['Portfolios'].cell(r,c).value for c in range(2,10)]
 for col,pc in [(10,14),(11,3),(12,4),(13,15)]:
  total=sum(b*w['Projects'].cell(k+2,pc).value for k,b in enumerate(bits))
  check(math.isclose(total,w['Portfolios'].cell(r,col).value,abs_tol=.01),f'Portfolio {r} column {col}')
for sh in form:
 for row in sh:
  for c in row:
   if c.data_type=='f':check(w[sh.title][c.coordinate].value is not None,f'Missing cache {sh.title}!{c.coordinate}')
check(math.isclose(w['Scenarios']['I2'].value,w['Executive Report']['B12'].value,abs_tol=.01),'Base scenario NPV reconciliation')
check(math.isclose(w['Scenarios']['H2'].value,w['Executive Report']['B13'].value,abs_tol=.01),'Base scenario NOI reconciliation')
feasible=[(w['Portfolios'].cell(r,12).value,w['Portfolios'].cell(r,13).value,r) for r in range(2,258) if w['Portfolios'].cell(r,15).value==1]
best=max(feasible);check(best[0]==w['Executive Report']['B11'].value,'Optimal affordable beds');check(math.isclose(best[1],w['Executive Report']['B12'].value,abs_tol=.01),'NPV tie break optimum')
with zipfile.ZipFile(p/'model.xlsx') as z:check(z.testzip() is None,'Zip integrity')
result={'errors':errors,'status':'PASS' if not errors else 'FAIL','sheets':len(w.sheetnames),'formula_cells':sum(c.data_type=='f' for s in form for row in s for c in row),'scenario_npvs':{w['Scenarios'].cell(r,1).value:w['Scenarios'].cell(r,9).value for r in range(2,5)},'SAVE_NOW':(p/'SAVE_NOW.txt').exists()}
(p/'verification.json').write_text(json.dumps(result,indent=2));print(json.dumps(result,indent=2));assert not errors
