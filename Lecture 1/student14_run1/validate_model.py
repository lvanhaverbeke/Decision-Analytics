from openpyxl import load_workbook
from pathlib import Path
import math
w=load_workbook('model.xlsx',data_only=True)
f=load_workbook('model.xlsx',data_only=False)
cf=[w['Cash flows'].cell(24,c).value for c in range(2,34)]
r=w['Summary']['B14'].value
npv=sum(v/(1+r)**t for t,v in enumerate(cf))
irr=w['Summary']['B5'].value
assert abs(npv-w['Summary']['B4'].value)<0.01
assert abs(sum(v/(1+irr)**t for t,v in enumerate(cf)))<0.01
assert abs(w['Sensitivity']['D5'].value-npv)<0.01
assert abs(w['Development']['B20'].value+cf[0]+cf[1])<0.01
assert w['Cash flows']['AG3'].value==30
assert w['Cash flows']['M17'].value<0 and w['Cash flows']['W17'].value<0 and w['Cash flows']['AG17'].value==0
assert w['Summary']['B20'].value=='OK'
for s in w:
 for row in s:
  for c in row:assert c.data_type!='e',(s.title,c.coordinate,c.value)
for s in f:
 for row in s:
  for c in row:
   if c.data_type=='f':
    assert w[s.title][c.coordinate].value is not None,(s.title,c.coordinate,'missing cached result')
assert abs(w['Cash flows']['AG28'].value-sum(cf))<0.01
print(f'PASS: NPV {npv:,.2f}, IRR {irr:.4%}, budget/cash-flow reconciliation, sensitivity, replacement timing and formula caches.')
print('SAVE_NOW:',Path('SAVE_NOW.txt').exists())
