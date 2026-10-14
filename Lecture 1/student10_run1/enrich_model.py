from pathlib import Path
import os
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
p=Path(__file__).resolve().parent
if (p/'SAVE_NOW.txt').exists(): print('SAVE_NOW: retaining saved model');raise SystemExit
w=load_workbook(p/'model.xlsx')
sources=[
('S1 Housing quality 2026','https://be.brussels/fr/logement/nouvelles-normes-minimales-de-qualite-partir-de-2026','New minimum standards apply from 1 Jan 2026; official summary gives 12 m² student minimum. Architect must assess actual layout, shared areas and all technical rules.'),
('S2 Safety and fire','https://be.brussels/fr/logement/location/conformite-logement/normes-et-securite','Housing must meet safety, health and equipment standards; smoke detection requirements updated from 2025. Building-specific fire strategy remains to be agreed.'),
('S3 Ixelles planning','https://www.ixelles.be/site/880-Guichet-electronique-Urbanisme','Municipal portal provides planning permit and urban-planning information requests; address-specific permitted use/capacity not verified.'),
('S4 Lease process','https://be.brussels/fr/logement/location/les-differentes-etapes-de-la-location','Regional portal explains disclosure, lease types and reference rents; check current student lease terms and registration procedure before letting.'),
('S5 Excessive rents','https://be.brussels/fr/logement/creation-de-la-commission-paritaire-locative','Landlords must not offer abusive rents from 1 May 2025; CPL offers free nonbinding opinions. A policy affordability ceiling alone does not establish compliance.'),
('S6 PEB works','https://environnement.brussels/pro/reglementation-et-inspection/obligations-et-autorisations/objectifs-peb-pour-chaque-bien-et-exigences-peb-en-cas-de-travaux','PEB requirements depend on use and nature of works; establish classification with PEB adviser and include required measures in budget.'),
('S7 Reference rents','https://loyers.brussels/a-propos-des-loyers-de-reference','Reference rent methodology indexed January 2026. Obtain address/layout-specific assessment and legal interpretation.'),
('S8 Ixelles market context','https://www.ixelles.be/site/downloads/publications/infoxl-2014-/infoxl20250708.pdf','Municipal July/August 2025 dossier describes strong student demand and insufficient supply. It mentions a 15% affordable-kot initiative for large projects; verify adoption, scope and site applicability. This is not an underwriting rent survey.'),
('S9 Housing label','https://be.brussels/en/about-region/structure-and-organisations/overview-administrations-and-institutions-region/brussels-housing/structure-brussels-housing/regional-housing-inspectorate-dirl','DIRL enforces housing standards and awards student quality label. Optional conformity certificate is distinct from mandatory compliance and post-ban attestation.')]
for key,url,note in sources:w['Sources'].append([key,url,'2026-10-14',note])
checks=[
('Planning / permitted use','Obtain title, planning history, approved plans, subdivision/change-of-use and heritage constraints.','Architect + Ixelles written planning position; S3','Capacity, delay, redesign and refusal risk'),
('Housing quality','Verify 2026 standards: net area, height, daylight, ventilation, sanitation, heating and electrical/gas safety.','Architect room schedule and technical reports; S1/S2','Room count and capex'),
('Fire strategy','Confirm evacuation, compartmentation, access, detection and any SIAMU review/certification applicable to this building.','Fire engineer + competent authority; S2','Design and opening date'),
('PEB and environment','Confirm building classification, certificates, works procedure, renovation duties and environmental permit applicability.','PEB adviser + permit assessment; S6','Initial and future capex, utilities'),
('Accessibility / amenity','Verify applicable accessibility, bicycle, waste, noise and shared-facility requirements.','Architect and municipality; S3','Usable area / costs'),
('Student leases','Review student eligibility, duration, termination, inventory, deposits, charges, registration and indexation conditions.','Brussels housing lawyer; S4','Turnover, revenue timing, enforceability'),
('Rent legality','Obtain reference-rent assessment; assess excessive-rent rules and applicable indexation restrictions.','Legal opinion and reference calculation; S5/S7','Rent cap may need reduction'),
('Local affordability obligation','Check whether municipal affordable-kot policy or permit agreement applies and on what terms.','Written Ixelles confirmation; S8','Unit mix / affordable rent obligations'),
('Affordability evidence','Validate rent + utilities against target student resources; confirm no hidden charges and summer contract assumptions.','Student survey, university partners and charge budget','Demand, ethical affordability'),
('Tax / VAT / ownership','Confirm duties, municipal/property taxes, VAT and recovery, transaction structure, corporate or personal income tax.','Notary and tax adviser; acquisition allowance is unverified','After-tax return not yet modelled'),
('Funding','Obtain term sheet and bridge interest/fees; test covenant DSCR and cost overruns.','Lender and quantity surveyor','Liquidity and equity requirement'),
('Demand and comparables','Collect dated comparable asking/achieved rents, vacancy, enrollment, pipeline and travel times.','University/Brik/Mykot data + local inspections','Scenario calibration')]
for a,b,c,d in checks:w['Due diligence'].append([a,b,c,'OPEN — not verified',d])
risk=w.create_sheet('Market and risks');risk.append(['Driver','Quantitative model link','Qualitative assessment / evidence to collect','Mitigation / gate'])
for row in [
('University proximity','Occupancy / rent','Measure walking/cycling/transit times to ULB and VUB; exact site not provided.','Verify catchment before purchase'),
('Demand segmentation','Rent ceiling / budget share','Local, international, grant-supported students have different budgets and contract needs.','Survey intended tenants; avoid dependence on premium rents'),
('Seasonality / turnover','Annual occupancy','12-month billing assumed. Summer vacancies and short contracts reduce effective occupancy.','Underwrite achieved annual collections; stress summer gap'),
('Competitive supply','Occupancy / exit yield','University residences, private kots, shared flats and pipeline projects compete by price and quality.','Dated comparable and pipeline study'),
('Planning uncertainty','Development years / overrun','Permits, objections, heritage and room-layout limits can delay opening or reduce beds.','Conditional acquisition and detailed feasibility'),
('Neighbourhood acceptance','Delay / owner costs','Noise, waste, cycling, visitor management and community relations affect operation.','Management plan and municipal consultation'),
('Energy and ageing assets','Inflation / reserve','Envelope, heating and compliance renewal may cost more than annual allowance.','20-year technical capex plan and contingency'),
('Affordability vs viability','Ceiling / NPV / DSCR','Returns may fail at affordable prices. Do not solve a poor investment by assuming unaffordable rent.','Negotiate purchase/build cost; consider evidenced grants or lower required return'),
('Exit liquidity','Exit yield / exit PV share','Single-use assets and regulation can narrow buyer pool; terminal value is uncertain.','Conservative exit yield; inspect no-sale returns')]:risk.append(row)
s=w.create_sheet('Sensitivity');s.append(['INTERMEDIATE NPV | works cost vs exit yield','Works -20%','Works base','Works +20%'])
s.append(['Exit yield -1pp',-.2,0,.2]);s.append(['Exit yield base',-.2,0,.2]);s.append(['Exit yield +1pp',-.2,0,.2])
# Headers contain assumptions; body is formula outputs.
for r,shift in [(2,-.01),(3,0),(4,.01)]:
 for c,change in [(2,-.2),(3,0),(4,.2)]:
  s.cell(r,c,f"=Dashboard!C2-({change})*SUM(Intermediate!I3:I4)+(Intermediate!G22*(Inputs!C19/(Inputs!C19+({shift}))-1))/(1+Inputs!C21)^20")
  # Development costs in years 1–2 are the only negative discounted CF in these rows.
  s.cell(r,c,f"=Dashboard!C2+({change})*SUM(Intermediate!I3:I4)+(Intermediate!G22*(Inputs!C19/(Inputs!C19+({shift}))-1))/(1+Inputs!C21)^20")
s['A7']='Holds rent and affordability ceiling fixed. Works multiplier includes proportional professional fees; acquisition unchanged.'
s['A8']='Delay and occupancy stress are represented jointly by the three main scenarios. Sensitivity is not a forecast probability.'
s.column_dimensions['A'].width=58
# Operating-only payback and financial checks
D=w['Dashboard'];D.append(['Operating-only payback','Beyond 20','Beyond 20','Beyond 20','Calculated below with sale excluded; full-year convention'])
for c,sc in enumerate(['Bad','Intermediate','Optimal'],2):
 ws=w[sc];ws['Q1']='Cumulative excluding sale';ws['R1']='Discounted cash flow check'
 for r in range(2,23):
  ws.cell(r,17,f'=J{r}-SUM(G$2:G{r})');ws.cell(r,18,f'=H{r}-(E{r}-F{r}-B{r}+G{r})')
 D.cell(19,c,f'=IF(COUNTIF(\'{sc}\'!Q3:Q22,">=0")=0,"Beyond 20",MATCH(TRUE,INDEX(\'{sc}\'!Q3:Q22>=0,0),0))')
D.append(['Regulatory decision gate','OPEN','OPEN','OPEN','No acquisition commitment before all required evidence is cleared'])
for ws in w:
 for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='16324F');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True)
 if ws.title in ['Sources','Due diligence','Market and risks']:
  for c in range(1,ws.max_column+1):ws.column_dimensions[L(c)].width=36 if c==1 else 65
  for r in range(2,ws.max_row+1):ws.row_dimensions[r].height=82
  ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for row in ws:
  for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
 for row in ws.iter_rows(min_row=2):
  for cell in row:
   if cell.data_type=='f':cell.number_format='0.0%' if ws.title=='Dashboard' and cell.row in [3,9,11,13,14,16] else '#,##0.00;[Red](#,##0.00)'
 for row in w['Sources'].iter_rows(min_row=2):row[1].hyperlink=row[1].value;row[1].style='Hyperlink'
tmp=p/'model.enriched.tmp.xlsx';w.save(tmp);os.replace(tmp,p/'model.xlsx')
print('Enriched workbook saved. SAVE_NOW=',(p/'SAVE_NOW.txt').exists())
