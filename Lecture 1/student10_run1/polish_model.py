from pathlib import Path
import os,json
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.formatting.rule import ColorScaleRule,CellIsRule
from openpyxl.chart import BarChart,Reference
from openpyxl.worksheet.datavalidation import DataValidation
p=Path(__file__).resolve().parent
if (p/'SAVE_NOW.txt').exists():print('SAVE_NOW: retaining current saved model');raise SystemExit
w=load_workbook(p/'model.xlsx');d=w['Dashboard'];read=w['Start']
read.append(['Illustrative conclusion','All three scenarios have negative NPV at their assumed discount rates. Intermediate and Optimal payback require the year-20 sale; operating-only payback exceeds 20 years. These results do not describe an identified property.'])
read.append(['Funding warning','All default scenarios have operating DSCR below 1.0, so annual operating cash after reserves does not cover scheduled debt service. Additional equity is required; obtain a revised funding plan.'])
read.append(['Next decision','Replace assumed acquisition price, construction cost and rents with evidence; renegotiate capital cost or assess documented grants. Keep affordability fixed while seeking a viable structure.'])
read.append(['Rebuild files','Run build_model.py, enrich_model.py, polish_model.py, then validate_and_cache.py. Each script saves atomically; validation_results.json records numeric checks.'])
read.append(['Reading payback','Annual end-year payback is deliberately conservative; year 20 includes the sale. Discounted payback uses scenario project discount rate. DSCR is after annual capital expenditure allowance.'])
read.append(['Current evidence boundary','Official web sources checked on 14 October 2026. No site, drawings, rent comparables, lender term sheet or construction quotes were provided.'])
read.append(['Scenario labels','Bad, Intermediate and Optimal are stress cases, not probabilities or best-possible achievable outcomes. Rooms and affordable headline rent stay constant across scenarios.'])
for r in range(2,read.max_row+1):
 read.row_dimensions[r].height=43
 read.cell(r,2).alignment=Alignment(wrap_text=True,vertical='top')
d.conditional_formatting.add('B2:D2',CellIsRule(operator='lessThan',formula=['0'],fill=PatternFill('solid',fgColor='FCE4D6'),font=Font(color='9C0006')))
d.conditional_formatting.add('B10:D10',CellIsRule(operator='lessThan',formula=['1'],fill=PatternFill('solid',fgColor='FCE4D6')))
w['Sensitivity'].conditional_formatting.add('B2:D4',ColorScaleRule(start_type='min',start_color='F8696B',end_type='max',end_color='63BE7B'))
chart=BarChart();chart.title='20-year project NPV (EUR)';chart.y_axis.title='EUR';chart.add_data(Reference(d,min_col=2,max_col=4,min_row=2,max_row=2),from_rows=True);chart.set_categories(Reference(d,min_col=2,max_col=4,min_row=1,max_row=1));chart.height=8;chart.width=18;d.add_chart(chart,'A23')
# Replace permissive generic checks with domain-specific validation.
ws=w['Inputs'];ws.data_validations.dataValidation=[]
for row in range(2,28):
 if row in [2,8,24]:kind='whole';lo=1;hi=19 if row==8 else 50 if row==24 else 10000
 elif row in [13,14,16,20,22,27]:kind='decimal';lo=0;hi=1
 elif row in [19,21]:kind='decimal';lo=.0001;hi=1
 elif row in [3,5,9,10,25]:kind='decimal';lo=.01;hi=100000000
 else:kind='decimal';lo=0;hi=100000000
 dv=DataValidation(type=kind,operator='between',formula1=lo,formula2=hi);dv.errorTitle='Assumption outside supported range';dv.error='Enter a value within the stated model range.';dv.showErrorMessage=True;dv.errorStyle='stop';ws.add_data_validation(dv);dv.add(f'B{row}:D{row}')
for sheet in w:
 sheet.sheet_properties.pageSetUpPr.fitToPage=True
 sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A3;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
 sheet.print_options.horizontalCentered=True;sheet.print_title_rows='1:1'
 sheet.sheet_properties.tabColor='007F82' if sheet.title in ['Dashboard','Start'] else '0070C0' if sheet.title=='Inputs' else '7F7F7F'
 sheet.sheet_view.showGridLines=False
w.active=w.sheetnames.index('Dashboard')
tmp=p/'model.polished.tmp.xlsx';w.save(tmp);os.replace(tmp,p/'model.xlsx')
print('Polished workbook saved. SAVE_NOW=',(p/'SAVE_NOW.txt').exists())
