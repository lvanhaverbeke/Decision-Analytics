from pathlib import Path
import os
from openpyxl import load_workbook
p=Path(__file__).resolve().parent
if (p/'SAVE_NOW.txt').exists():print('SAVE_NOW: retained saved workbook');raise SystemExit
w=load_workbook(p/'model.xlsx')
for c,sc in zip('BCD',['Bad','Intermediate','Optimal']):
 w['Dashboard'][f'{c}10']=f"=MIN(INDEX('{sc}'!P2:P22,Inputs!{c}8+2):'{sc}'!P22)"
for r,shift in [(2,-.01),(3,0),(4,.01)]:
 for c,change in zip('BCD',[-.2,0,.2]):
  w['Sensitivity'][f'{c}{r}']=f'=Dashboard!C2-({change})*SUMPRODUCT(Intermediate!B3:B22/(1+Inputs!C21)^Intermediate!A3:A22)+(Intermediate!G22*(Inputs!C19/(Inputs!C19+({shift}))-1))/(1+Inputs!C21)^20'
for row in w['Start'].iter_rows():
 if row[0].value == 'Rebuild files':row[1].value='Run build_model.py, enrich_model.py, polish_model.py, finalize_model.py, then validate_and_cache.py. Atomic saves; validation_results.json records checks.'
# Persist final repair for repeatable rebuilds.
for c in 'BCD':w['Dashboard'][f'{c}10'].number_format='0.00x'
tmp=p/'model.final.tmp.xlsx';w.save(tmp);os.replace(tmp,p/'model.xlsx')
print('Dynamic input references saved; SAVE_NOW=',(p/'SAVE_NOW.txt').exists())
