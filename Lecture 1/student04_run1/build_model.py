from pathlib import Path
import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
ROOT=Path(__file__).resolve().parent
wb=Workbook(); wb.remove(wb.active)
ws=wb.create_sheet('Read me')
for row in [ ['VUB / ULB student housing — decision model'],['Reference date','1 October 2026 (year assumed from assignment context)'],['Status','Screening model; hypothetical acquisition, no selected property'],['Instructions','Edit blue cells on Inputs. All amounts EUR nominal; unlevered project cash flows.'],['Decision','Proceed to due diligence only; do not commit to acquisition until permits, costs and rental demand are verified.'],['Scope','24 rooms in an existing building near VUB/ULB; acquisition and renovation.'],['Boundary','Rent excludes recoverable utilities; tenant charges are pass-through and excluded from owner revenue.'],['Important','Formulas recalculate in Excel/LibreOffice. No property valuation or investment guarantee.'],['Outputs','Annual cash flow, NPV, IRR, simple payback, affordability and scenario comparison.'] ]: ws.append(row)
s=wb.create_sheet('Inputs'); s.append(['Input','Value','Unit / interpretation','Evidence status'])
data=[('Rooms',24,'rooms','Design assumption'),('Purchase price',1800000,'EUR','Placeholder; obtain comparable transactions'),('Acquisition tax',0.125,'share of purchase price','Brussels standard rate; verify applicability'),('Other acquisition costs',35000,'EUR excluding acquisition tax','Assumption; notary quote needed'),('Renovation per room',30000,'EUR/room incl. nonrecoverable VAT','Assumption; contractor quote needed'),('Furniture per room',3500,'EUR/room incl. VAT','Assumption'),('Contingency',0.15,'share of renovation + furniture','Assumption'),('Initial rent',650,'EUR/room/month excluding utilities','Market assumption'),('Stabilised occupancy',0.95,'share of available room months','Assumption'),('Annual rent growth',0.02,'annual nominal escalation','Scenario assumption; indexation legally conditional'),('Fixed operating cost',25000,'EUR/year','Insurance, maintenance, admin, property tax allowance'),('Variable operating cost',0.2,'share of rent revenue','Management, routine repairs, reserve'),('Operating cost growth',0.025,'annual nominal escalation','Assumption'),('Discount rate',0.065,'nominal unlevered required return','Assumption'),('Exit cap rate',0.055,'terminal NOI yield','Assumption'),('Selling costs',0.03,'share of terminal value','Assumption'),('Year 1 occupancy',0.75,'ramp-up after renovation at time zero','Assumption'),('Annual utilities charge',1200,'EUR/room/year paid by tenant','Assumption; metering needed'),('Student monthly budget',1300,'EUR/month available for living','Illustrative; not a statutory living wage'),('Annual replacement capex',10000,'EUR/year from year 1','Separate from routine repairs reserve'),('Replacement capex growth',0.025,'annual escalation','Assumption')]
for r in data:s.append(r)
refs={r[0]:f"'Inputs'!$B${i+2}" for i,r in enumerate(data)}
def v(n):return refs[n]
s=wb.create_sheet('Investment');s.append(['Item','EUR']);
items=[('Purchase',f'={v("Purchase price")}'),('Acquisition tax',f'=B2*{v("Acquisition tax")}'),('Other acquisition costs',f'={v("Other acquisition costs")}'),('Renovation',f'={v("Rooms")}*{v("Renovation per room")}'),('Furniture',f'={v("Rooms")}*{v("Furniture per room")}'),('Contingency',f'=(B5+B6)*{v("Contingency")}'),('Total initial investment','=SUM(B2:B7)')]
for r in items:s.append(r)
s=wb.create_sheet('Cash flow');s.append(['Year','Rent / room / month','Occupancy','Rent revenue','Fixed costs','Variable costs','NOI','Replacement capex','Terminal proceeds','Project cash flow','Discounted cash flow','Cumulative operating cash flow'])
s.append([0,0,0,0,0,0,0,0,0,"=-'Investment'!B8",'=J2','=J2'])
for year in range(1,21):
 r=year+2
 s.append([year,f'={v("Initial rent")}*(1+{v("Annual rent growth")})^(A{r}-1)',f'={v("Year 1 occupancy")}' if year==1 else f'={v("Stabilised occupancy")}',f'={v("Rooms")}*12*B{r}*C{r}',f'={v("Fixed operating cost")}*(1+{v("Operating cost growth")})^(A{r}-1)',f'=D{r}*{v("Variable operating cost")}',f'=D{r}-E{r}-F{r}',f'={v("Annual replacement capex")}*(1+{v("Replacement capex growth")})^(A{r}-1)',f'=((D{r}*(1+{v("Annual rent growth")}))-(E{r}*(1+{v("Operating cost growth")}))-(F{r}*(1+{v("Annual rent growth")})))/{v("Exit cap rate")}*(1-{v("Selling costs")})' if year==20 else 0,f'=G{r}-H{r}+I{r}',f'=J{r}/(1+{v("Discount rate")})^A{r}',f'=L{r-1}+G{r}-H{r}'])
s=wb.create_sheet('Dashboard');s.append(['Metric','Value','Interpretation']);
for r in [('Initial investment',"='Investment'!B8",'All-in cash outlay at year 0'),('20-year NPV',"=SUM('Cash flow'!K2:K22)",'Includes resale; positive required for financial acceptance'),('Project IRR',"=IRR('Cash flow'!J2:J22)",'Unlevered, before income/corporate tax'),('Year 2 net operating income',"='Cash flow'!G4",'Before replacement capex and financing'),('Year 2 net cash yield',"=('Cash flow'!G4-'Cash flow'!H4)/B2",'Annual operating cash / initial investment'),('Simple operating payback', '=IF(\'Cash flow\'!L22<0,"Not reached in 20 years",COUNTIF(\'Cash flow\'!L3:L22,"<0")+1)','Whole-year approximation; excludes resale'),('Monthly tenant housing cost',f'={v("Initial rent")}+{v("Annual utilities charge")}/12','Rent plus utilities'),('Housing share of student budget',f'=B8/{v("Student monthly budget")}','Illustrative budget; compare affordability'),('Terminal PV share of benefits',"='Cash flow'!I22/(1+'Inputs'!B15)^20/(SUM('Cash flow'!K3:K22))",'Sensitivity to resale assumptions')]:s.append(r)
s=wb.create_sheet('Executive report');
for r in [['Executive report — screening recommendation'],['Proposal','Acquire and renovate an existing 24-room building within walking/cycling distance of VUB/ULB and a frequent tram, metro or rail connection.'],['Recommendation','Conditional due diligence only. Financial feasibility must be demonstrated after quotes and site checks.'],['Financial approach','20-year nominal unlevered DCF with full acquisition costs, renovation, contingency, annual replacement capex and terminal sale.'],['Environmental','Reuse the existing structure; verify PEB, insulation, heat pumps, ventilation, flood risk, bicycle provision and metered utilities.'],['Economic','Test rent against furnished student room comparables; tenant utilities shown separately. Contractor and furnishing costs are placeholders.'],['Social','Balance student affordability, accessible rooms, quiet hours, neighbour engagement and safe routes to campus.'],['Legal','Confirm planning use, number of units, fire approval, room habitability, lease conditions and municipal requirements before purchase.'],['Limitations','No identified property; no debt or investor income tax. Cost and rent assumptions require verification.']]:s.append(r)
s=wb.create_sheet('Sources');s.append(['ID','Topic','URL / source','As-of / status','Use in model']);s.append(['A0','Assumptions','Team screening assumptions','2026-10-01 assumed reference date','Blue input cells require verification'])
s=wb.create_sheet('Due diligence');s.append(['Area','Required evidence','Gate / action']);
for r in [('Site','Address, floor plans, cadastral records, transit walking time','Select property and verify route on operator timetable'),('Planning','Authorised use and subdivision; municipal planning advice','Written confirmation before binding purchase'),('Fire','SIAMU advice, evacuation routes, alarms and compartmentation','Approval and contractor design'),('Habitation','Surface, daylight, ventilation, sanitation and safety','Assess each room against Brussels Housing Code'),('Energy','PEB certificate, audit, bills, renovation estimates','Costed retrofit and energy savings'),('Financial','Three contractor quotes, notary estimate, market comparables','Replace placeholders and rerun scenarios'),('Social','Student budgets and disability access review','Set affordable room mix and support procedures')]:s.append(r)
wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
def save():
 for sh in wb:
  sh.freeze_panes='B2';sh.auto_filter.ref=sh.dimensions
  for cell in sh[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='183B56')
  for row in sh.iter_rows(min_row=2):
   for c in row:
    c.alignment=Alignment(vertical='top',wrap_text=True)
    if c.data_type=='f':c.number_format='#,##0.00;[Red](#,##0.00)'
  for col in range(1,sh.max_column+1):sh.column_dimensions[get_column_letter(col)].width=24 if sh.title=='Cash flow' else (34 if col==1 else 65)
  if sh.title!='Cash flow':
   for rr in range(2,sh.max_row+1):sh.row_dimensions[rr].height=42
 for r in range(2,23):
  wb['Inputs'].cell(r,2).font=Font(color='0000FF');wb['Inputs'].cell(r,2).number_format='0.00%' if isinstance(wb['Inputs'].cell(r,2).value,float) and wb['Inputs'].cell(r,2).value<1 else '#,##0.00'
 for r in [4,6,9,10]:wb['Dashboard'].cell(r,2).number_format='0.0%'
 tmp=ROOT/'model.tmp.xlsx';wb.save(tmp);load_workbook(tmp);os.replace(tmp,ROOT/'model.xlsx')
save()
print('Saved basic model.xlsx; SAVE_NOW:',(ROOT/'SAVE_NOW.txt').exists())
