from pathlib import Path
import re, math, os, zipfile, xml.etree.ElementTree as ET
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.chart import LineChart,Reference
from openpyxl.worksheet.datavalidation import DataValidation
ROOT=Path(__file__).resolve().parent
wb=load_workbook(ROOT/'model.xlsx');cache={};busy=set()
if not (ROOT/'SAVE_NOW.txt').exists():
 wb['Executive report'].delete_rows(10,100)
 wb['Dashboard'].delete_rows(11,100)
 wb['Dashboard']._charts=[]
def irr(vals):
 lo,hi=-.95,1.
 def npv(x):return sum(v/(1+x)**i for i,v in enumerate(vals))
 while npv(hi)>0:hi*=2
 for _ in range(150):
  mid=(lo+hi)/2
  if npv(mid)>0:lo=mid
  else:hi=mid
 return (lo+hi)/2
def flat(x):return x if isinstance(x,list) else [x]
def SUM(*args):return sum(v for a in args for v in flat(a) if isinstance(v,(int,float)))
def COUNTIF(a,c):return sum(v<0 for v in a) if c=='<0' else sum(v==c for v in a)
def SUMPRODUCT(a,b):return sum(x*y for x,y in zip(a,b))
def val(sheet,addr):
 key=(sheet,addr)
 if key in cache:return cache[key]
 x=wb[sheet][addr].value
 if not isinstance(x,str) or not x.startswith('='):return x if x is not None else 0
 assert key not in busy,key
 busy.add(key);expr=x[1:];tokens={}
 def hold(value):
  name=f'placeholder_{len(tokens)}';tokens[name]=value;return name
 expr=re.sub(r'"([^"\n]*)"',lambda m:hold(m.group(1)),expr)
 pat=r"(?:(?:'([^']+)'|([A-Za-z_][A-Za-z_ ]*))!)?(\$?[A-Z]{1,3}\$?\d+)(?::(\$?[A-Z]{1,3}\$?\d+))?"
 def ref(m):
  sh=m.group(1) or m.group(2) or sheet;a=m.group(3).replace('$','');b=m.group(4)
  if b:
   cells=wb[sh][a+':'+b.replace('$','')];return hold([val(sh,c.coordinate) for row in cells for c in row])
  return hold(val(sh,a))
 expr=re.sub(pat,ref,expr);expr=expr.replace('^','**')
 env={'MAX':max,'SUM':SUM,'IRR':irr,'IF':lambda c,a,b:a if c else b,'COUNTIF':COUNTIF,'SUMPRODUCT':SUMPRODUCT,**tokens}
 result=eval(expr,{'__builtins__':{}},env);busy.remove(key);cache[key]=result;return result
for sh in wb:
 for row in sh:
  for c in row:
   if c.data_type=='f':val(sh.title,c.coordinate)
metrics={wb['Dashboard'].cell(r,1).value:val('Dashboard',f'B{r}') for r in range(2,11)}
print(metrics)
if not (ROOT/'SAVE_NOW.txt').exists():
 s=wb['Executive report'];s.append(['Base-case results',f"Investment EUR{metrics['Initial investment']:,.0f}; NPV EUR{metrics['20-year NPV']:,.0f}; IRR {metrics['Project IRR']:.2%}; operating payback {metrics['Simple operating payback']}; housing EUR{metrics['Monthly tenant housing cost']:,.0f}/month."])
 s.append(['Financial decision','Reject acquisition at the assumed price: negative NPV. Renegotiate price, reduce construction costs or seek a transparent subsidy; do not rely on raising rent above affordability limits.'])
 s.append(['Scenario outcomes', '; '.join(f"{wb['Scenarios'].cell(r,1).value}: NPV EUR{val('Scenarios',f'H{r}'):,.0f}, IRR {val('Scenarios',f'I{r}'):.2%}" for r in range(2,5))])
 s.append(['Affordability','EUR650 total is 50% of the illustrative EUR1300 budget. This is a concern despite being within the broad ULB room-cost range. Subsidised university rooms are cheaper.'])
 s.append(['Preferred search area','Prioritise the Etterbeek station / Petillon / La Plaine corridor for VUB/ULB access. Compare Ixelles cemetery and Solbosch alternatives using address-specific quotes and walking routes. Matrix scores are illustrative.'])
 s.append(['Investment gates','Positive NPV at the chosen return, validated demand at an affordable rent, confirmed planning/fire/housing compliance, and a funded PEB retrofit. All legal gates remain unverified.'])
 s.append(['Presentation sequence','1. Need and location; 2. Assumptions and sources; 3. Costs and cash flow; 4. Scenarios; 5. Environmental/social tradeoffs; 6. Legal gates and conditional decision.'])
 s.append(['Next 4–8 weeks','Shortlist addresses; obtain planning records and fire advice; commission energy/building surveys; get contractor/notary quotes; collect comparable rents and student budgets. Do not contact sellers or authorities on behalf of the team without separate instruction.'])
 # Break-even price holds costs and cash flows fixed; registration scales with purchase.
 d=wb['Dashboard'];d.append(['Break-even purchase price',"=MAX(0,('Inputs'!B3*(1+'Inputs'!B4)+B3)/(1+'Inputs'!B4))",'NPV=0 holding rent, costs and exit yield fixed; negotiation screen'])
 # MAX support for new formula
 price=max(0,(wb['Inputs']['B3'].value*1.125+metrics['20-year NPV'])/1.125);cache[('Dashboard','B11')]=price
 s.append(['Negotiation screen',f'Break-even purchase price approximately EUR{price:,.0f}, holding all other base assumptions fixed.'])
 chart=LineChart();chart.title='Annual project cash flow (includes year 20 sale)';chart.y_axis.title='EUR';chart.x_axis.title='Year';chart.add_data(Reference(wb['Cash flow'],min_col=10,min_row=1,max_row=22),titles_from_data=True);chart.set_categories(Reference(wb['Cash flow'],min_col=1,min_row=2,max_row=22));d.add_chart(chart,'E2')
 report='''# Student housing near VUB/ULB — executive report\n\nReference date: 1 October 2026 (year assumed). Research checked 14 October 2026; live sources do not prove their historical content on 1 October.\n\nThe proposal is to buy and renovate an existing building for 24 furnished rooms near the VUB/ULB La Plaine corridor. This is a screening design, with no selected property or verified purchase price.\n\n'''
 report+=f"At EUR{metrics['Initial investment']:,.0f} initial investment and EUR550 monthly rent plus EUR100 utilities, the 20-year model produces NPV EUR{metrics['20-year NPV']:,.0f} at 6.5% and an IRR of {metrics['Project IRR']:.2%}. Operating payback: {metrics['Simple operating payback']}. The assumed acquisition is financially unattractive. The break-even purchase price is approximately EUR{price:,.0f}, with all other assumptions fixed.\n\n"
 report+='The model includes purchase duty, acquisition costs, renovation, furniture, contingency, operating costs, replacement capex and resale. It excludes borrowing and investor income tax. Occupancy measures paid room-months: 95% requires summer demand. The downside uses ten paid months, lower rent and a construction overrun.\n\n'
 report+='EUR650 monthly housing consumes 50% of the illustrative EUR1300 student budget. Affordability needs direct student evidence and a subsidised room mix. ULB guidance gives a broad EUR500–680 range including charges; its Lucia De Brouckere residence charges EUR491 including EUR61 charges for 2026–27. These are distinct market segments.\n\n'
 report+='Reuse the building, insulate and ventilate well, meter utilities and provide bicycle storage. Illustrative energy consumption falls from 150000 to 72000 kWh/year; this is an engineering assumption, not a PEB certificate. Savings accrue to tenants and are excluded from owner revenue. Survey the structure, flood exposure and heating system before setting retrofit costs.\n\n'
 report+='Planning use, permitted room count, the housing standards applicable from January 2026, fire approval, PEB and compliant leases are mandatory gates. Municipality-specific rules depend on the exact address. No gate is presently verified. Seek conditional due diligence, renegotiate the price and rerun the workbook before making an acquisition decision.\n\nSources and limitations are recorded in the workbook. Selected primary evidence:\n\n'
 for sid,topic,url,*_ in [tuple(wb['Sources'].cell(r,c).value for c in range(1,6)) for r in range(2,wb['Sources'].max_row+1)]:report+=f'- [{sid}: {topic}]({url})\n'
 (ROOT/'executive_report.md').write_text(report)
for sh in wb:
 sh.freeze_panes='B2';sh.auto_filter.ref=sh.dimensions
 for c in sh[1]:c.fill=PatternFill('solid',fgColor='183B56');c.font=Font(bold=True,color='FFFFFF')
 for row in sh.iter_rows(min_row=2):
  for c in row:c.alignment=Alignment(wrap_text=True,vertical='top');c.number_format='#,##0.00;[Red](#,##0.00)' if c.data_type=='f' else c.number_format
  if sh.title not in ['Cash flow','Scenario cash']:sh.row_dimensions[row[0].row].height=55
 for col in sh.columns:sh.column_dimensions[col[0].column_letter].width=25 if sh.title in ['Cash flow','Scenario cash'] else (34 if col[0].column==1 else 48)
 if sh.title in ['Scenarios','Environment','Decision matrix']:
  for row in sh.iter_rows(min_row=2):
   for c in row:
    if isinstance(c.value,(int,float)):c.font=Font(color='0000FF')
for r in [4,6,9,10]:wb['Dashboard'].cell(r,2).number_format='0.0%'
for r in range(2,5):
 for c in [3,5,6,9]:wb['Scenarios'].cell(r,c).number_format='0.0%'
for r in [4,7,8]:wb['Social and services'].cell(r,2).number_format='0.0%'
# Save and inject formula caches so initial values are visible without an Excel engine.
tmp=ROOT/'model.tmp.xlsx';wb.save(tmp)
ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with zipfile.ZipFile(tmp) as zin:
 members={name:zin.read(name) for name in zin.namelist()}
for i,sh in enumerate(wb,1):
 name=f'xl/worksheets/sheet{i}.xml';root=ET.fromstring(members[name])
 for cell in root.findall('.//m:c',ns):
  if cell.find('m:f',ns) is None:continue
  k=(sh.title,cell.attrib['r'])
  if k not in cache:continue
  value=cache[k];node=cell.find('m:v',ns)
  if node is None:node=ET.SubElement(cell,'{'+ns['m']+'}v')
  if isinstance(value,str):cell.set('t','str');node.text=value
  else:cell.attrib.pop('t',None);node.text=str(value)
 members[name]=ET.tostring(root,encoding='utf-8',xml_declaration=True)
with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as zout:
 for name,content in members.items():zout.writestr(name,content)
check=load_workbook(tmp,data_only=True)
assert abs(check['Dashboard']['B3'].value-metrics['20-year NPV'])<.01
assert all(math.isfinite(v) for v in cache.values() if isinstance(v,(int,float)))
os.replace(tmp,ROOT/'model.xlsx')
(ROOT/'validation.txt').write_text(f'Validated {len(cache)} formula results; finite values; cached NPV matches.\n'+str(metrics)+'\n')
print('Final model saved atomically; formula caches validated. SAVE_NOW:',(ROOT/'SAVE_NOW.txt').exists())
