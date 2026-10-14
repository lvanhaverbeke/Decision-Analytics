from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.properties import CalcProperties
import os
ROOT=Path(__file__).resolve().parent
wb=Workbook(); wb.remove(wb.active)
for name in ['Read me','Inputs','Room mix','Development','Cash flow','Dashboard','Sources']:
 wb.create_sheet(name)
wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
def row(s,n,values):
 for c,v in enumerate(values,1): wb[s].cell(n,c,v)
row('Read me',1,['VUB student housing | investment feasibility'])
row('Read me',3,['Working baseline saved before research. All initial numbers are illustrative assumptions.'])
row('Read me',4,['Location: assumed site within walking/cycling distance of VUB Etterbeek campus; no land secured.'])
row('Read me',5,['No classroom workbook supplied in the working folder; model built from first principles.'])
row('Read me',6,['Shared studio means two students per studio. Rents exclude separately billed actual utilities.'])
row('Read me',7,['Blue cells = editable inputs; black cells = formulas. EUR nominal, annual periods.'])
row('Inputs',1,['Assumption','Value','Basis / unit'])
inputs=[('Normal rooms',60,'units'),('Shared studios',20,'units; 2 beds each'),('Single studios',20,'units'),('Normal monthly rent',480,'EUR per room; provisional'),('Shared monthly rent',390,'EUR per person; provisional'),('Single studio monthly rent',650,'EUR per unit; provisional'),('Letting months',12,'months/year'),('Occupancy',.95,'stabilised'),('Rent growth',.02,'annual'),('Operating expense ratio',.28,'share of effective rent'),('Reserve per bed',150,'EUR/year'),('Hard cost per gross sqm',2200,'EUR; assumption'),('Land cost',1800000,'EUR; assumption'),('Soft cost ratio',.15,'share of hard costs'),('Contingency',.10,'share hard + soft'),('VAT / nonrecoverable tax',.21,'conservative assumption on hard + soft + contingency'),('Debt share',.55,'share of total development cost'),('Interest rate',.045,'annual'),('Loan term',25,'years'),('Discount rate',.08,'unlevered'),('Exit cap rate',.055,'NOI / terminal value'),('Selling cost',.02,'share exit value'),('Construction years',2,'fixed 2-year baseline'),('Normal room net area',18,'sqm'),('Shared studio net area',32,'sqm per unit'),('Single studio net area',25,'sqm'),('Grossing factor',1.4,'gross / net'),('Opening occupancy',.85,'first operating year'),('Operating years',20,'fixed baseline'),('Equity discount rate',.10,'annual')]
for n,v in enumerate(inputs,2): row('Inputs',n,v)
row('Room mix',1,['Type','Units','Beds per unit','Beds','Net sqm/unit','Gross sqm','Rent/person/month','Annual gross rent'])
for n,ir,beds,area,rent in [(2,2,1,25,5),(3,3,2,26,6),(4,4,1,27,7)]:
 row('Room mix',n,[['Normal dorm','Shared studio','Single studio'][n-2],f'=Inputs!B{ir}',beds,f'=B{n}*C{n}',f'=Inputs!B{area}',f'=B{n}*E{n}*Inputs!B28',f'=Inputs!B{rent}',f'=D{n}*G{n}*Inputs!B8'])
row('Room mix',5,['Total','=SUM(B2:B4)',None,'=SUM(D2:D4)',None,'=SUM(F2:F4)',None,'=SUM(H2:H4)'])
row('Development',1,['Development budget','EUR'])
for n,v in {2:['Land','=Inputs!B14'],3:['Hard costs',"='Room mix'!F5*Inputs!B13"],4:['Soft costs','=B3*Inputs!B15'],5:['Contingency','=SUM(B3:B4)*Inputs!B16'],6:['Nonrecoverable VAT','=SUM(B3:B5)*Inputs!B17'],7:['Total development cost','=SUM(B2:B6)'],8:['Debt','=B7*Inputs!B18'],9:['Equity','=B7-B8'],10:['Annual debt service','=-PMT(Inputs!B19,Inputs!B20,B8)']}.items():row('Development',n,v)
row('Cash flow',1,['Year','Gross rent','Occupancy','Effective rent','Operating cost','Reserve','NOI','Development spend','Debt draw','Debt service','Exit proceeds','Debt outstanding','Unlevered CF','Equity CF'])
for n in range(2,24):
 y=n-2; op=y-1
 vals=[y,0,0,0,0,0,0,'=Development!B7*0.5' if y<2 else 0,'=Development!B8*0.5' if y<2 else 0,0,0,'=SUM(I$2:I%d)'%n, f'=G{n}-H{n}+K{n}',f'=M{n}+I{n}-J{n}']
 if y>=2:
  vals[1]=f"='Room mix'!H5*(1+Inputs!B10)^({op}-1)"
  vals[2]='=Inputs!B29' if y==2 else '=Inputs!B9'
  vals[3]=f'=B{n}*C{n}';vals[4]=f'=D{n}*Inputs!B11';vals[5]=f"='Room mix'!D5*Inputs!B12*(1+Inputs!B10)^({op}-1)";vals[6]=f'=D{n}-E{n}-F{n}'
  vals[9]='=Development!B10';vals[11]=f'=MAX(0,L{n-1}*(1+Inputs!B19)-J{n})'
 if y==21:
  vals[10]=f'=G{n}*(1+Inputs!B10)/Inputs!B22*(1-Inputs!B23)';vals[13]=f'=M{n}+I{n}-J{n}-L{n}'
 row('Cash flow',n,vals)
row('Dashboard',1,['Decision dashboard','Value'])
for n,v in {3:['Total beds',"='Room mix'!D5"],4:['Total units',"='Room mix'!B5"],5:['Total gross area',"='Room mix'!F5"],6:['Total development cost','=Development!B7'],7:['Stabilised NOI',"='Cash flow'!G5"],8:['Yield on cost','=B7/B6'],9:['Unlevered NPV',"=NPV(Inputs!B21,'Cash flow'!M3:M23)+'Cash flow'!M2"],10:['Unlevered IRR',"=IFERROR(IRR('Cash flow'!M2:M23),0)"],11:['Equity NPV',"=NPV(Inputs!B31,'Cash flow'!N3:N23)+'Cash flow'!N2"],12:['Equity IRR',"=IFERROR(IRR('Cash flow'!N2:N23),0)"],13:['Stabilised DSCR','=B7/Development!B10'],14:['Preliminary decision','=IF(B9>0,"Financially viable under assumptions","Redesign / lower land and construction cost")']}.items(): row('Dashboard',n,v)
row('Sources',1,['Topic','Source / URL','Interpretation','Status'])
row('Sources',2,['Initial version','User brief','Research pending; inputs are assumptions','Baseline'])
def save():
 for ws in wb:
  ws.freeze_panes='B2';ws.sheet_view.showGridLines=False
  for c in ws[1]: c.fill=PatternFill('solid',fgColor='17365D');c.font=Font(color='FFFFFF',bold=True)
  ws.column_dimensions['A'].width=38
  for col in 'BCDEFGHIJKLMN':ws.column_dimensions[col].width=20
  if ws.title in ['Read me','Sources']:ws.column_dimensions['B'].width=65;ws.column_dimensions['C'].width=80
  for rows in ws.iter_rows(min_row=2):
   for c in rows:
    if c.data_type=='f':c.font=Font(color='172D42');c.number_format='#,##0.00;[Red](#,##0.00);–'
    elif isinstance(c.value,(int,float)):c.font=Font(color='0070C0');c.number_format='#,##0.00'
    c.alignment=Alignment(vertical='top',wrap_text=True)
  ws.auto_filter.ref=ws.dimensions
 for r in [9,10,11,15,16,17,18,19,21,22,23,29,31]: wb['Inputs'].cell(r,2).number_format='0.0%'
 for r in [8,10,12]:wb['Dashboard'].cell(r,2).number_format='0.0%'
 tmp=ROOT/'model.tmp.xlsx';wb.save(tmp);load_workbook(tmp).close();os.replace(tmp,ROOT/'model.xlsx')
save()
print('Saved model.xlsx; SAVE_NOW:',(ROOT/'SAVE_NOW.txt').exists())
if (ROOT/'SAVE_NOW.txt').exists(): raise SystemExit('SAVE_NOW: baseline saved; stopped')
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import ColorScaleRule
for name in ['Demand & competition','Rent rules','Scenarios','Sensitivity','Checks']:wb.create_sheet(name)
row('Read me',3,['Research version | 14 October 2026 | educational feasibility, pre-tax cash flows'])
notes={8:'Recommendation: test 120 beds in 100 units (60 normal, 20 two-person studios, 20 single studios). Demand case is plausible; financial acceptance depends on dashboard.',9:'120 beds limits exposure to the 465-room Usquare pipeline. No verified VUB waiting-list count or secured plot: capacity remains a testable proposal.',10:'Model timeline: year 0 = 2027, year 1 = 2028 construction, year 2 = 2029 opening; 20 operating years; sale at end of year 21.',11:'Cash flows are year-end; debt draws are half in each construction year. Construction interest uses midyear draws; first debt repayment at opening.',12:'Nonrecoverable VAT 21%, land transaction allowance 15%, construction price, land price and financing are assumptions requiring professional quotations.',13:'Utilities are reimbursed at actual cost, outside revenue and expenses. Operating ratio includes management, insurance, owner property tax, maintenance and administration.',14:'Rent growth of 2% is a financial scenario, not permission to index: check lease registration, EPC rules, health index and reference rent annually.',15:'No corporate income tax, depreciation tax shield, grant or subsidy. No plot-specific planning, fire safety, accessibility or shared-occupancy approval obtained.',16:'Fixed 2 construction years / 20 operating years. Change the Python generator to change these periods; all other blue financial inputs can be edited in Excel.',17:'Official reference rents are missing: proposed rents are planning inputs. Workbook blocks legal clearance until references are supplied. Do not interpret affordability caps as government tariffs.',18:'Shared studio is let as one unit to two occupants; divide the unit rent by two for per-person affordability. Validate compatibility with actual tenant arrangements.'}
for n,s in notes.items():row('Read me',n,[s])
wb['Inputs']['B24']='2 years fixed';wb['Inputs']['B30']='20 years fixed'
row('Inputs',32,['Land acquisition allowance',.15,'assumption; land price includes no transaction fees'])
row('Inputs',33,['Furniture per bed',1800,'EUR; initial furnishing'])
row('Inputs',34,['Initial working capital',50000,'EUR; conservatively unrecovered at exit'])
row('Inputs',35,['Reference rent policy factor',1,'stay at/below reference when supplied; max 1.2 screening threshold'])
row('Inputs',36,['Candidate beds',120,'planning target; mix on rows 2:4 drives actual capacity'])
row('Rent rules',1,['Room type','Desired rent/unit/month','Official reference/unit/month','Planning affordability cap/unit','Applied unit rent','Applied per person','Verification'])
for n,ir,b in [(2,5,1),(3,6,2),(4,7,1)]:
 row('Rent rules',n,[['Normal dorm','Shared studio','Single studio'][n-2],f'=Inputs!B{ir}*{b}',0,[480,780,650][n-2],f'=MIN(B{n},D{n},IF(C{n}>0,C{n}*Inputs!B35,D{n}))',f'=E{n}/{b}',f'=IF(C{n}>0,"Reference supplied - review quality and lease","NOT VERIFIED")'])
 wb['Room mix'].cell(n,7,f"='Rent rules'!F{n}")
row('Rent rules',6,['Rule','Reference rent depends on address, area, condition and amenities. No universal Brussels student-room tariff.'])
row('Rent rules',7,['20% screen','Above reference +20% is presumed excessive, subject to specific justifications. Lower rents may also be excessive due to quality defects.'])
row('Rent rules',8,['Model policy','At/below supplied reference by default. Enter official calculator output per rented unit; a zero means missing, not zero allowable rent.'])
row('Rent rules',9,['Lease review','Reference must appear in lease; verify student lease, indexation and separate utility settlement before offering units.'])
row('Rent rules',10,['Official calculator','https://loyers.brussels/'])
row('Rent rules',11,['Source','https://be.brussels/fr/logement/location/bail-dhabitation/loyers-de-reference'])
# Financing: debt funds base development, construction interest is funded by equity.
row('Development',11,['Land transaction allowance','=B2*Inputs!B32'])
row('Development',12,['Furniture',"='Room mix'!D5*Inputs!B33"])
row('Development',13,['Working capital','=Inputs!B34'])
row('Development',14,['Construction interest (equity-funded)','=B8*Inputs!B19'])
row('Development',15,['Unlevered project cost','=SUM(B2:B6,B11:B13)'])
row('Development',7,['Total funding need','=B15+B14'])
row('Development',8,['Debt','=B15*Inputs!B18'])
row('Development',16,['Cost per bed',"=B7/'Room mix'!D5"])
for n in [2,3]:
 wb['Cash flow'].cell(n,8,'=Development!B15*0.5')
 wb['Cash flow'].cell(n,14,f'=M{n}+I{n}-J{n}-Development!B14*0.5')
row('Cash flow',25,['Development costs exclude financing interest; equity cash flow includes it. NOI is after a replacement reserve.'])
row('Cash flow',26,['First operating NOI is G4. G5 is the second operating year (stabilised occupancy). Exit uses next-year NOI.'])
row('Demand & competition',1,['Evidence / planning assumption','Value','Interpretation / source ID'])
demand=[('Brussels students, 2022-23',137000,'Lower-bound rounded figure; S1'),('Students using accommodation, 2022-23',37000,'At least; regional demand indicator, not new vacant demand; S1'),('Dedicated student units, 2025',11305,'Units are not necessarily beds; cross-period comparison; S1'),('Indicative non-dedicated demand','=B3-B4','Proxy only; many already housed in ordinary rentals'),('Usquare proposed rooms',465,'Permit submitted 25 Aug 2026; S2'),('Pipeline overlap factor',1,'Conservative assumption: all compete with project'),('Residual proxy','=MAX(0,B5-B6*B7)','Not a verified local waiting list'),('Local addressable capture assumption',.006,'0.6% of proxy; replace with local survey'),('Indicative addressable beds','=B8*B9','Screening estimate only'),('Proposed beds',"='Room mix'!D5",'100 units / 120 beds baseline'),('Beds versus addressable estimate','=IF(B11<=B10,"Within planning envelope","Capacity exceeds assumed envelope")','Demand gate; validate pre-letting'),('Expected stabilised occupied beds','=B11*Inputs!B9','95% assumption, not observed market occupancy'),('Usquare opening target',2029,'S3 provisional schedule; older 2026 announcement superseded')]
for n,v in enumerate(demand,2):row('Demand & competition',n,v)
row('Demand & competition',17,['Competitor / status','Beds or units','Rent/month','Charges','Source / implication'])
for n,v in enumerate([['Usquare / proposed',465,'Not published','Unknown','S2/S3: affordable pipeline, 2029 target; monitor permit and delivery'],['Bru Upkot / operating',216,589,119,'S4: entry room; 0.75km to VUB; existing stock, do not subtract again'],['Bru Upkot / 1p studio',None,769,134,'S4: 2026-27 asking rent'],['Bru Upkot / 2p studio',None,859,164,'S4: unit asking rent; per person half if 2 tenants'],['VUB / operating network',1540,449.3,'Included','S5/S6: across Etterbeek + Jette; not all local'],['VUB / restricted single studio',12,561.6,'Included','S6: reserved for medical need; not general private market'],['VUB / two-person flats',35,674,'Included','S6: unit price; limited institutional allocation']],18):row('Demand & competition',n,v)
row('Demand & competition',26,['Positioning','Base rent below private Bru Upkot examples; VUB institutional all-inclusive prices are lower and not directly comparable.'])
row('Demand & competition',27,['Sizing rationale','60 normal beds + 40 shared-studio beds + 20 single-studio beds; 83% of beds in lower-price products.'])
row('Demand & competition',28,['Site concept','Approximately 3,108 gross sqm; at 4 floors, ~777 sqm footprint, before site circulation and external space. Assumption, no planning assurance.'])
row('Sources',1,['ID / topic','Primary URL','Evidence / model treatment','Publication / access'])
sources=[('S1 Demand','https://perspective.brussels/fr/actualites/un-etudiant-sur-trois-trouve-un-kot-dans-un-logement-qui-lui-est-destine','137,000+ students; 37,000+ housing users; 11,305 dedicated units. Cross-year regional proxy only.','6 Jan 2026 / 14 Oct 2026'),('S2 Usquare permit','https://msi.brussels/nieuws/usquare-de-msi-dient-vergunningsaanvraag-voor-465-studentenwoningen','465 student rooms in permit application; replaces older 460-room concept.','25 Aug 2026 / 14 Oct 2026'),('S3 Usquare schedule','https://usquare.brussels/en/schedule','Provisional September 2029 commissioning; 30-36 month works. Not a completion guarantee.','Access 14 Oct 2026'),('S4 Bru Upkot','https://www.upkot.be/en/locations/brussels/bru-upkot','216 rooms; 2026-27 prices: entry 589+119; 1p studio 769+134; 2p studio 859+164 EUR/month.','2026-27 / 14 Oct 2026'),('S5 VUB stock','https://www.vub.be/en/studying-vub/practical-info-for-students/student-housing','1,540 rooms across campus network, not Etterbeek alone.','Access 14 Oct 2026'),('S6 VUB rents','https://www.vub.be/en/studying-vub/practical-info-for-students/student-housing/practical-housing-information','2026-27: room 449.30, single studio 561.60, two-person flat 674 EUR/month including listed services; demand exceeds supply.','Updated 18 Sep 2026 / 14 Oct 2026'),('S7 Rent regulation','https://be.brussels/fr/logement/location/bail-dhabitation/loyers-de-reference','Reference must be in lease; >20% above reference presumed excessive with potential justification; quality issues can affect lower rents; obligation from 1 May 2025.','Access 14 Oct 2026'),('S8 Calculator','https://loyers.brussels/','Exact site and property details required. No official reference amount assumed in this model.','Not calculated: missing site'),('A1 Cost / finance','Analyst assumptions, not market quotes','Hard costs 2200 EUR/gross sqm, land 1.8m, 21% nonrecoverable VAT, financing 4.5%, land fees 15%; obtain quotations.','14 Oct 2026')]
for n,v in enumerate(sources,2):row('Sources',n,v)
for n in range(2,10):
 c=wb['Sources'].cell(n,2);c.hyperlink=c.value;c.style='Hyperlink'
row('Dashboard',16,['Rent reference status','=IF(COUNTIF(\'Rent rules\'!C2:C4,">0")=3,"Reference inputs supplied; final legal review required","NOT VERIFIED - enter official references")'])
row('Dashboard',17,['Demand screen',"='Demand & competition'!B12"])
row('Dashboard',18,['Opening calendar year',2029])
row('Dashboard',19,['Funding per bed','=Development!B16'])
row('Dashboard',20,['Break-even occupancy (NOI covers debt)',"=(Development!B10+'Cash flow'!F5)/('Cash flow'!B5*(1-Inputs!B11))"])
row('Dashboard',21,['Operating rent multiplier for NPV=0','See sensitivity and scenarios; increases remain subject to reference-rent rules.'])
row('Dashboard',23,['Investment gate','=IF(AND(B9>0,B13>=1.2,COUNTIF(\'Rent rules\'!C2:C4,">0")=3),"Proceed to detailed due diligence","HOLD: financial and/or rent-reference gate unmet")'])
row('Scenarios',1,['Scenario','Cost multiplier','Stabilised occupancy','Rent multiplier','Exit cap','Project NPV','Project IRR','Stabilised DSCR']+[f'CF year {y}' for y in range(22)])
from openpyxl.utils import get_column_letter as col
for n,v in enumerate([['Downside',1.10,.85,.95,.06],['Base',1,.95,1,.055],['Upside',.90,.98,1.05,.05]],2):
 row('Scenarios',n,v)
 for y in range(22):
  cr=y+2; letter=col(y+9)
  if y<2: formula=f'=-Development!B15*B{n}/2'
  else:
   occ=f'Inputs!B29*C{n}/Inputs!B9' if y==2 else f'C{n}'
   noi=f"('Cash flow'!B{cr}*D{n}*({occ})*(1-Inputs!B11)-'Cash flow'!F{cr})"
   formula='='+noi
   if y==21:formula+=f'+{noi}*(1+Inputs!B10)/E{n}*(1-Inputs!B23)'
  wb['Scenarios'].cell(n,y+9,formula)
 for c,f in [(6,f'=NPV(Inputs!B21,J{n}:AD{n})+I{n}'),(7,f'=IFERROR(IRR(I{n}:AD{n}),0)'),(8,f"=('Cash flow'!B5*D{n}*C{n}*(1-Inputs!B11)-'Cash flow'!F5)/(-PMT(Inputs!B19,Inputs!B20,Development!B8*B{n}))")]:wb['Scenarios'].cell(n,c,f)
row('Scenarios',7,['Scenario occupancy replaces stabilised input; opening occupancy scales proportionally. Debt scales with capex for DSCR. Rent upside is a stress case, not legal permission.'])
row('Sensitivity',1,['Unlevered NPV: cost vs occupancy',.85,.95,.98])
for n,m in enumerate([.9,1,1.1],2):
 wb['Sensitivity'].cell(n,1,m)
 for c in range(2,5):
  oc=f'{col(c)}$1'; terms=['-Development!B15*$A%d/2'%n,'-Development!B15*$A%d/2/(1+Inputs!B21)'%n]
  for y in range(2,22):
   cr=y+2; occ=f'Inputs!B29*{oc}/Inputs!B9' if y==2 else oc
   noi=f"('Cash flow'!B{cr}*{occ}*(1-Inputs!B11)-'Cash flow'!F{cr})"
   term=noi
   if y==21:term=f'({noi}+{noi}*(1+Inputs!B10)/Inputs!B22*(1-Inputs!B23))'
   terms.append(f'{term}/(1+Inputs!B21)^{y}')
  wb['Sensitivity'].cell(n,c,'='+'+'.join(terms))
row('Sensitivity',6,['Rows = development cost multiplier. Columns = stabilised occupancy. Opening occupancy scales with stabilised occupancy. Rents unchanged.'])
row('Checks',1,['Control','Value','Expected'])
checks=[['Beds reconcile',"='Room mix'!D5-(Inputs!B2+Inputs!B3*2+Inputs!B4)",0],['Sources = funding','=Development!B7-Development!B8-Development!B9',0],['Debt fully drawn',"=SUM('Cash flow'!I2:I23)-Development!B8",0],['Unlevered spend matches budget',"=SUM('Cash flow'!H2:H23)-Development!B15",0],['Equity invested',"=-SUM('Cash flow'!N2:N3)-Development!B9",0],['Base scenario NPV reconciles',"=Scenarios!F3-Dashboard!B9",0],['Rent references supplied','=COUNTIF(\'Rent rules\'!C2:C4,">0")',3],['Occupancy bounds','=IF(AND(Inputs!B9>0,Inputs!B9<=1,Inputs!B29>0,Inputs!B29<=1),0,1)',0],['Candidate capacity reconciliation',"='Room mix'!D5-Inputs!B36",0]]
for n,v in enumerate(checks,2):row('Checks',n,v)
for sheet,ran,kind,lo,hi in [('Inputs','B9:B11','decimal',0,1),('Inputs','B2:B4','whole',0,1000),('Rent rules','C2:C4','decimal',0,10000),('Inputs','B35','decimal',0,1.2)]:
 dv=DataValidation(type=kind,operator='between',formula1=lo,formula2=hi);dv.errorTitle='Input outside model bounds';dv.error='Enter a value within the permitted range.';dv.showErrorMessage=True;wb[sheet].add_data_validation(dv);dv.add(ran)
chart=BarChart();chart.title='Project NPV by scenario';chart.add_data(Reference(wb['Scenarios'],min_col=6,min_row=1,max_row=4),titles_from_data=True);chart.set_categories(Reference(wb['Scenarios'],min_col=1,min_row=2,max_row=4));chart.width=20;chart.height=9;wb['Dashboard'].add_chart(chart,'D3')
wb['Sensitivity'].conditional_formatting.add('B2:D4',ColorScaleRule(start_type='min',start_color='F8696B',mid_type='percentile',mid_value=50,mid_color='FFEB84',end_type='max',end_color='63BE7B'))
save()
print('Enhanced model saved; SAVE_NOW:',(ROOT/'SAVE_NOW.txt').exists())
if (ROOT/'SAVE_NOW.txt').exists(): raise SystemExit('SAVE_NOW: enhanced spreadsheet saved; stopped')
wb.create_sheet('Capacity options')
row('Capacity options',1,['Design','Normal rooms','Shared studios','Single studios','Beds','Gross sqm','Gross annual rent','Project cost ex interest','Funding incl interest','Stabilised NOI','Yield on cost','DSCR','Demand screen'])
for n,v in enumerate([['Compact',40,12,16],['Proposed',60,20,20],['Larger',80,28,24]],2):
 row('Capacity options',n,v)
 fs={5:f'=B{n}+2*C{n}+D{n}',6:f'=(B{n}*Inputs!B25+C{n}*Inputs!B26+D{n}*Inputs!B27)*Inputs!B28',7:f"=(B{n}*'Rent rules'!F2+C{n}*2*'Rent rules'!F3+D{n}*'Rent rules'!F4)*Inputs!B8",8:f'=Inputs!B14*(1+Inputs!B32)+F{n}*Inputs!B13*(1+Inputs!B15)*(1+Inputs!B16)*(1+Inputs!B17)+E{n}*Inputs!B33+Inputs!B34',9:f'=H{n}*(1+Inputs!B18*Inputs!B19)',10:f'=(G{n}*Inputs!B9*(1-Inputs!B11)-E{n}*Inputs!B12)*(1+Inputs!B10)',11:f'=J{n}/I{n}',12:f'=J{n}/(-PMT(Inputs!B19,Inputs!B20,H{n}*Inputs!B18))',13:f'=IF(E{n}<=\'Demand & competition\'!B10,"Within assumed demand","Above assumed demand")'}
 for c,f in fs.items():wb['Capacity options'].cell(n,c,f)
row('Capacity options',7,['Selection rationale','120 beds is the largest tested design within the assumed local demand envelope (~151 beds). The 80-bed design limits lease-up exposure; 160 beds exceeds the envelope.'])
row('Capacity options',8,['Financial conclusion','Demand alone does not justify construction. All sizes require yield and debt-coverage review; base 120-bed case fails the investment gate.'])
row('Checks',11,['Capacity option cost reconciliation',"='Capacity options'!I3-Development!B7",0])
for c,v in enumerate(['Interest paid','Principal repaid','Calendar year','DSCR'],15):wb['Cash flow'].cell(1,c,v)
for n in range(2,24):
 y=n-2
 wb['Cash flow'].cell(n,17,f'=A{n}+2027')
 if y>=2:
  wb['Cash flow'].cell(n,10,f'=IF(A{n}-1<=Inputs!B20,Development!B10,0)')
  wb['Cash flow'].cell(n,15,f'=L{n-1}*Inputs!B19')
  wb['Cash flow'].cell(n,16,f'=J{n}-O{n}')
  wb['Cash flow'].cell(n,18,f'=IF(J{n}>0,G{n}/J{n},0)')
 else:
  wb['Cash flow'].cell(n,15,0);wb['Cash flow'].cell(n,16,0);wb['Cash flow'].cell(n,18,0)
# Finish presentation with legible notes, input shading, filters and print setup.
for ws in wb:
 ws.sheet_properties.pageSetUpPr.fitToPage=True
 ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
 ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
 ws.print_title_rows='1:1'
 ws.sheet_properties.outlinePr.summaryRight=False
 for r in range(1,ws.max_row+1):ws.row_dimensions[r].height=30
 if ws.title in ['Read me']:
  ws.column_dimensions['A'].width=145
  for r in range(1,19):ws.row_dimensions[r].height=36
 if ws.title=='Sources':
  ws.column_dimensions['B'].width=65;ws.column_dimensions['C'].width=85;ws.column_dimensions['D'].width=30
  for r in range(2,11):ws.row_dimensions[r].height=76
 if ws.title in ['Inputs','Dashboard','Checks','Development']:ws.column_dimensions['A'].width=46
 if ws.title=='Inputs':ws.column_dimensions['C'].width=65
 if ws.title=='Dashboard':
  ws.column_dimensions['B'].width=43
  for r in [14,16,21,23]:ws.row_dimensions[r].height=60
 if ws.title=='Demand & competition':
  ws.column_dimensions['C'].width=65;ws.column_dimensions['E'].width=80
  for r in range(18,25):ws.row_dimensions[r].height=47
 if ws.title=='Rent rules':
  ws.column_dimensions['G'].width=38
  for r in range(6,12):
   ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=7);ws.row_dimensions[r].height=43
 if ws.title=='Capacity options':
  for r in [7,8]:ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=13);ws.row_dimensions[r].height=45
 if ws.title=='Sensitivity':ws.merge_cells('A6:D6');ws.row_dimensions[6].height=65
 if ws.title=='Scenarios':ws.merge_cells('A7:H7');ws.row_dimensions[7].height=65
 if ws.title=='Cash flow':
  for r in [25,26]:ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=14)
  ws.column_dimensions.group('O','R',hidden=False)
 for c in ws[1]:c.alignment=Alignment(vertical='center',wrap_text=True)
 ws.row_dimensions[1].height=48
for wsname,coords in [('Inputs',[f'B{r}' for r in range(2,37) if r not in [24,30]]),('Rent rules',[f'{c}{r}' for c in ['C','D'] for r in range(2,5)]),('Scenarios',[f'{c}{r}' for c in ['B','C','D','E'] for r in range(2,5)])]:
 for a in coords:wb[wsname][a].fill=PatternFill('solid',fgColor='E8F2FC')
# save() applies consistent base styles; apply number formats afterwards then atomic save.
save()
for r in [9,10,11,15,16,17,18,19,21,22,23,29,31,32,35]:wb['Inputs'].cell(r,2).number_format='0.0%'
for r in [3,20]:wb['Dashboard'].cell(r,2).number_format='0.0%' if r==20 else '0'
for r in [9]:wb['Demand & competition'].cell(r,2).number_format='0.0%'
for r in range(2,5):
 for c in [2,3,4,5,7]:wb['Scenarios'].cell(r,c).number_format='0.0%'
 wb['Scenarios'].cell(r,8).number_format='0.00"x"'
 wb['Capacity options'].cell(r,11).number_format='0.0%'
 wb['Capacity options'].cell(r,12).number_format='0.00"x"'
for c in range(2,5):wb['Sensitivity'].cell(1,c).number_format='0%'
for r in range(2,5):wb['Sensitivity'].cell(r,1).number_format='0%'
for n in range(2,24):
 wb['Cash flow'].cell(n,3).number_format='0.0%';wb['Cash flow'].cell(n,17).number_format='0';wb['Cash flow'].cell(n,18).number_format='0.00"x"'
wb['Dashboard']['B13'].number_format='0.00"x"'
wb.active=wb.sheetnames.index('Dashboard')
tmp=ROOT/'model.tmp.xlsx';wb.save(tmp);load_workbook(tmp).close();os.replace(tmp,ROOT/'model.xlsx')
print('Final model saved; SAVE_NOW:',(ROOT/'SAVE_NOW.txt').exists())
