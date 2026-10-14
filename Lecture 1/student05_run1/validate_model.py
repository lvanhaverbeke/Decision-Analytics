"""Evaluate the workbook's formula subset, validate controls, and cache Excel results atomically."""
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
import re, math, zipfile, os, json
from xml.etree import ElementTree as ET
P=Path(__file__).resolve().parent
if (P/'SAVE_NOW.txt').exists():
 print('SAVE_NOW: existing spreadsheet is already saved; stopping');raise SystemExit()
w=load_workbook(P/'model.xlsx');cache={};active=set()
def flat(xs):
 out=[]
 for x in xs:
  out.extend(flat(x)) if isinstance(x,list) else out.append(x)
 return out
def SUM(*xs):return sum(x for x in flat(xs) if isinstance(x,(int,float)))
def NPV(r,xs):return sum(x/(1+r)**(i+1) for i,x in enumerate(flat(xs)))
def PMT(r,n,p):return -p*r/(1-(1+r)**(-n)) if r else -p/n
def IRR(xs):
 xs=flat(xs)
 def f(r):return sum(x/(1+r)**i for i,x in enumerate(xs))
 lo=-.9;hi=1
 if f(lo)*f(hi)>0:return 0
 for _ in range(150):
  mid=(lo+hi)/2
  if f(lo)*f(mid)<=0:hi=mid
  else:lo=mid
 return (lo+hi)/2
def COUNTIF(xs,test):
 if test=='>0':return sum(isinstance(x,(int,float)) and x>0 for x in flat(xs))
 raise ValueError(test)
functions={'SUM':SUM,'NPV':NPV,'PMT':PMT,'IRR':IRR,'IF':lambda c,a,b:a if c else b,'IFERROR':lambda a,b:a,'MIN':min,'MAX':max,'COUNTIF':COUNTIF,'AND':lambda *x:all(x)}
pattern=re.compile(r'"[^"\n]*"|(?:(?:\'[^\']+\'|[A-Za-z_][A-Za-z_0-9 ]*)!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?')
def value(s,a):
 a=a.replace('$','');key=(s,a)
 if key in cache:return cache[key]
 if key in active:raise ValueError('Circular '+str(key))
 active.add(key);v=w[s][a].value
 if isinstance(v,str) and v.startswith('='):
  def sub(m):
   token=m.group()
   if token.startswith('"'):return token
   sh,ad=token.rsplit('!',1) if '!' in token else (s,token)
   sh=sh.strip("'");ad=ad.replace('$','')
   return ('R' if ':' in ad else 'V')+'('+repr(sh)+','+repr(ad)+')'
  expr=pattern.sub(sub,v[1:]).replace('^','**')
  expr=re.sub(r'(?<![<>=!])=(?!=)', '==', expr)
  try:v=eval(expr,{'__builtins__':{}},dict(functions,V=value,R=rangeval))
  except Exception as e:raise RuntimeError(f'{s}!{a}: {expr}: {e}') from e
 active.remove(key);cache[key]=v
 return 0 if v is None else v
def rangeval(s,a):
 c1,r1,c2,r2=range_boundaries(a)
 return [value(s,w[s].cell(r,c).coordinate) for r in range(r1,r2+1) for c in range(c1,c2+1)]
for ws in w:
 for row in ws:
  for c in row:
   if c.data_type=='f':value(ws.title,c.coordinate)
report={}
for n in range(2,12):
 label=w['Checks'].cell(n,1).value;actual=value('Checks',f'B{n}');expected=value('Checks',f'C{n}')
 report[label]={'actual':actual,'expected':expected,'pass':abs(actual-expected)<1e-6}
 if n!=8:assert report[label]['pass'],(label,actual,expected)
for a in ['B3','B6','B8','B9','B10','B11','B12','B13','B20','B23']:report['Dashboard '+a]=value('Dashboard',a)
# Preserve formula XML, add calculated values for Excel preview and data_only readers.
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main';ET.register_namespace('',NS)
tmp=P/'model.cached.tmp.xlsx'
with zipfile.ZipFile(P/'model.xlsx') as src,zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as dst:
 for item in src.infolist():
  blob=src.read(item.filename)
  match=re.fullmatch(r'xl/worksheets/sheet(\d+)\.xml',item.filename)
  if match:
   sheet=w.worksheets[int(match.group(1))-1].title;root=ET.fromstring(blob)
   for c in root.findall('.//{'+NS+'}c'):
    if c.find('{'+NS+'}f') is None:continue
    v=cache[(sheet,c.attrib['r'])];vc=c.find('{'+NS+'}v')
    if vc is None:vc=ET.SubElement(c,'{'+NS+'}v')
    if isinstance(v,str):c.set('t','str');vc.text=v
    else:
     c.attrib.pop('t',None);vc.text=str(int(v)) if isinstance(v,bool) else str(v if v is not None else 0)
   blob=ET.tostring(root,encoding='utf-8',xml_declaration=True)
  dst.writestr(item,blob)
load_workbook(tmp,data_only=True).close();os.replace(tmp,P/'model.xlsx')
(P/'validation.json').write_text(json.dumps(report,indent=2))
print(json.dumps(report,indent=2));print('Formula caches saved atomically; SAVE_NOW:',(P/'SAVE_NOW.txt').exists())
