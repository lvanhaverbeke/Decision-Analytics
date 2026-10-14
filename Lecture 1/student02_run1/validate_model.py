from pathlib import Path
from openpyxl import load_workbook
import math,zipfile,json
p=Path('model.xlsx');w=load_workbook(p,data_only=True);f=load_workbook(p,data_only=False)
flows=[w['Cash flow'].cell(r,12).value for r in range(3,35)]
discount=w['Inputs']['B15'].value
npv=sum(x/(1+discount)**t for t,x in enumerate(flows))
assert math.isclose(npv,w['Overview']['B5'].value,abs_tol=.001)
irr=w['Overview']['B6'].value
assert abs(sum(x/(1+irr)**t for t,x in enumerate(flows)))<.001
assert math.isclose(sum(w['Cash flow'].cell(r,3).value for r in range(3,35)),-w['Inputs']['B31'].value,abs_tol=.001)
assert math.isclose(w['Rent sensitivity']['F3'].value,npv,abs_tol=.001)
assert math.isclose(w['Risk sensitivity']['C5'].value,npv,abs_tol=.001)
assert math.isclose(w['Overview']['B7'].value,sum(w['Cash flow'].cell(r,18).value for r in range(3,35)),abs_tol=.001)
assert math.isclose(w['Overview']['B25'].value,sum(w['Cash flow'].cell(r,19).value for r in range(3,35)),abs_tol=.001)
assert w['Overview']['B13'].value==.16
assert len(f['Rent sensitivity']._charts)==1
assert len(f['Rent sensitivity']._charts[0]._charts)==2
errors=[(s.title,c.coordinate,c.value) for s in w for row in s for c in row if c.data_type=='e']
assert not errors,errors
assert not any('#REF!' in str(c.value) for s in f for row in s for c in row)
assert len([c for s in f for row in s for c in row if c.data_type=='f'])>400
with zipfile.ZipFile(p) as z:assert z.testzip() is None
result={'status':'PASS','sheets':w.sheetnames,'NPV_EUR':npv,'IRR':irr,'simple_PBP_years':w['Overview']['B7'].value,'discounted_PBP':w['Overview']['B26'].value,'formula_errors':errors,'note':'Python independently checks cached economics and workbook structure; no Excel or LibreOffice engine available for native recalculation.','SAVE_NOW':Path('SAVE_NOW.txt').exists()}
Path('validation.json').write_text(json.dumps(result,indent=2));print(json.dumps(result,indent=2))
