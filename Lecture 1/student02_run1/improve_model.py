from pathlib import Path
import json
p=Path('research.json');d=json.loads(p.read_text())
d['sources'] += [
['S10','VUB: current practical housing information','https://www.vub.be/en/studying-vub/practical-info-for-students/student-housing/practical-housing-information','2026–2027; updated 2026-09-18; accessed 2026-10-14','Standard VUB room EUR449.30/month including energy, internet, fire insurance, repairs and communal cleaning. Demand exceeds availability; allocation conditions apply. Not a private new-build market comparable.'],
['S11','FPS Employment: student employment remuneration','https://emploi.belgique.be/fr/themes/contrats-de-travail/contrats-de-travail-particuliers/contrat-doccupation-detudiants','Rates effective 2026-07-01; accessed 2026-10-14','Sector collective agreement governs pay. Where not applicable, age-adjusted interprofessional minimum; indicative hourly rates at 38h/week: age18 EUR10.7158;19 EUR11.5297;20 EUR12.2079;21 EUR13.5644. Annual student income depends on actual hours.']]
d['inputs']['10']['note']='EUR650 private new-build rent assumption, excluding utilities. S10: VUB standard room 2026–2027 EUR449.30 including listed costs; S2 historical private average EUR550. Premium and achievable rent require local comparables.'
p.write_text(json.dumps(d,indent=2,ensure_ascii=False))
p=Path('build_model.py');s=p.read_text()
s=s.replace("'Regulations','Sources']", "'Regulations','Sources','Student budget','Risk sensitivity']")
s=s.replace("'First operating cash flow at end of year 2; delays unmodelled'", "'Fixed schedule: first operating CF at end of year 2; rebuild code to change timing'")
s=s.replace("'First operating cash flow", "'First operating cash flow")
s=s.replace("for sh in wb:\n", '''# Supplementary cash-flow controls and formula-based payback calculations.
cf.cell(2,16,'Cumulative CF excl. exit');cf.cell(2,17,'Cumulative discounted CF');cf.cell(2,18,'PBP crossing');cf.cell(2,19,'Operating PBP crossing');cf.cell(2,20,'Discounted PBP crossing')
noexit=0;dpv=0;discount_pb=None
for r in range(3,35):
 t=r-3;previous_noexit=noexit;previous_dpv=dpv
 terminal_value=cache[('Cash flow',f'K{r}')];noexit+=flows[t]-terminal_value;dpv+=pvs[t]
 put('Cash flow',f'P{r}',f'=SUM(L$3:L{r})-SUM(K$3:K{r})',noexit)
 put('Cash flow',f'Q{r}',f'=SUM(N$3:N{r})',dpv)
 for col,prevcol,amountcol,cur,prev in [('R','O','L',sum(flows[:t+1]),sum(flows[:t])),('S','P','J',noexit,previous_noexit),('T','Q','N',dpv,previous_dpv)]:
  crossing=t-1-prev/(flows[t] if col=='R' else flows[t]-terminal_value if col=='S' else pvs[t]) if t>0 and cur>=0 and prev<0 else 0
  formula=f'=IF(AND({prevcol}{r}>=0,{prevcol}{r-1}<0),A{r}-1-{prevcol}{r-1}/{amountcol}{r},0)' if t>0 else '=0'
  put('Cash flow',f'{col}{r}',formula,crossing)
  if col=='T' and crossing:discount_pb=crossing
put('Overview','B7', '=IF(COUNTIF(\\'Cash flow\\'!R3:R34,">0")=0,"Not reached",SUM(\\'Cash flow\\'!R3:R34))',payback or 'Not reached')
put('Overview','B25','=IF(COUNTIF(\\'Cash flow\\'!S3:S34,">0")=0,"Not reached",SUM(\\'Cash flow\\'!S3:S34))',oppb or 'Not reached')
put('Overview','B26','=IF(COUNTIF(\\'Cash flow\\'!T3:T34,">0")=0,"Not reached",SUM(\\'Cash flow\\'!T3:T34))',discount_pb or 'Not reached')
ov['A28']='Minimum proof of means / year';put('Overview','B28',"='Student budget'!B5*12",12744);ov['C28']='Regulatory benchmark; continuous; ratio; EUR; certain non-EU students (S3)'
ov['A29']='Minimum student wage, age 21';put('Overview','B29',"='Student budget'!B12",13.5644);ov['C29']='Regulatory benchmark; continuous; ratio; EUR/hour; fallback absent sector CBA (S11)'
ov['A30']='Illustrative 650h gross pay, age 21';put('Overview','B30',"='Student budget'!C12",13.5644*650);ov['C30']='Derived scenario; continuous; ratio; EUR/year; not VUB average'
ov['A32']='Model controls';ov['B32']='Blue cells are assumptions. Development duration and horizon are fixed to 2 and 30 years; rebuild the saved code for another schedule. No debt, grants, income tax or financing fees; costs gross of nonrecoverable tax.'
budget=wb['Student budget'];budget.append(['STUDENT RESOURCES AND AFFORDABILITY']);budget.append(['Item / age','Value / hourly pay EUR','Annual gross EUR','After solidarity contribution EUR','Classification / evidence'])
budget['A4']='Illustrative annual total resources';put('Student budget','B4','=Inputs!B21',v[21]);budget['E4']='Scenario; continuous; ratio. Family support + grants + work; no observed VUB average.'
budget['A5']='Monthly proof of means';budget['B5']=1062;budget['E5']='S3: 2026–2027 benchmark for relevant non-EU student residence applicants; not wages.'
budget['A6']='Reduced-contribution annual hours';budget['B6']=650;budget['E6']='S4: count, discrete, ratio; scenario assumes entire quota worked, not guaranteed employment.'
budget['A7']='Student solidarity contribution';budget['B7']=.0271;budget['E7']='S4: continuous ratio; 2.71%; tax treatment depends on personal circumstances.'
for rr,age,rate in [(9,18,10.7158),(10,19,11.5297),(11,20,12.2079),(12,21,13.5644)]:
 budget.cell(rr,1,age);budget.cell(rr,2,rate);put('Student budget',f'C{rr}',f'=B{rr}*$B$6',rate*650);put('Student budget',f'D{rr}',f'=C{rr}*(1-$B$7)',rate*650*(1-.0271));budget.cell(rr,5,'S11: fallback hourly rate effective July 2026, indicative 38h-week basis; sector CBA may apply. Gross annual pay assumes rate unchanged for all 650h; D is before any income tax.')
budget['A15']='VUB standard rent incl. listed costs';budget['B15']=449.30;budget['E15']='S10: EUR/month; continuous ratio; 2026–2027; subsidised university offering with eligibility conditions.'
budget['A16']='Project rent excl. utilities';put('Student budget','B16','=Inputs!B10',v[10]);budget['E16']='Decision; continuous ratio; EUR/month.'
budget['A17']='Rent share of proof-of-means amount';put('Student budget','B17','=B16/B5',v[10]/1062);budget['E17']='Derived; continuous ratio; excludes utilities, tuition and other living expenses.'
budget['A19']='Minimum annual earnings';budget['B19']='No guaranteed amount: a student with no paid work may earn EUR0. Actual average VUB earnings were not established.'
risk=wb['Risk sensitivity'];risk.append(['NPV SENSITIVITY — required return versus exit capitalisation']);risk.append(['Discount rate / exit cap',.045,.055,.065])
for rr,discount in enumerate([.05,.06,.07,.08,.09],3):
 risk.cell(rr,1,discount)
 for cc,cap in enumerate([.045,.055,.065],2):
  terminal=cache[('Cash flow','H34')]*(1+v[11])/cap*(1-v[17]);value=sum(f/(1+discount)**t for t,f in enumerate(flows))+(terminal-cache[('Cash flow','K34')])/(1+discount)**31
  letter=chr(64+cc);formula=f"=NPV($A{rr},'Cash flow'!L4:L34)+'Cash flow'!L3+('Cash flow'!H34*(1+Inputs!B11)/{letter}$2*(1-Inputs!B17)-'Cash flow'!K34)/(1+$A{rr})^31"
  put('Risk sensitivity',f'{letter}{rr}',formula,value)
risk['A10']='Break-even opening rent, no extra departures';rent_pv=sum(pvs[2:])+sum(cache[('Cash flow',f'I{r}')]*cache[('Cash flow',f'M{r}')] for r in range(5,35));breakeven=v[10]*(1-sum(pvs)/rent_pv)
put('Risk sensitivity','B10',"=Inputs!B10*(1-Overview!B5/(SUM('Cash flow'!N5:N34)+SUMPRODUCT('Cash flow'!I5:I34,'Cash flow'!M5:M34)))",breakeven)
risk['A12']='Interpretation';risk['B12']='Discount and exit cap are continuous ratio assumptions; NPV is a derived continuous EUR output. Break-even rent holds occupancy fixed; verify affordability, lawfulness and departures separately.'
from openpyxl.worksheet.datavalidation import DataValidation
for rr,num in [(18,2),(19,30)]:
 fixed=DataValidation(type='whole',operator='equal',formula1=str(num));fixed.errorTitle='Fixed schedule';fixed.error='Rebuild build_model.py to change this schedule assumption.';fixed.showErrorMessage=True;fixed.errorStyle='stop';wb['Inputs'].add_data_validation(fixed);fixed.add(f'B{rr}')
for sh in wb:
''')
s=s.replace("wb.calculation=CalcProperties", """for rr in [18,19]:
 wb['Inputs'].cell(rr,2).font=Font(color='595959');wb['Inputs'].cell(rr,2).fill=PatternFill('solid',fgColor='EEEEEE')
wb['Inputs']['E18']='Fixed 2-year schedule incl. approvals; change code and rebuild for alternative timing.'
wb['Student budget']['B7'].number_format='0.00%';wb['Student budget']['B17'].number_format='0.0%'
for rr in range(3,8):wb['Risk sensitivity'].cell(rr,1).number_format='0.0%'
for cc in range(2,5):wb['Risk sensitivity'].cell(2,cc).number_format='0.0%'
wb.calculation=CalcProperties""")
p.write_text(s)
