from pathlib import Path
import math,json
from openpyxl import load_workbook
P=Path(__file__).resolve().parent
f=load_workbook(P/'model.xlsx',data_only=False);v=load_workbook(P/'model.xlsx',data_only=True)
errors=[];count=0
for s in f:
 for row in s:
  for c in row:
   if c.data_type=='f':
    count+=1;x=v[s.title][c.coordinate].value
    if not isinstance(x,(int,float)) or not math.isfinite(x):errors.append((s.title,c.coordinate,x))
assert not errors,errors
assert abs(v['Sources & Uses']['B15'].value)<1e-8
for col in range(2,12):
 d=v['Debt'];assert abs(d.cell(2,col).value-d.cell(5,col).value-d.cell(6,col).value)<1e-6
assert abs(v['Sensitivity']['C3'].value-v['Returns']['B10'].value)<1e-6
cash=[v['Returns'].cell(8,c).value for c in range(2,13)];r=v['Returns']['B9'].value
assert abs(sum(x/(1+r)**i for i,x in enumerate(cash)))<1e-5
# Independent base-case cross-check using explicit inputs.
rev=40*650*12*.95;noi=rev*.94-18000-14000-6000-14000-5000
payment=1920000*.045/(1-(1.045)**-25)
assert abs(noi-v['Operations']['B13'].value)<1e-6
assert abs(noi-12000-payment-v['Returns']['C3'].value)<1e-6
assert abs(v['Sources & Uses']['B10'].value-4686320)<1e-6
result={'status':'PASS','formula_cells_with_numeric_caches':count,'checks':['sources equal uses','debt balances reconcile','base sensitivity equals NPV','IRR residual zero','independent Year 1 NOI and cash-flow calculation','project cost cross-check'],'excel_recalculation':'No Excel/LibreOffice runtime available; formulas and cached values checked with Python.'}
(P/'validation.json').write_text(json.dumps(result,indent=2));print(json.dumps(result));print('SAVE_NOW:',(P/'SAVE_NOW.txt').exists())
