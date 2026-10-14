from pathlib import Path
import os, json, math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
ROOT=Path(__file__).resolve().parent
w=Workbook(); w.remove(w.active)
for n in ['Read Me','Assumptions','Sources & Uses','Operations','Debt','Returns','Sensitivity','Legal Checklist','Sources'] : w.create_sheet(n)
A=w['Assumptions']; A.append(['Assumption','Value','Unit','Basis / action needed'])
rows=[
('Rooms',40,'rooms','Illustrative; confirm lawful room capacity'),
('Purchase price',3200000,'EUR','Illustrative; obtain appraisal and vendor terms'),
('Registration duty',0.125,'%','Provisional Brussels investment acquisition rate; confirm transaction structure'),
('Other acquisition costs',0.02,'% of purchase','Notary, diligence, adviser allowance; excludes duty'),
('Refurbishment per room',18000,'EUR/room','Illustrative contractor budget; gross including nonrecoverable VAT'),
('Capex contingency',0.10,'% of refurbishment','Illustrative'),
('Setup & professional costs',80000,'EUR','Architect, permit and setup allowance, gross'),
('Initial reserve',100000,'EUR','Funded at closing; released at sale'),
('Monthly rent per room',650,'EUR','Illustrative base rent; excludes tenant utility reimbursements'),
('Occupancy',0.95,'%','Economic occupancy; includes voids and rent losses'),
('Paid months per year',12,'months','Confirm lease terms; test 10 months'),
('Other annual income',0,'EUR','No assumed ancillary income'),
('Annual rent growth',0.02,'%','Forecast only; legal indexation conditions must be met'),
('Management fee',0.06,'% of effective revenue','External manager allowance'),
('Annual utilities/common areas',18000,'EUR','Owner borne; no assumed recovery'),
('Annual maintenance',14000,'EUR','Owner borne'),
('Annual insurance',6000,'EUR','Illustrative quotation required'),
('Annual property tax',14000,'EUR','Placeholder; confirm indexed cadastral income and municipal charges'),
('Annual administration',5000,'EUR','Accounting, compliance and letting'),
('Annual cost growth',0.025,'%','Illustrative'),
('Annual replacement reserve',12000,'EUR','Deducted below NOI'),
('Debt / purchase price',0.60,'%','Debt based on purchase price only; no capex financing'),
('Loan interest',0.045,'%','Illustrative; lender quotation required'),
('Amortisation term',25,'years','Annual payments approximation'),
('Debt arrangement fee',0.01,'% of debt','Cash use at closing'),
('Holding period',10,'years','Fixed 10-year projection'),
('Exit cap rate',0.055,'%','Illustrative valuation assumption; not market evidence'),
('Selling costs',0.02,'% of sale price','Illustrative'),
('Equity discount rate',0.09,'%','Illustrative required return'),
('Construction delay',0,'months','Year 1 rent reduction; 0 means stabilised at acquisition'),
]
for r in rows:A.append(r)
def af(r): return f"'Assumptions'!$B${r}"
S=w['Sources & Uses']; S.append(['Sources / uses','EUR'])
for r in [('Purchase',f'={af(3)}'),('Registration duty',f'={af(3)}*{af(4)}'),('Other acquisition costs',f'={af(3)}*{af(5)}'),('Refurbishment',f'={af(2)}*{af(6)}'),('Contingency',f'=B5*{af(7)}'),('Setup costs',f'={af(8)}'),('Initial reserve',f'={af(9)}'),('Debt fee',f'=B12*{af(26)}'),('Total uses','=SUM(B2:B9)'),('',''),('Debt proceeds',f'={af(3)}*{af(23)}'),('Equity required','=B10-B12'),('Total sources','=SUM(B12:B13)'),('Funding check','=B14-B10')]: S.append(r)
O=w['Operations']; O.append(['Year']+list(range(1,12)))
labels=['Potential room rent','Effective room rent','Other revenue','Total revenue','Management','Utilities','Maintenance','Insurance','Property tax','Administration','Total operating costs','Net operating income (NOI)','Replacement reserve','Cash flow before debt']
for lab in labels: O.append([lab])
for col in range(2,13):
 c=get_column_letter(col); y=f'{c}$1'; prev=get_column_letter(col-1)
 formulas={2:f'{af(2)}*{af(10)}*{af(12)}*(1+{af(14)})^({y}-1)',3:f'{c}2*{af(11)}'+(f'*(1-{af(31)}/12)' if col==2 else ''),4:f'{af(13)}*(1+{af(14)})^({y}-1)',5:f'SUM({c}3:{c}4)',6:f'{c}5*{af(15)}',7:f'{af(16)}*(1+{af(21)})^({y}-1)',8:f'{af(17)}*(1+{af(21)})^({y}-1)',9:f'{af(18)}*(1+{af(21)})^({y}-1)',10:f'{af(19)}*(1+{af(21)})^({y}-1)',11:f'{af(20)}*(1+{af(21)})^({y}-1)',12:f'SUM({c}6:{c}11)',13:f'{c}5-{c}12',14:f'{af(22)}*(1+{af(21)})^({y}-1)',15:f'{c}13-{c}14'}
 for r,f in formulas.items():O.cell(r,col,'='+f)
D=w['Debt']; D.append(['Year']+list(range(1,11)))
for l in ['Opening balance','Debt service','Interest','Principal','Closing balance','DSCR']:D.append([l])
for col in range(2,12):
 c=get_column_letter(col); p=get_column_letter(col-1)
 for r,f in {2:"'Sources & Uses'!B12" if col==2 else f'{p}6',3:f'MIN({c}2*(1+{af(24)}),-PMT({af(24)},{af(25)},\'Sources & Uses\'!$B$12))',4:f'{c}2*{af(24)}',5:f'{c}3-{c}4',6:f'MAX(0,{c}2-{c}5)',7:f"'Operations'!{c}13/{c}3"}.items():D.cell(r,col,'='+f)
R=w['Returns'];R.append(['Year']+list(range(0,11)))
for l in ['Initial equity','Cash after debt','Gross sale value','Selling costs','Debt repayment','Reserve released','Total equity cash flow']:R.append([l])
R['B2']="=-'Sources & Uses'!B13";R['B8']='=B2'
for col in range(3,13):
 c=get_column_letter(col); op=get_column_letter(col-1)
 R.cell(3,col,f"='Operations'!{op}15-'Debt'!{op}3")
 if col==12:
  for r,f in {4:f"'Operations'!L13/{af(28)}",5:f'{c}4*{af(29)}',6:"'Debt'!K6",7:af(9)}.items():R.cell(r,col,'='+f)
 R.cell(8,col,f'={c}3+{c}4-{c}5-{c}6+{c}7')
for r in [('Equity IRR','=IRR(B8:L8)'),('Equity NPV',f'=NPV({af(30)},C8:L8)+B8'),('Equity multiple','=SUM(C8:L8)/-B8'),('Year 1 cash yield','=C3/-B8'),('Year 1 NOI yield on cost',"='Operations'!B13/'Sources & Uses'!B10"),('Year 1 DSCR',"='Debt'!B7"),('Break-even occupancy (cash)',f"=('Operations'!B12-'Operations'!B6+'Operations'!B14+'Debt'!B3)/('Operations'!B2*(1-{af(15)}))")]:R.append(r)
T=w['Sensitivity'];T.append(['Exit cap rate / rent change',-0.10,0,0.10]);
for i,cap in enumerate([0.045,0.055,0.065],2):
 T.cell(i,1,cap)
 for col in range(2,5):
  c=get_column_letter(col)
  # NPV change: incremental rental margin each year + changed terminal value
  inc='+'.join([f"('Operations'!{get_column_letter(y+1)}3*{c}$1*(1-{af(15)}))/(1+{af(30)})^{y}" for y in range(1,11)])
  term=f"((('Operations'!L13+'Operations'!L3*{c}$1*(1-{af(15)}))/$A{i})*(1-{af(29)})-'Returns'!L4+'Returns'!L5)/(1+{af(30)})^10"
  T.cell(i,col,f"='Returns'!B10+{inc}+{term}")
T.append(['Values show equity NPV (EUR), with fixed costs and debt unchanged.'])
L=w['Legal Checklist']; L.append(['Topic','Required verification','Status','Source'])
legal=[('Municipality & zoning','Identify exact address, commune, lawful use and permitted room count. VUB/ULB campus area may involve Ixelles or Etterbeek, not necessarily City of Brussels.','OPEN',''),('Planning permission','Check conversion/change of use, works, planning permit and local rules before purchase.','OPEN',''),('Housing quality','Confirm Brussels minimum safety, health and equipment standards; inspect every unit.','OPEN',''),('Fire safety','Obtain applicable fire service/municipal requirements, occupancy limit and certificates.','OPEN',''),('Student tenancy','Check student lease regime, student status, duration, termination, deposit and registration.','OPEN',''),('Rent & indexation','Check reference rent / excessive rent rules and EPC-related indexation conditions. Forecast growth is conditional.','OPEN',''),('Energy','Obtain valid PEB/EPC and price required renovation works; assess regional obligations.','OPEN',''),('Taxes & VAT','Confirm registration duty, property tax, ownership vehicle, VAT recoverability and income/capital gains taxes. Model is pre-income-tax.','OPEN','')]
for r in legal:L.append(r)
Q=w['Sources'];Q.append(['Source','URL','Access date','Use / reliability'])
Q.append(['Initial model','User request','2026-10-14','All financial assumptions illustrative, pending research'])
M=w['Read Me'];
for r in [['Student housing investment — VUB/ULB, Brussels'],['Status','Usable baseline; illustrative acquisition + refurbishment; not a confirmed property'],['Instructions','Edit blue cells in Assumptions; formulas recalculate in Excel. All currency EUR.'],['Investment basis','40 rooms, 10-year hold; annual debt payments; sale valued on forward Year 11 NOI.'],['Model scope','Pre-income-tax equity returns; costs and capex gross of nonrecoverable VAT.'],['Critical limitation','Exact address, planning permission, fire/quality compliance and investment taxes require verification.'],['Timing','Construction delay reduces Year 1 rent; full debt service and fixed costs apply throughout.'],['Reserve','Initial funded reserve released on exit; annual replacement reserve treated as spent.'],['Sensitivity','Equity NPV sensitivity to base rent and exit cap; no price evidence implied.'],['Decision','Do not commit until property-specific legal, technical and lender diligence is complete.']]:M.append(r)
if (ROOT/'research.json').exists():
 data=json.loads((ROOT/'research.json').read_text())
 for r in data.get('sources',[]):Q.append(r)
 for r in data.get('legal',[]):L.append(r)
 M['B2']='Research-informed screening model; financial inputs remain illustrative.'
for ws in w:
 ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
 ws.row_dimensions[1].height=28
 for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='183B56');cell.font=Font(color='FFFFFF',bold=True)
 for row in ws.iter_rows(min_row=2):
  for cell in row:
   cell.alignment=Alignment(vertical='top',wrap_text=True)
   if cell.data_type=='f':cell.font=Font(color='111111');cell.number_format='#,##0.00;[Red](#,##0.00)'
   elif isinstance(cell.value,(float,int)):cell.number_format='#,##0.00';cell.font=Font(color='0066CC' if ws==A else '111111')
 ws.column_dimensions['A'].width=34
 for col in range(2,ws.max_column+1):ws.column_dimensions[get_column_letter(col)].width=18 if ws in [O,D,R,T] else 32
 if ws in [L,Q]:ws.column_dimensions['B'].width=85;ws.column_dimensions['D'].width=65
 if ws==M:ws.column_dimensions['B'].width=110
for r in [4,5,7,11,14,15,21,23,24,26,28,29,30]:A.cell(r,2).number_format='0.0%'
for r in [9,12,13]:R.cell(r,2).number_format='0.0%'
R['B11'].number_format='0.00x';R['B14'].number_format='0.00x';R['B15'].number_format='0.0%'
for row in D.iter_rows(min_row=7,max_row=7,min_col=2):
 for c in row:c.number_format='0.00x'
for c in T[1][1:]:c.number_format='0%'
for r in range(2,5):T.cell(r,1).number_format='0.0%'
w.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
tmp=ROOT/'model.tmp.xlsx';w.save(tmp);os.replace(tmp,ROOT/'model.xlsx')
print('Saved model.xlsx atomically. SAVE_NOW:',(ROOT/'SAVE_NOW.txt').exists())
